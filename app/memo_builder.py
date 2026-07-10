from __future__ import annotations

import html
import filecmp
import csv
import re
import shutil
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from .cli_support import PROJECT_ROOT, case_path
from .workflow import extract_docx_text, folder_role_map, load_case, load_yaml_file, safe_path_component


DEFAULT_TEMPLATE_BY_TASK_TYPE = {
    "eb1a_petition": "templates/EB1A/EB1A_unified_template_LLM.docx",
    "eb1a_rfe_response": "templates/RFE/EB1/EB1A_RFE_response_unified_LLM_template.yaml",
    "o1b_petition": "templates/O1B/MEMO O-1В_ver.1.0.docx",
}

DEFAULT_WORKING_STRUCTURE_BY_TASK_TYPE = {
    "eb1a_petition": "templates/EB1A/EB1A_working_document_structure.yaml",
    "o1b_petition": "templates/O1B/O1B_working_document_structure.yaml",
}

TECHNICAL_MEMO_PREFIXES = (
    "Drafting direction (internal):",
    "[DRAFTING PLACEHOLDER:",
    "[SCRIPT PLACEHOLDER:",
    "[LLM SECTION PLACEHOLDER:",
    "[SCRIPT-CONTROLLED CONTENT:",
)

# Regulatory criteria only. Employment-plan components are assembled through
# their dedicated workflow destination and _draft_step_xml() path below.
EB1A_CRITERION_STEP_BY_ROLE = {
    "awards": ("criterion_awards_episode", "Awards"),
    "memberships": ("criterion_memberships_episode", "Memberships / associations"),
    "media": ("criterion_media_episode", "Published material"),
    "judging": ("criterion_judging_episode", "Judging"),
    "original_contribution": ("criterion_original_contribution_fact", "Original contribution"),
    "scholarly_articles": ("criterion_scholarly_articles_episode", "Scholarly articles"),
    "exhibitions": ("criterion_exhibitions_episode", "Exhibitions / showcases"),
    "leading_critical_role": ("criterion_leading_critical_role_fact", "Leading or critical role"),
    "high_salary": ("criterion_high_salary_fact", "High salary / remuneration"),
    "commercial_success": ("criterion_commercial_success_episode", "Commercial success"),
}

O1B_CRITERION_STEP_BY_ROLE = {
    "lead_starring_productions": ("o1b_criterion_i_episode", "Lead or starring productions/events"),
    "published_recognition": ("o1b_criterion_ii_episode", "Published recognition"),
    "organization_role": ("o1b_criterion_iii_role", "Lead, starring, or critical organizational role"),
    "commercial_critical_success": ("o1b_criterion_iv_episode", "Commercial or critically acclaimed success"),
    "significant_recognition": ("o1b_criterion_v_episode", "Significant recognition"),
    "high_salary": ("o1b_criterion_vi_compensation", "High salary or substantial remuneration"),
    "comparable_evidence": ("o1b_comparable_evidence_episode", "Comparable evidence (Arts only)"),
}

RFE_CRITERION_TEMPLATE_KEY_BY_ROLE = {
    "awards": "criterion_1_awards",
    "memberships": "criterion_2_memberships",
    "media": "criterion_3_published_material",
    "judging": "criterion_4_judging",
    "original_contribution": "criterion_5_original_contributions",
    "scholarly_articles": "criterion_6_scholarly_articles",
    "exhibitions": "criterion_7_display",
    "leading_critical_role": "criterion_8_leading_or_critical_role",
    "high_salary": "criterion_9_high_salary",
    "commercial_success": "criterion_10_commercial_success_performing_arts",
}

RFE_CRITERION_ROMAN_BY_ROLE = {
    "awards": "i",
    "memberships": "ii",
    "media": "iii",
    "judging": "iv",
    "original_contribution": "v",
    "scholarly_articles": "vi",
    "exhibitions": "vii",
    "leading_critical_role": "viii",
    "high_salary": "ix",
    "commercial_success": "x",
}

RFE_CRITERION_HEADING_BY_ROLE = {
    "awards": "Criterion 1. Evidence of receipt of lesser nationally or internationally recognized prizes or awards for excellence",
    "memberships": "Criterion 2. Evidence of petitioner's membership in associations in the field for which classification is sought that require outstanding achievement of their members, as judged by recognized national or international experts in their disciplines or fields.",
    "media": "Criterion 3. Published material about the petitioner in professional or major trade publications or other major media. The materials must relate to the petitioner's work in the field for which classification is sought.",
    "judging": "Criterion 4. Evidence of the petitioner's participation on a panel, or individually, as a judge of the work of others in the same or in an allied field of specialization for which classification is sought.",
    "original_contribution": "Criterion 5. Evidence of the petitioner's original scientific, scholarly, or business-related contributions of major significance in the field.",
    "scholarly_articles": "Criterion 6. The person's authorship of scholarly articles in the field, in professional or major trade publications or other major media.",
    "exhibitions": "Criterion 7. Display of the person's work in the field at artistic exhibitions or showcases.",
    "leading_critical_role": "Criterion 8. The person has performed in a leading or critical role for organizations or establishments that have a distinguished reputation.",
    "high_salary": "Criterion 9. The person has commanded a high salary, or other significantly high remuneration for services, in relation to others in the field.",
    "commercial_success": "Criterion 10. Commercial successes in the performing arts, as shown by box office receipts or record, cassette, compact disk, or video sales.",
}

# Backward-compatible alias for EB-1A-specific callers and tests.
CRITERION_STEP_BY_ROLE = EB1A_CRITERION_STEP_BY_ROLE


@dataclass(frozen=True)
class TemplateSection:
    level: int
    title: str
    line_number: int


@dataclass(frozen=True)
class TemplateParseReport:
    template_path: Path
    source_kind: str
    placeholder_count: int
    placeholders: list[str]
    section_count: int
    sections: list[TemplateSection]


@dataclass(frozen=True)
class WorkingMemoSummary:
    markdown_path: Path
    docx_path: Path
    template_report_path: Path
    sections_written: int
    placeholders_seen: int
    inferred_items: int


@dataclass(frozen=True)
class IntakeSummary:
    config_path: Path
    fields_updated: int
    source_files_copied: int
    source_files_skipped: int


def refresh_case_sources(case_id: str) -> IntakeSummary:
    """Refresh remembered external source folders without changing intake data."""
    loaded = load_case(case_id)
    config = dict(loaded.config)
    source_imports = config.get("source_imports", {})
    if not isinstance(source_imports, dict):
        source_imports = {}
    copied = 0
    skipped = 0
    for key in ("source_originals", "source_translations", "source_other"):
        raw_path = str(source_imports.get(key, "")).strip()
        if not raw_path:
            continue
        refreshed, unchanged = _copy_source_folder(
            Path(_clean_user_path(raw_path)),
            loaded.case_dir / _path_from_config(config, key),
            refresh_existing=True,
        )
        copied += refreshed
        skipped += unchanged
    return IntakeSummary(
        config_path=loaded.case_dir / "case_config.yaml",
        fields_updated=0,
        source_files_copied=copied,
        source_files_skipped=skipped,
    )


def apply_case_intake(
    case_id: str,
    fields: dict[str, str],
    *,
    case_info_file: str = "",
    source_folder_path: str = "",
    source_target_key: str = "source_originals",
    source_folder_paths: dict[str, str] | None = None,
    claimed_criteria: list[str] | None = None,
) -> IntakeSummary:
    loaded = load_case(case_id)
    config = dict(loaded.config)
    merged = dict(_parse_case_info_file(case_info_file))
    merged.update({key: value for key, value in fields.items() if value})
    updates = _normalize_intake_fields(merged)
    fields_updated = _apply_updates(config, updates)
    fields_updated += _apply_gender_defaults(config)
    if case_info_file.strip():
        intake_sources = config.get("intake_sources")
        if not isinstance(intake_sources, dict):
            intake_sources = {}
            config["intake_sources"] = intake_sources
        remembered_case_info = _clean_user_path(case_info_file)
        if intake_sources.get("case_info_file") != remembered_case_info:
            intake_sources["case_info_file"] = remembered_case_info
            fields_updated += 1
    if claimed_criteria is not None:
        criterion_map = _criterion_step_map(config)
        normalized_criteria = [role for role in criterion_map if role in claimed_criteria]
        if str(config.get("task_type", "")) == "o1b_petition" and str(config.get("o1b_track", "")) == "mptv":
            normalized_criteria = [role for role in normalized_criteria if role != "comparable_evidence"]
        if config.get("claimed_criteria") != normalized_criteria:
            config["claimed_criteria"] = normalized_criteria
            fields_updated += 1
    source_files_copied = 0
    source_files_skipped = 0
    source_requests = dict(source_folder_paths or {})
    if source_folder_path.strip():
        source_requests[source_target_key] = source_folder_path
    if source_folder_paths is not None:
        source_imports = config.get("source_imports")
        if not isinstance(source_imports, dict):
            source_imports = {}
            config["source_imports"] = source_imports
        for key, raw_path in source_folder_paths.items():
            cleaned_path = _clean_user_path(raw_path)
            if cleaned_path:
                if source_imports.get(key) != cleaned_path:
                    source_imports[key] = cleaned_path
                    fields_updated += 1
            elif key in source_imports:
                del source_imports[key]
                fields_updated += 1
    for key, raw_path in source_requests.items():
        cleaned_path = _clean_user_path(raw_path)
        if not cleaned_path:
            continue
        copied, skipped = _copy_source_folder(
            Path(cleaned_path),
            loaded.case_dir / _path_from_config(config, key),
        )
        source_files_copied += copied
        source_files_skipped += skipped

    config_path = loaded.case_dir / "case_config.yaml"
    _write_yaml_file(config_path, config)
    return IntakeSummary(
        config_path=config_path,
        fields_updated=fields_updated,
        source_files_copied=source_files_copied,
        source_files_skipped=source_files_skipped,
    )


def _clean_user_path(value: str) -> str:
    cleaned = value.strip()
    quote_pairs = {('"', '"'), ("'", "'"), ("“", "”"), ("«", "»")}
    while len(cleaned) >= 2 and (cleaned[0], cleaned[-1]) in quote_pairs:
        cleaned = cleaned[1:-1].strip()
    return cleaned


def build_working_memo(case_id: str, *, template_path: str = "") -> WorkingMemoSummary:
    loaded = load_case(case_id)
    template = _resolve_template_path(loaded.config, template_path)
    report = parse_machine_template(template)
    skeleton = build_memo_skeleton(loaded.config, report, loaded.case_dir)
    final_dir = loaded.case_dir / _path_from_config(loaded.config, "final_memo")
    final_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = final_dir / "working_memo.md"
    docx_path = final_dir / "working_memo.docx"
    report_path = loaded.case_dir / "reports" / "template_parse_report.md"
    report_path.parent.mkdir(parents=True, exist_ok=True)

    markdown_path.write_text(render_markdown_skeleton(loaded.case_id, loaded.config, skeleton), encoding="utf-8")
    report_path.write_text(render_template_report(report), encoding="utf-8")
    write_docx(docx_path, loaded.case_id, loaded.config, skeleton)
    return WorkingMemoSummary(
        markdown_path=markdown_path,
        docx_path=docx_path,
        template_report_path=report_path,
        sections_written=len(skeleton),
        placeholders_seen=report.placeholder_count,
        inferred_items=sum(1 for item in skeleton if item.get("inferred")),
    )


def parse_machine_template(path: Path) -> TemplateParseReport:
    text = _read_template_text(path)
    placeholders = sorted(
        set(re.findall(r"__[A-Z0-9_]+__", text))
        | set(re.findall(r"\[[A-Z][A-Z0-9_ /.-]{2,}\]", text))
    )
    sections = _parse_sections(text)
    return TemplateParseReport(
        template_path=path,
        source_kind=path.suffix.lower().lstrip(".") or "text",
        placeholder_count=len(placeholders),
        placeholders=placeholders,
        section_count=len(sections),
        sections=sections,
    )


def build_memo_skeleton(
    config: dict[str, Any], report: TemplateParseReport, case_dir: Path
) -> list[dict[str, Any]]:
    task_type = str(config.get("task_type", ""))
    if task_type == "eb1a_rfe_response":
        return _rfe_skeleton(config, case_dir, report)
    if task_type == "o1b_petition":
        return _o1b_skeleton(config, case_dir, report)
    return _eb1a_skeleton(config, case_dir, report)


def render_template_report(report: TemplateParseReport) -> str:
    lines = [
        "# Template parse report",
        "",
        f"- template: `{report.template_path}`",
        f"- source_kind: `{report.source_kind}`",
        f"- placeholders: {report.placeholder_count}",
        f"- sections: {report.section_count}",
        "",
        "## Placeholders",
        "",
    ]
    lines.extend(f"- `{placeholder}`" for placeholder in report.placeholders[:200])
    if len(report.placeholders) > 200:
        lines.append(f"- ... and {len(report.placeholders) - 200} more")
    lines.extend(["", "## Parsed sections", ""])
    for section in report.sections[:300]:
        indent = "  " * max(section.level - 1, 0)
        lines.append(f"- {indent}{section.title} (line {section.line_number})")
    if len(report.sections) > 300:
        lines.append(f"- ... and {len(report.sections) - 300} more")
    return "\n".join(lines) + "\n"


def render_markdown_skeleton(
    case_id: str, config: dict[str, Any], skeleton: list[dict[str, Any]]
) -> str:
    lines = [
        f"# Working memorandum: {case_id}",
        "",
        *(_metadata_lines(config)),
        "",
        "> This file is generated by scripts. Draft section placeholders are intentionally visible.",
        "",
    ]
    for item in skeleton:
        level = int(item.get("level", 1))
        heading = str(item.get("title", "Untitled section"))
        lines.append("#" * min(max(level, 1), 4) + " " + heading)
        lines.append("")
        for paragraph in item.get("paragraphs", []):
            lines.append(str(paragraph))
            lines.append("")
        for bullet in item.get("bullets", []):
            lines.append(f"- {bullet}")
        if item.get("bullets"):
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def write_docx(path: Path, case_id: str, config: dict[str, Any], skeleton: list[dict[str, Any]]) -> None:
    task_type = str(config.get("task_type", ""))
    if task_type == "eb1a_rfe_response":
        _write_rfe_company_docx(path, config, path.parent.parent)
        return
    if task_type == "eb1a_petition":
        document_xml = _eb1a_document_xml(config, path.parent.parent)
    elif task_type == "o1b_petition":
        document_xml = _o1b_document_xml(config, path.parent.parent)
    else:
        document_xml = _document_xml(case_id, config, skeleton)
    styles_xml = _o1b_styles_xml() if task_type == "o1b_petition" else _styles_xml()
    content_types = _content_types_xml(include_footer=task_type == "o1b_petition")
    rels = _rels_xml()
    doc_rels = _document_rels_xml(include_footer=task_type == "o1b_petition")
    app_xml = _app_xml()
    core_xml = _core_xml(case_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("docProps/app.xml", app_xml)
        archive.writestr("docProps/core.xml", core_xml)
        archive.writestr("word/_rels/document.xml.rels", doc_rels)
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", styles_xml)
        archive.writestr("word/numbering.xml", _numbering_xml())
        if task_type == "o1b_petition":
            archive.writestr("word/footer1.xml", _o1b_footer_xml())


def _write_rfe_company_docx(path: Path, config: dict[str, Any], case_dir: Path) -> None:
    """Build the RFE working file inside the company DOCX's native style system."""
    try:
        from docx import Document  # type: ignore
        from docx.enum.style import WD_STYLE_TYPE  # type: ignore
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.shared import Inches, Pt, RGBColor  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "python-docx is required to build the company-formatted RFE working memorandum."
        ) from exc

    from .rfe_strategy import effective_strategy_units, load_strategy_manifest

    response_config = config.get("rfe_response", {})
    if not isinstance(response_config, dict):
        response_config = {}
    human_template_value = str(
        response_config.get(
            "human_template_file", "templates/RFE/EB1/rfe draft template.docx"
        )
    )
    human_template = _resolve_project_or_case_template(case_dir, human_template_value)
    if not human_template.exists():
        raise SystemExit(f"RFE company Word template not found: {human_template}")
    yaml_value = str(
        response_config.get(
            "template_file",
            "templates/RFE/EB1/EB1A_RFE_response_unified_LLM_template.yaml",
        )
    )
    yaml_path = _resolve_project_or_case_template(case_dir, yaml_value)
    yaml_template = load_yaml_file(yaml_path) if yaml_path.exists() else {}
    manifest = load_strategy_manifest(case_dir)
    manifest["units"] = effective_strategy_units(case_dir, config)

    document = Document(str(human_template))
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for heading_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = document.styles[heading_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def ensure_style(name: str, *, italic: bool = False, gray: bool = False, indent: float = 0.0):
        if name in [style.name for style in document.styles]:
            style = document.styles[name]
        else:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.italic = italic
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        if gray:
            style.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        style.paragraph_format.left_indent = Inches(indent)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(8)
        return style

    ensure_style("RFE Quote", italic=True, indent=0.25)
    ensure_style("Drafting Note", italic=True, gray=True)
    ensure_style("Script Placeholder", gray=True)
    ensure_style("Evidence Index Item", indent=0.25)

    metadata = manifest.get("case_metadata", {})
    if not isinstance(metadata, dict):
        metadata = {}
    full_name = str(metadata.get("beneficiary_full_name") or _get(config, "beneficiary.full_name"))
    preferred = str(metadata.get("preferred_reference") or _get(config, "beneficiary.preferred_reference"))
    field = str(metadata.get("field") or config.get("field", ""))
    specialization = str(metadata.get("specialization") or config.get("specialization", ""))
    case_number = str(metadata.get("case_number") or _get(config, "rfe_metadata.case_number"))

    def add_text(text: str = "", *, style: str = "Normal", bold: bool = False, italic: bool = False, align=None):
        paragraph = document.add_paragraph(style=style)
        if align is not None:
            paragraph.alignment = align
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        return paragraph

    def add_label_value(label: str, value: str):
        paragraph = document.add_paragraph(style="Normal")
        label_run = paragraph.add_run(label)
        label_run.bold = True
        label_run.font.name = "Times New Roman"
        value_run = paragraph.add_run(value)
        value_run.font.name = "Times New Roman"
        return paragraph

    def add_criterion_line(role: str):
        roman = RFE_CRITERION_ROMAN_BY_ROLE.get(role, "")
        title = _rfe_criterion_heading(yaml_template, role)
        description = re.sub(r"^Criterion\s+\d+\.\s*", "", title).strip()
        paragraph = document.add_paragraph(style="Normal")
        prefix = paragraph.add_run(f"Criterion ({roman}).")
        prefix.bold = True
        paragraph.add_run(f" {description}")
        return paragraph

    response_date = str(metadata.get("rfe_response_date", "")).strip()
    add_text(response_date or "[RFE RESPONSE DATE TO BE CONFIRMED]")
    add_text("TO USCIS", bold=True)
    add_label_value(
        "RE: ",
        f"I-140 Petition for Alien of Extraordinary Ability in {field or '[FIELD]'} (EB-1A) - {specialization or '[SPECIALIZATION]'}",
    ).paragraph_format.space_before = Pt(10)
    add_label_value("Petitioner: ", full_name or "[BENEFICIARY NAME]")
    add_label_value("Case No.: ", case_number or "[CASE NUMBER]")
    office = str(metadata.get("uscis_office_or_service_center", "")).strip()
    if office:
        add_label_value("USCIS Office: ", office)
    rfe_date = str(metadata.get("rfe_date", "")).strip()
    if rfe_date:
        add_label_value("RFE Date: ", rfe_date)
    deadline = str(metadata.get("response_deadline", "")).strip()
    if deadline:
        add_label_value("Response Deadline: ", deadline)
    add_text(str(metadata.get("salutation", "")).strip() or "Dear Officer:")
    add_text(
        f"Please accept this response to the Request for Evidence regarding the Form I-140 petition filed on behalf of {full_name or '[BENEFICIARY NAME]'} under INA § 203(b)(1)(A), who is a specialist in the field of {field or '[FIELD]'}, and especially in {specialization or '[SPECIALIZATION]'}."
    )
    add_text(
        "We respectfully submit the enclosed additional evidence and explanations in response to the issues raised in the Request for Evidence. This response provides additional documentary evidence and legal explanation in support of the remaining criteria addressed in the RFE, namely:"
    )
    addressed_roles = list(
        dict.fromkeys(
            str(unit.get("criterion_role", ""))
            for unit in manifest.get("units", [])
            if unit.get("criterion_role")
        )
    )
    for role in addressed_roles:
        add_criterion_line(role)
    accepted = [
        str(role)
        for role in manifest.get("accepted_criteria", [])
        if str(role) in RFE_CRITERION_ROMAN_BY_ROLE
    ]
    if accepted:
        paragraph = add_text(
            f"USCIS recognized that {preferred or full_name or '[BENEFICIARY]'} satisfies the following criteria:",
            bold=True,
        )
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(12)
        for role in accepted:
            add_criterion_line(role)
    add_text(
        "For the convenience of the adjudicating officer, this response also cites and references evidence previously submitted with the original petition where relevant. Many supporting documents referenced throughout this response were already included in the initial filing and are therefore not duplicated herein in order to avoid making the present response unnecessarily voluminous and duplicative."
    )
    add_text(
        "Accordingly, we respectfully request that USCIS consider this response together with the evidence, exhibits, and legal arguments submitted in the original petition as part of the total evidentiary record."
    )
    closing = add_text("Respectfully submitted,")
    closing.paragraph_format.space_before = Pt(12)
    submitter = str(metadata.get("submitter_name", "")).strip() or full_name or "[SUBMITTER NAME]"
    title = str(metadata.get("submitter_title", "")).strip()
    add_text(submitter, bold=True)
    if title:
        add_text(title, bold=True)

    units = [unit for unit in manifest.get("units", []) if isinstance(unit, dict)]
    body_units = [
        unit
        for unit in units
        if str(unit.get("section_type", "")) not in {"cover_letter", "attachments"}
    ]
    grouped: list[tuple[str, list[dict[str, Any]]]] = []
    group_index: dict[str, int] = {}
    for unit in body_units:
        role = str(unit.get("criterion_role", ""))
        key = f"criterion:{role}" if role else f"section:{unit.get('section_id') or unit.get('unit_id')}"
        if key not in group_index:
            group_index[key] = len(grouped)
            grouped.append((key, []))
        grouped[group_index[key]][1].append(unit)

    for key, group in grouped:
        document.add_page_break()
        first = group[0]
        role = str(first.get("criterion_role", ""))
        section_type = str(first.get("section_type", ""))
        heading = _rfe_criterion_heading(yaml_template, role) if role else str(first.get("section_title") or first.get("title") or "RFE Response Section")
        add_text(heading, style="Heading 1")
        if role:
            add_text(
                "Exhibits from XX to XX, Pages from XX to XX.",
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )

        issues: list[dict[str, Any]] = []
        seen_issues: set[str] = set()
        for unit in group:
            for issue in unit.get("rfe_issues", []):
                issue_id = str(issue.get("issue_id", ""))
                if issue_id and issue_id not in seen_issues:
                    issues.append(issue)
                    seen_issues.add(issue_id)
        if issues:
            label = add_text("In the RFE, the officer states:", bold=True)
            label.paragraph_format.space_before = Pt(12)
            label.paragraph_format.space_after = Pt(6)
            for issue in issues:
                quote = str(issue.get("exact_rfe_quote", "")).strip()
                if quote:
                    cleaned_quote = quote.strip("“”\"")
                    add_text(f"“{cleaned_quote}”", style="RFE Quote")
            answer = add_text("Answer:", bold=True)
            answer.paragraph_format.line_spacing = 1.15

        for unit in group:
            unit_id = str(unit.get("unit_id", ""))
            draft_path = case_dir / "draft_sections/rfe/sections" / f"{unit_id}.md"
            if not draft_path.exists():
                continue
            title_text = str(unit.get("title", "")).strip()
            if len(group) > 1 or role or title_text != heading:
                add_text(title_text or str(unit.get("unit_id", "")), style="Heading 2")
            draft_text = draft_path.read_text(encoding="utf-8-sig", errors="replace")
            draft_paragraphs = _clean_memo_paragraphs(draft_text)
            if draft_paragraphs and _normalized_heading(draft_paragraphs[0]) in {
                _normalized_heading(title_text),
                _normalized_heading(heading),
            }:
                draft_paragraphs = draft_paragraphs[1:]
            for paragraph_text in draft_paragraphs:
                add_text(paragraph_text)

    document.add_page_break()
    add_text("Attachments / Evidence Index", style="Heading 1")
    evidence_entries = _evidence_index_entries(case_dir, config)
    for exhibit_heading, document_titles in evidence_entries:
        add_text(exhibit_heading, style="Heading 2")
        for document_title in document_titles:
            add_text(document_title, style="Evidence Index Item")

    for section in document.sections:
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _resolve_project_or_case_template(case_dir: Path, value: str) -> Path:
    candidate = Path(value)
    if candidate.is_absolute():
        return candidate
    project_candidate = PROJECT_ROOT / candidate
    if project_candidate.exists():
        return project_candidate
    return case_dir / candidate


def _rfe_criterion_heading(yaml_template: dict[str, Any], role: str) -> str:
    key = RFE_CRITERION_TEMPLATE_KEY_BY_ROLE.get(role, "")
    sections = yaml_template.get("sections", {}) if isinstance(yaml_template, dict) else {}
    criteria = sections.get("criteria", {}) if isinstance(sections, dict) else {}
    data = criteria.get(key, {}) if isinstance(criteria, dict) else {}
    if isinstance(data, dict) and data.get("heading"):
        return str(data["heading"])
    if role in RFE_CRITERION_HEADING_BY_ROLE:
        return RFE_CRITERION_HEADING_BY_ROLE[role]
    fallback = EB1A_CRITERION_STEP_BY_ROLE.get(role, ("", role.replace("_", " ").title()))[1]
    return fallback


def _read_template_text(path: Path) -> str:
    if not path.exists():
        raise SystemExit(f"Template not found: {path}")
    if path.suffix.lower() == ".docx":
        return extract_docx_text(path)
    return path.read_text(encoding="utf-8-sig", errors="replace")


def _parse_sections(text: str) -> list[TemplateSection]:
    sections: list[TemplateSection] = []
    lines = text.splitlines()
    previous_was_rule = False
    for index, raw_line in enumerate(lines, start=1):
        line = raw_line.strip()
        if not line:
            previous_was_rule = False
            continue
        if re.fullmatch(r"={8,}", line):
            previous_was_rule = True
            continue
        level = _section_level(line, previous_was_rule)
        previous_was_rule = False
        if level:
            sections.append(TemplateSection(level=level, title=line, line_number=index))
    return sections


def _section_level(line: str, previous_was_rule: bool) -> int:
    if previous_was_rule and re.match(r"^\d+(?:\.\d+)*\.\s+\S", line):
        return line.split(".", 1)[0].count(".") + 1
    if re.match(r"^\d+\.\s+[A-ZА-Я0-9]", line):
        return 1
    if re.match(r"^\d+\.\d+\.\s+\S", line):
        return 2
    if re.match(r"^[A-Z]\.\s+\S", line):
        return 1
    if re.match(r"^\[[A-Z]\]\s+\S", line):
        return 2
    if line.isupper() and 8 <= len(line) <= 90 and not line.startswith("__"):
        return 1
    return 0


def _eb1a_skeleton(
    config: dict[str, Any], case_dir: Path, report: TemplateParseReport
) -> list[dict[str, Any]]:
    structure = _load_working_structure(config)
    skeleton: list[dict[str, Any]] = [
        {
            "level": 1,
            "title": "First Page / Table of Contents",
            "paragraphs": ["[SCRIPT-CONTROLLED CONTENT: generated from EB1A_working_document_structure.yaml]"],
        },
        {
            "level": 1,
            "title": "Cover Letter",
            "paragraphs": [
                "[SCRIPT-CONTROLLED CONTENT: fixed legal text with beneficiary, area, specialization, gender, and claimed criteria substituted.]",
            ],
        },
        {
            "level": 1,
            "title": "Section 1. Initial Evidence of Extraordinary Ability",
            "paragraphs": ["[SCRIPT PLACEHOLDER: overview drafted near the end, after criteria are systematized.]"],
        },
        {
            "level": 2,
            "title": "[A] General information about the beneficiary and professional biography, overview of achievements, recommendations and education",
            "paragraphs": [
                "[LLM SECTION PLACEHOLDER: professional_biography]",
                "[LLM SECTION PLACEHOLDER: recommendation_letters_roster]",
                "[LLM SECTION PLACEHOLDER: recommendation_letters_quotes]",
            ],
        },
        {"level": 2, "title": "[B] Evidence of eligibility", "paragraphs": []},
    ]
    for item in _eb1a_criteria_from_config_and_folders(config, case_dir, structure):
        skeleton.append(item)
    skeleton.extend(
        [
            {
                "level": 1,
                "title": "Section 2. Beneficiary’s entry into the United States will substantially benefit prospectively the United States.",
                "paragraphs": ["[LLM SECTION PLACEHOLDER: specialization_essay]"],
            },
            {
                "level": 1,
                "title": "Conclusions / Final Merits",
                "paragraphs": ["[LLM SECTION PLACEHOLDER: final_overview]"],
            },
            {
                "level": 1,
                "title": "Statement of Beneficiary on Work Plans in the United States",
                "paragraphs": ["[LLM SECTION PLACEHOLDER: employment_plan]"],
            },
            {
                "level": 1,
                "title": "Exhibit List",
                "paragraphs": ["[SCRIPT PLACEHOLDER: generated from indexes/exhibit_index.csv]"],
            },
        ]
    )
    if not any(item.get("inferred") for item in skeleton):
        skeleton.extend(_template_fallback_sections(report, prefix="[TEMPLATE STRUCTURE PLACEHOLDER]"))
    return skeleton


def _o1b_skeleton(
    config: dict[str, Any], case_dir: Path, report: TemplateParseReport
) -> list[dict[str, Any]]:
    structure = _load_working_structure(config)
    skeleton: list[dict[str, Any]] = [
        {
            "level": 1,
            "title": "INDEX",
            "paragraphs": ["[SCRIPT-CONTROLLED CONTENT: generated from the document and exhibit indexes.]"],
        },
        {
            "level": 1,
            "title": "O-1B Cover Letter",
            "paragraphs": ["[SCRIPT-CONTROLLED CONTENT: petitioner, beneficiary, O-1B track, and claimed criteria.]"],
        },
        {
            "level": 1,
            "title": "OVERVIEW OF THE BENEFICIARY'S QUALIFICATIONS AND ACHIEVEMENTS",
            "paragraphs": [
                "[LLM SECTION PLACEHOLDER: o1b_final_overview]",
                "[LLM SECTION PLACEHOLDER: o1b_professional_biography]",
                "[OPTIONAL LLM SECTION PLACEHOLDER: o1b_recommendation_letter_episode]",
            ],
        },
        {
            "level": 1,
            "title": "OVERVIEW OF THE INDUSTRY",
            "paragraphs": ["[LLM SECTION PLACEHOLDER: o1b_industry_overview]"],
        },
        {"level": 1, "title": "EVIDENTIAL CRITERIA OF ELIGIBILITY", "paragraphs": []},
    ]
    skeleton.extend(_o1b_criteria_from_config_and_folders(config, case_dir, structure))
    skeleton.extend(
        [
            {
                "level": 1,
                "title": "ADVISORY OPINION",
                "paragraphs": ["[LLM SECTION PLACEHOLDER: o1b_advisory_opinion]"],
            },
            {
                "level": 1,
                "title": "Continuing to Work in the Area of Extraordinary Ability or Achievement",
                "paragraphs": ["[LLM SECTION PLACEHOLDER: o1b_continuing_to_work]"],
            },
            {
                "level": 1,
                "title": "Conclusions",
                "paragraphs": ["[LLM SECTION PLACEHOLDER: o1b_conclusions]"],
            },
            {
                "level": 1,
                "title": "Exhibit List",
                "paragraphs": ["[SCRIPT PLACEHOLDER: generated from indexes/exhibit_index.csv]"],
            },
        ]
    )
    if not any(item.get("inferred") for item in skeleton):
        skeleton.extend(_template_fallback_sections(report, prefix="[TEMPLATE STRUCTURE PLACEHOLDER]"))
    return skeleton


def _rfe_skeleton(
    config: dict[str, Any], case_dir: Path, report: TemplateParseReport
) -> list[dict[str, Any]]:
    skeleton: list[dict[str, Any]] = [
        {
            "level": 1,
            "title": "RFE Response",
            "paragraphs": [
                _substitute("Case No. __CASE_NO__", config),
                _substitute("Accepted for review on __RECEIPT_DATE__", config),
                _substitute("Beneficiary: __BENEFICIARY_FULL_NAME__", config),
                _substitute("Area: __FIELD__", config),
                _substitute("Specialization: __SPECIALIZATION__", config),
            ],
        },
    ]
    from .rfe_strategy import effective_strategy_units

    units = effective_strategy_units(case_dir, config)
    for unit in units:
        if str(unit.get("section_type", "")) in {"cover_letter", "attachments"}:
            continue
        unit_id = str(unit.get("unit_id", ""))
        draft_path = case_dir / "draft_sections" / "rfe" / "sections" / f"{unit_id}.md"
        if not draft_path.exists():
            continue
        draft_text = draft_path.read_text(encoding="utf-8-sig", errors="replace")
        paragraphs = _clean_memo_paragraphs(draft_text)
        unit_title = str(unit.get("title", unit_id))
        if paragraphs and _normalized_heading(paragraphs[0]) == _normalized_heading(unit_title):
            paragraphs = paragraphs[1:]
        skeleton.append(
            {
                "level": 1 if not unit.get("criterion_role") else 2,
                "title": unit_title,
                "paragraphs": paragraphs,
                "inferred": True,
                "unit_id": unit_id,
                "criterion_role": str(unit.get("criterion_role", "")),
            }
        )
    skeleton.append({"level": 1, "title": "Attachments / Evidence Index", "paragraphs": []})
    for exhibit_heading, document_titles in _evidence_index_entries(case_dir, config):
        skeleton.append(
            {
                "level": 2,
                "title": exhibit_heading,
                "paragraphs": [],
                "bullets": document_titles,
            }
        )
    if not any(item.get("inferred") for item in skeleton):
        skeleton.append(
            {
                "level": 1,
                "title": "Strategy bootstrap required",
                "paragraphs": [
                    "[SCRIPT PLACEHOLDER: accept the RFE strategy JSON to generate the case-specific structure]"
                ],
            }
        )
    return skeleton


def _eb1a_criteria_from_config_and_folders(
    config: dict[str, Any], case_dir: Path, structure: dict[str, Any]
) -> list[dict[str, Any]]:
    roles = folder_role_map(config)
    selected = _selected_criteria(config, case_dir)
    source_root = case_dir / _path_from_config(config, "source_originals")
    criteria_template = structure.get("criteria", {})
    if not isinstance(criteria_template, dict):
        criteria_template = {}
    tokens = _beneficiary_tokens(config)
    result: list[dict[str, Any]] = []
    for role, folder in roles.items():
        role = str(role)
        if role not in EB1A_CRITERION_STEP_BY_ROLE or role not in selected:
            continue
        folder_path = source_root / str(folder)
        criterion = criteria_template.get(role, {})
        if not isinstance(criterion, dict):
            criterion = {}
        section = str(criterion.get("section", ""))
        title = _format_case_text(str(criterion.get("title", EB1A_CRITERION_STEP_BY_ROLE[role][1])), tokens)
        step_ids = criterion.get("step_ids", [EB1A_CRITERION_STEP_BY_ROLE[role][0]])
        if not isinstance(step_ids, list):
            step_ids = [str(step_ids)]
        paragraphs = [f"[LLM SECTION PLACEHOLDER: {step_id}]" for step_id in step_ids]
        result.append(
            {
                "level": 3,
                "title": f"[{section}] {title}" if section else title,
                "paragraphs": paragraphs,
                "inferred": True,
                "criterion_role": role,
                "episode_folders": _episode_names(folder_path),
            }
        )
    return result


def _o1b_criteria_from_config_and_folders(
    config: dict[str, Any], case_dir: Path, structure: dict[str, Any]
) -> list[dict[str, Any]]:
    roles = folder_role_map(config)
    selected = _selected_criteria(config, case_dir)
    source_root = case_dir / _path_from_config(config, "source_originals")
    criteria_template = structure.get("criteria", {})
    if not isinstance(criteria_template, dict):
        criteria_template = {}
    tokens = _beneficiary_tokens(config)
    result: list[dict[str, Any]] = []
    for role in O1B_CRITERION_STEP_BY_ROLE:
        if role not in selected:
            continue
        criterion = criteria_template.get(role, {})
        if not isinstance(criterion, dict):
            criterion = {}
        title = _format_case_text(
            str(criterion.get("title", O1B_CRITERION_STEP_BY_ROLE[role][1])), tokens
        )
        step_ids = criterion.get("step_ids", [O1B_CRITERION_STEP_BY_ROLE[role][0]])
        if not isinstance(step_ids, list):
            step_ids = [str(step_ids)]
        folder_path = source_root / str(roles.get(role, role))
        result.append(
            {
                "level": 2,
                "title": title,
                "paragraphs": [f"[LLM SECTION PLACEHOLDER: {step_id}]" for step_id in step_ids],
                "inferred": True,
                "criterion_role": role,
                "episode_folders": _episode_names(folder_path),
            }
        )
    return result


def _rfe_issues_from_folders(case_dir: Path) -> list[str]:
    candidates: list[str] = []
    roots = [
        case_dir / "source_documents" / "rfe" / "issues",
        case_dir / "source_documents" / "rfe_response" / "new_documents" / "issues",
        case_dir / "source_documents" / "initial_filing" / "issues",
    ]
    seen: set[str] = set()
    for root in roots:
        if not root.exists():
            continue
        for child in sorted(root.iterdir()):
            if child.is_dir() and child.name not in seen and child.name != ".gitkeep":
                candidates.append(child.name)
                seen.add(child.name)
    return candidates


def _episode_names(folder: Path) -> list[str]:
    if not folder.exists():
        return []
    subfolders = [
        child.name
        for child in sorted(folder.iterdir())
        if child.is_dir() and _folder_has_documents(child)
    ]
    if subfolders:
        return subfolders
    files = [child for child in folder.iterdir() if child.is_file() and child.name != ".gitkeep"]
    return ["1"] if files else []


def _folder_has_documents(folder: Path) -> bool:
    return folder.exists() and any(
        child.is_file() and child.name != ".gitkeep" for child in folder.rglob("*")
    )


def _selected_criteria(config: dict[str, Any], case_dir: Path) -> list[str]:
    criterion_map = _criterion_step_map(config)
    claimed = [str(item) for item in config.get("claimed_criteria", []) or []]
    if claimed:
        return [role for role in criterion_map if role in claimed]
    roles = folder_role_map(config)
    source_root = case_dir / _path_from_config(config, "source_originals")
    return [
        role
        for role in criterion_map
        if role in roles and _folder_has_documents(source_root / str(roles[role]))
    ]


def _criterion_step_map(config: dict[str, Any]) -> dict[str, tuple[str, str]]:
    if str(config.get("task_type", "")) == "o1b_petition":
        return O1B_CRITERION_STEP_BY_ROLE
    return EB1A_CRITERION_STEP_BY_ROLE


def _load_working_structure(config: dict[str, Any]) -> dict[str, Any]:
    task_type = str(config.get("task_type", ""))
    value = config.get("working_document_template") or DEFAULT_WORKING_STRUCTURE_BY_TASK_TYPE.get(task_type, "")
    if not value:
        return {}
    path = Path(str(value))
    resolved = path if path.is_absolute() else PROJECT_ROOT / path
    if not resolved.exists():
        raise SystemExit(f"Working document structure template not found: {resolved}")
    return load_yaml_file(resolved)


def _template_fallback_sections(report: TemplateParseReport, *, prefix: str) -> list[dict[str, Any]]:
    return [
        {
            "level": min(max(section.level, 1), 3),
            "title": section.title,
            "paragraphs": [prefix],
            "inferred": False,
        }
        for section in report.sections[:80]
    ]


def _resolve_template_path(config: dict[str, Any], template_path: str) -> Path:
    task_type = str(config.get("task_type", ""))
    if template_path.strip():
        path = Path(template_path.strip())
        resolved = path if path.is_absolute() else PROJECT_ROOT / path
        if task_type != "eb1a_rfe_response" or resolved.exists():
            return resolved
        # Existing browser forms/case configs may still submit the retired TXT path.
        return PROJECT_ROOT / DEFAULT_TEMPLATE_BY_TASK_TYPE["eb1a_rfe_response"]
    configured = config.get("rfe_response", {}).get("template_file", "") if isinstance(config.get("rfe_response"), dict) else ""
    value = configured or DEFAULT_TEMPLATE_BY_TASK_TYPE.get(task_type, DEFAULT_TEMPLATE_BY_TASK_TYPE["eb1a_petition"])
    path = Path(str(value))
    resolved = path if path.is_absolute() else PROJECT_ROOT / path
    if task_type == "eb1a_rfe_response" and not resolved.exists():
        return PROJECT_ROOT / DEFAULT_TEMPLATE_BY_TASK_TYPE["eb1a_rfe_response"]
    return resolved


def _parse_case_info_file(path_value: str) -> dict[str, str]:
    if not path_value.strip():
        return {}
    path = Path(path_value.strip())
    if not path.exists():
        raise SystemExit(f"Case info file not found: {path}")
    if path.suffix.lower() in {".yaml", ".yml"}:
        data = load_yaml_file(path)
        return _flatten_mapping(data)
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    result: dict[str, str] = {}
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or ":" not in stripped:
            continue
        key, value = stripped.split(":", 1)
        result[key.strip()] = value.strip()
    return result


def _flatten_mapping(data: dict[str, Any], prefix: str = "") -> dict[str, str]:
    result: dict[str, str] = {}
    for key, value in data.items():
        full_key = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            result.update(_flatten_mapping(value, full_key))
        elif value is not None:
            result[full_key] = str(value)
    return result


def _normalize_intake_fields(fields: dict[str, str]) -> dict[str, str]:
    aliases = {
        "beneficiary_full_name": "beneficiary.full_name",
        "full_name": "beneficiary.full_name",
        "beneficiary_short_name": "beneficiary.preferred_reference",
        "preferred_reference": "beneficiary.preferred_reference",
        "gender": "beneficiary.gender",
        "honorific": "beneficiary.honorific",
        "citizenship": "beneficiary.citizenship",
        "area": "field",
        "field": "field",
        "specialization": "specialization",
        "soc": "soc_code",
        "soc_code": "soc_code",
        "petition_date": "petition_date",
        "case_number": "rfe_metadata.case_number",
        "receipt_date": "rfe_metadata.receipt_date",
        "rfe_date": "rfe_metadata.rfe_date",
        "response_deadline": "rfe_metadata.response_deadline",
        "uscis_address": "rfe_metadata.uscis_address",
        "o1b_track": "o1b_track",
        "petitioner_company_name": "petitioner.company_name",
        "petitioner_company_address": "petitioner.company_address",
        "petitioner_type": "petitioner.petitioner_type",
        "authorized_signatory": "petitioner.authorized_signatory",
        "filing_processing": "filing.processing",
        "validity_start": "filing.validity_start",
        "validity_end": "filing.validity_end",
        "filing_uscis_address": "filing.uscis_address",
        "position_or_role": "us_work.position_or_role",
        "compensation": "us_work.compensation",
        "work_location": "us_work.work_location",
        "duties_summary": "us_work.duties_summary",
    }
    result: dict[str, str] = {}
    for key, value in fields.items():
        normalized = aliases.get(key.strip(), key.strip())
        if value.strip():
            result[normalized] = value.strip()
    return result


def _apply_gender_defaults(config: dict[str, Any]) -> int:
    beneficiary = config.get("beneficiary")
    if not isinstance(beneficiary, dict):
        beneficiary = {}
        config["beneficiary"] = beneficiary
    gender = str(beneficiary.get("gender", "")).strip().lower()
    if gender not in {"male", "female", "neutral"}:
        preferred = str(beneficiary.get("preferred_reference", "")).strip().lower()
        if preferred.startswith("mr.") or preferred.startswith("mr "):
            gender = "male"
        elif preferred.startswith(("ms.", "ms ", "mrs.", "mrs ")):
            gender = "female"
        elif preferred.startswith("mx.") or preferred.startswith("mx "):
            gender = "neutral"
        else:
            return 0
    defaults = {
        "male": ("Mr.", {"subject": "he", "object": "him", "possessive": "his"}),
        "female": ("Ms.", {"subject": "she", "object": "her", "possessive": "her"}),
        "neutral": ("Mx.", {"subject": "they", "object": "them", "possessive": "their"}),
    }
    honorific, pronouns = defaults[gender]
    changed = 0
    for key, value in (("gender", gender), ("honorific", honorific)):
        if beneficiary.get(key) != value:
            beneficiary[key] = value
            changed += 1
    existing_pronouns = beneficiary.get("pronouns")
    if not isinstance(existing_pronouns, dict):
        existing_pronouns = {}
        beneficiary["pronouns"] = existing_pronouns
    for key, value in pronouns.items():
        if existing_pronouns.get(key) != value:
            existing_pronouns[key] = value
            changed += 1
    return changed


def _apply_updates(config: dict[str, Any], updates: dict[str, str]) -> int:
    count = 0
    for dotted, value in updates.items():
        target = config
        parts = dotted.split(".")
        for part in parts[:-1]:
            current = target.get(part)
            if not isinstance(current, dict):
                current = {}
                target[part] = current
            target = current
        if target.get(parts[-1]) != value:
            target[parts[-1]] = value
            count += 1
    return count


def _copy_source_folder(
    source: Path, destination: Path, *, refresh_existing: bool = False
) -> tuple[int, int]:
    if not source.exists() or not source.is_dir():
        raise SystemExit(f"Source folder not found: {source}")
    copied = 0
    skipped = 0
    destination.mkdir(parents=True, exist_ok=True)
    for item in sorted(source.rglob("*")):
        if not item.is_file():
            continue
        rel = item.relative_to(source)
        target = destination / rel
        if target.exists():
            if (
                not refresh_existing
                or item.resolve() == target.resolve()
                or filecmp.cmp(item, target, shallow=False)
            ):
                skipped += 1
                continue
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(item, target)
        copied += 1
    return copied, skipped


def _write_yaml_file(path: Path, data: dict[str, Any]) -> None:
    try:
        import yaml  # type: ignore

        text = yaml.safe_dump(data, allow_unicode=True, sort_keys=False)
    except ModuleNotFoundError:
        text = _dump_simple_yaml(data)
    path.write_text(text, encoding="utf-8")


def _dump_simple_yaml(data: dict[str, Any], indent: int = 0) -> str:
    lines: list[str] = []
    prefix = " " * indent
    for key, value in data.items():
        if isinstance(value, dict):
            lines.append(f"{prefix}{key}:")
            lines.append(_dump_simple_yaml(value, indent + 2).rstrip())
        elif isinstance(value, list):
            if not value:
                lines.append(f"{prefix}{key}: []")
            else:
                lines.append(f"{prefix}{key}:")
                for item in value:
                    lines.append(f"{prefix}  - {item}")
        else:
            lines.append(f"{prefix}{key}: {value}")
    return "\n".join(lines) + "\n"


def _path_from_config(config: dict[str, Any], key: str) -> Path:
    paths = config.get("paths", {})
    if not isinstance(paths, dict) or not paths.get(key):
        raise SystemExit(f"case_config.yaml is missing paths.{key}")
    return Path(str(paths[key]))


def _substitute(text: str, config: dict[str, Any]) -> str:
    replacements = {
        "__BENEFICIARY_FULL_NAME__": _get(config, "beneficiary.full_name"),
        "__BENEFICIARY_SHORT_NAME__": _get(config, "beneficiary.preferred_reference"),
        "__FIELD__": _get(config, "field"),
        "__AREA__": _get(config, "field"),
        "__SPECIALIZATION__": _get(config, "specialization"),
        "__SOC_CODE__": _get(config, "soc_code"),
        "__CASE_NO__": _get(config, "rfe_metadata.case_number"),
        "__RECEIPT_DATE__": _get(config, "rfe_metadata.receipt_date"),
    }
    for placeholder, value in replacements.items():
        text = text.replace(placeholder, value or placeholder)
    return text


def _get(config: dict[str, Any], dotted: str) -> str:
    value: Any = config
    for part in dotted.split("."):
        if not isinstance(value, dict):
            return ""
        value = value.get(part, "")
    return "" if value is None else str(value)


def _beneficiary_tokens(config: dict[str, Any]) -> dict[str, str]:
    def clean(value: object) -> str:
        text = "" if value is None else str(value).strip()
        return "" if text.startswith("__") else text

    full_name = clean(_get(config, "beneficiary.full_name")) or "[Beneficiary Full Name]"
    honorific = clean(_get(config, "beneficiary.honorific"))
    preferred = clean(_get(config, "beneficiary.preferred_reference"))
    if not honorific:
        lowered = preferred.lower()
        for candidate in ("Mr.", "Ms.", "Mrs.", "Mx."):
            if lowered.startswith(candidate.lower()):
                honorific = candidate
                break
    honorific = honorific or "Mx."
    if not preferred:
        surname = full_name.split()[-1] if full_name.split() else full_name
        preferred = f"{honorific} {surname}"
    formal_name = full_name if full_name.lower().startswith(honorific.lower()) else f"{honorific} {full_name}"
    possessive_reference = preferred + ("’" if preferred.endswith(("s", "S")) else "’s")
    gender = clean(_get(config, "beneficiary.gender")).lower()
    if gender not in {"male", "female", "neutral"}:
        gender = "male" if honorific == "Mr." else "female" if honorific in {"Ms.", "Mrs."} else "neutral"
    default_pronouns = {
        "male": {"subject": "he", "object": "him", "possessive": "his"},
        "female": {"subject": "she", "object": "her", "possessive": "her"},
        "neutral": {"subject": "they", "object": "them", "possessive": "their"},
    }[gender]
    return {
        "full_name": full_name,
        "formal_name": formal_name,
        "honorific": honorific,
        "preferred_reference": preferred,
        "possessive_reference": possessive_reference,
        "subject_pronoun": clean(_get(config, "beneficiary.pronouns.subject")) or default_pronouns["subject"],
        "object_pronoun": clean(_get(config, "beneficiary.pronouns.object")) or default_pronouns["object"],
        "possessive_pronoun": clean(_get(config, "beneficiary.pronouns.possessive")) or default_pronouns["possessive"],
        "field": clean(config.get("field", "")) or "[Area]",
        "specialization": clean(config.get("specialization", "")) or "[Specialization]",
        "soc_code": clean(config.get("soc_code", "")) or "[SOC Code]",
        "beneficiary_address": clean(_get(config, "beneficiary.address")) or "[Address]",
        "uscis_address": clean(_get(config, "filing.uscis_address")) or "[USCIS Address and processing center]",
        "petition_date": clean(config.get("petition_date", "")) or f"___/__/{datetime.now(UTC).year}",
        "citizenship": clean(_get(config, "beneficiary.citizenship")) or "[Citizenship]",
        "o1b_track": clean(config.get("o1b_track", "")) or "arts",
        "petitioner_company": clean(_get(config, "petitioner.company_name")) or "[Petitioner Company]",
        "petitioner_address": clean(_get(config, "petitioner.company_address")) or "[Petitioner Address]",
        "petitioner_type": clean(_get(config, "petitioner.petitioner_type")) or "us_employer",
        "authorized_signatory": clean(_get(config, "petitioner.authorized_signatory")) or "[Authorized Signatory]",
        "processing": clean(_get(config, "filing.processing")) or "Regular Processing",
        "validity_start": clean(_get(config, "filing.validity_start")) or "[Start Date]",
        "validity_end": clean(_get(config, "filing.validity_end")) or "[End Date]",
        "position_or_role": clean(_get(config, "us_work.position_or_role")) or "[Position / Role]",
        "compensation": clean(_get(config, "us_work.compensation")) or "[Compensation]",
        "work_location": clean(_get(config, "us_work.work_location")) or "[Work Location]",
        "duties_summary": clean(_get(config, "us_work.duties_summary")) or "[Duties Summary]",
    }


def _format_case_text(text: str, tokens: dict[str, str]) -> str:
    try:
        return text.format_map(tokens)
    except KeyError:
        return text


def _metadata_lines(config: dict[str, Any]) -> list[str]:
    return [
        f"- Task type: `{config.get('task_type', '')}`",
        f"- Beneficiary: `{_get(config, 'beneficiary.full_name')}`",
        f"- Area: `{config.get('field', '')}`",
        f"- Specialization: `{config.get('specialization', '')}`",
        f"- SOC code: `{config.get('soc_code', '')}`",
    ]


def _clean_memo_paragraphs(text: str) -> list[str]:
    return [
        paragraph
        for paragraph in (
            part.strip() for part in re.split(r"\n\s*\n", text.strip())
        )
        if paragraph
        and not any(paragraph.startswith(prefix) for prefix in TECHNICAL_MEMO_PREFIXES)
    ]


def _normalized_heading(value: str) -> str:
    return " ".join(
        "".join(character if character.isalnum() else " " for character in value.casefold()).split()
    )


def _evidence_index_entries(
    case_dir: Path, config: dict[str, Any]
) -> list[tuple[str, list[str]]]:
    paths = config.get("paths", {}) if isinstance(config.get("paths"), dict) else {}
    document_path = case_dir / str(paths.get("document_index", "indexes/document_index.csv"))
    exhibit_path = case_dir / str(paths.get("exhibit_index", "indexes/exhibit_index.csv"))
    if not document_path.exists() or not exhibit_path.exists():
        return []

    with document_path.open("r", encoding="utf-8-sig", newline="") as handle:
        documents = [dict(row) for row in csv.DictReader(handle)]
    with exhibit_path.open("r", encoding="utf-8-sig", newline="") as handle:
        exhibits = [dict(row) for row in csv.DictReader(handle)]
    by_id = {row.get("document_id", ""): row for row in documents if row.get("document_id")}

    entries: list[tuple[str, list[str]]] = []
    for exhibit in sorted(exhibits, key=_exhibit_index_sort_key):
        number = exhibit.get("exhibit_number", "").strip()
        if not number:
            continue
        title = exhibit.get("display_title", "").strip()
        heading = f"Exhibit {number}"
        if title and title.casefold() != heading.casefold():
            heading += f": {title}"
        document_ids = [
            item.strip() for item in exhibit.get("document_ids", "").split(";") if item.strip()
        ]
        selected_ids = set(document_ids)
        translations_by_parent: dict[str, list[dict[str, str]]] = {}
        for document_id in document_ids:
            document = by_id.get(document_id, {})
            parent = document.get("parent_document_id", "").strip()
            is_translation = (
                document.get("translation_status", "").strip().lower() == "translation"
                or document.get("relationship_type", "").strip().lower() == "translation"
            )
            if is_translation and parent in selected_ids:
                translations_by_parent.setdefault(parent, []).append(document)
        titles: list[str] = []
        for document_id in document_ids:
            document = by_id.get(document_id)
            if not document:
                continue
            parent = document.get("parent_document_id", "").strip()
            is_translation = (
                document.get("translation_status", "").strip().lower() == "translation"
                or document.get("relationship_type", "").strip().lower() == "translation"
            )
            if is_translation and parent in selected_ids:
                continue
            document_title = (
                document.get("display_title", "").strip()
                or document.get("original_file_name", "").strip()
            )
            if translations_by_parent.get(document_id):
                document_title += "; English translation"
            if document_title:
                titles.append(document_title)
        entries.append((heading, titles))
    return entries


def _exhibit_index_sort_key(row: dict[str, str]) -> tuple[int, list[object], str]:
    order = row.get("final_bundle_order", "").strip()
    if order.isdigit():
        return (0, [int(order)], row.get("exhibit_number", ""))
    return (1, _natural_key(row.get("exhibit_number", "")), row.get("exhibit_id", ""))


def _natural_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def _o1b_document_xml(config: dict[str, Any], case_dir: Path) -> str:
    structure = _load_working_structure(config)
    tokens = _beneficiary_tokens(config)
    criteria_data = structure.get("criteria", {})
    if not isinstance(criteria_data, dict):
        criteria_data = {}
    selected_roles = _selected_criteria(config, case_dir)
    selected_criteria = [
        (role, criteria_data[role])
        for role in selected_roles
        if role in criteria_data and isinstance(criteria_data[role], dict)
    ]
    track = tokens["o1b_track"].lower()
    is_mptv = track == "mptv"
    classification = (
        "O-1B extraordinary achievement in motion picture or television"
        if is_mptv
        else "O-1B extraordinary ability in the arts"
    )
    regulation = "8 C.F.R. § 214.2(o)(3)(v)" if is_mptv else "8 C.F.R. § 214.2(o)(3)(iv)"
    standard = "extraordinary achievement" if is_mptv else "distinction"
    body: list[str] = []

    body.append(_rich_paragraph([("INDEX:", {"bold": True})], style="Heading1"))
    index_rows = structure.get("index_rows", [])
    if isinstance(index_rows, list):
        for row in index_rows:
            body.append(_rich_paragraph([(_format_case_text(str(row), tokens), {})]))
    body.append(_placeholder_paragraph("[SCRIPT PLACEHOLDER: final exhibit titles and page ranges are generated from indexes after PDF assembly.]"))
    body.append(_page_break_paragraph())

    body.append(_rich_paragraph([(tokens["petition_date"], {})]))
    body.append(_rich_paragraph([("TO USCIS", {"bold": True})], style="Heading1"))
    body.append(_rich_paragraph([(tokens["processing"], {})]))
    body.append(_rich_paragraph([(f"RE: I-129 Petition for {classification}", {"bold": True})]))
    body.append(_rich_paragraph([(f'Petitioner: {tokens["petitioner_company"]}', {"bold": True})]))
    body.append(_rich_paragraph([(f'Beneficiary: {tokens["full_name"]}', {"bold": True})]))
    body.append(_rich_paragraph([(f'Citizenship: {tokens["citizenship"]}', {"bold": True})]))
    body.append(_rich_paragraph([("Dear Officer,", {})]))
    body.append(
        _rich_paragraph(
            [
                (
                    f'{tokens["petitioner_company"]} petitions for {tokens["full_name"]} to be classified '
                    f'as a nonimmigrant of {classification}, specifically in {tokens["specialization"]}, '
                    f'pursuant to {regulation}.',
                    {},
                )
            ]
        )
    )
    count = len([role for role, _ in selected_criteria if role != "comparable_evidence"])
    body.append(
        _rich_paragraph(
            [
                (
                    f'As demonstrated below, {tokens["preferred_reference"]} satisfies {count} of the six '
                    f'regulatory criteria claimed for this petition:',
                    {},
                )
            ]
        )
    )
    for _role, criterion in selected_criteria:
        cover_text = _format_case_text(str(criterion.get("cover_text", criterion.get("title", ""))), tokens)
        body.append(_rich_paragraph([(cover_text, {})], num_id=2))
    body.append(_page_break_paragraph())

    body.append(_rich_paragraph([("OVERVIEW OF THE BENEFICIARY'S QUALIFICATIONS AND ACHIEVEMENTS", {"bold": True})], style="Heading1"))
    body.append(_rich_paragraph([("Exhibits from XX to XX, Pages from XX to XX.", {"italic": True})]))
    body.append(
        _draft_step_xml(case_dir, "o1b_final_overview")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: o1b_final_overview — draft after the criteria.]"))
    body.append(
        _draft_step_xml(case_dir, "o1b_professional_biography", strip_opening_headings=True)
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: education, specialization, and professional biography.]"))
    recommendation_xml = _draft_step_xml(
        case_dir, "o1b_recommendation_letter_episode", episode_style="Heading2", episode_label=""
    )
    if recommendation_xml:
        body.append(_rich_paragraph([("Recommendation Letters", {"bold": True})], style="Heading2"))
        body.append(recommendation_xml)

    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("OVERVIEW OF THE INDUSTRY", {"bold": True})], style="Heading1"))
    body.append(_rich_paragraph([("Exhibits from XX to XX, Pages from XX to XX.", {"italic": True})]))
    body.append(
        _draft_step_xml(case_dir, "o1b_industry_overview")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: concise O-1B industry overview.]"))

    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("EVIDENTIAL CRITERIA OF ELIGIBILITY", {"bold": True})], style="Heading1"))
    body.append(
        _rich_paragraph(
            [
                (
                    f'The evidence below demonstrates that {tokens["preferred_reference"]} satisfies the claimed '
                    f'O-1B criteria and meets the governing {standard} standard. Each criterion is supported by '
                    "identified documentary evidence.",
                    {},
                )
            ]
        )
    )
    for role, criterion in selected_criteria:
        body.append(_page_break_paragraph())
        title = _format_case_text(str(criterion.get("title", "")), tokens)
        body.append(_rich_paragraph([(title, {"bold": True})], style="Heading1"))
        exhibit_range = str(criterion.get("exhibit_range", "Exhibits from XX to XX, Pages from XX to XX."))
        body.append(_rich_paragraph([(exhibit_range, {"italic": True})]))
        step_ids = criterion.get("step_ids", [])
        if not isinstance(step_ids, list):
            step_ids = [step_ids]
        for step_id in step_ids:
            body.append(
                _draft_step_xml(
                    case_dir,
                    str(step_id),
                    episode_style="Heading2",
                    episode_label="",
                )
                or _placeholder_paragraph(f"[LLM SECTION PLACEHOLDER: {step_id}]")
            )

    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("ADVISORY OPINION", {"bold": True})], style="Heading1"))
    body.append(_rich_paragraph([("Exhibit 9.1, pages from XX to XX.", {"italic": True})]))
    body.append(
        _draft_step_xml(case_dir, "o1b_advisory_opinion")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: O-1 consultation under 8 C.F.R. § 214.2(o)(5).]"))

    body.append(_page_break_paragraph())
    body.append(
        _rich_paragraph(
            [("Continuing to Work in the Area of Extraordinary Ability or Achievement", {"bold": True})],
            style="Heading1",
        )
    )
    body.append(_rich_paragraph([("Exhibits from XX to XX, Pages from XX to XX.", {"italic": True})]))
    body.append(
        _draft_step_xml(case_dir, "o1b_continuing_to_work")
        or _placeholder_paragraph(
            f'[LLM SECTION PLACEHOLDER: {tokens["position_or_role"]}, {tokens["work_location"]}, '
            f'{tokens["validity_start"]}–{tokens["validity_end"]}, {tokens["compensation"]}.]'
        )
    )

    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("Conclusions", {"bold": True})], style="Heading1"))
    body.append(
        _draft_step_xml(case_dir, "o1b_conclusions")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: conclusion listing only the claimed and supported criteria.]"))
    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("Exhibit List", {"bold": True})], style="Heading1"))
    body.append(_placeholder_paragraph("[SCRIPT PLACEHOLDER: generated from indexes/exhibit_index.csv]"))
    body.append(_o1b_section_properties())
    return _wrap_document_xml("".join(body))


def _eb1a_document_xml(config: dict[str, Any], case_dir: Path) -> str:
    structure = _load_working_structure(config)
    tokens = _beneficiary_tokens(config)
    document_data = structure.get("document", {})
    if not isinstance(document_data, dict):
        document_data = {}
    criteria_data = structure.get("criteria", {})
    if not isinstance(criteria_data, dict):
        criteria_data = {}
    selected_roles = _selected_criteria(config, case_dir)
    selected_criteria = [
        (role, criteria_data[role])
        for role in selected_roles
        if role in criteria_data and isinstance(criteria_data[role], dict)
    ]

    body: list[str] = []
    petition_title = str(
        document_data.get(
            "petition_title", "Petition for Permanent Residence under Extraordinary Ability (EB-1A)"
        )
    )

    # Page 1: deterministic title page and filing table of contents.
    body.append(_rich_paragraph([(tokens["full_name"], {"bold": True})], align="center"))
    body.append(_rich_paragraph([(tokens["beneficiary_address"], {})], align="center"))
    body.append(_rich_paragraph([(petition_title, {"bold": True})], align="center"))
    body.append(_rich_paragraph([("Table of Contents:", {"italic": True})]))
    toc_rows = structure.get("table_of_contents", [])
    formatted_toc: list[list[str]] = []
    if isinstance(toc_rows, list):
        for row in toc_rows:
            if isinstance(row, list) and len(row) >= 2:
                formatted_toc.append(
                    [_format_case_text(str(row[0]), tokens), _format_case_text(str(row[1]), tokens)]
                )
    body.append(_table_xml(formatted_toc, [8640, 720]))
    body.append(_page_break_paragraph())

    # Cover letter: fixed text; only case facts and selected criteria vary.
    cover_heading = _format_case_text(
        str(document_data.get("cover_heading", "Permanent residence petition for {preferred_reference}")),
        tokens,
    )
    body.append(_rich_paragraph([(cover_heading, {"bold": True})]))
    body.append(_rich_paragraph([(tokens["petition_date"], {})]))
    body.append(
        _table_xml(
            [[f'To:\n{tokens["uscis_address"]}', "", f'From:\n{tokens["formal_name"]}\n{tokens["beneficiary_address"]}']],
            [3600, 1440, 4320],
        )
    )
    body.append(_rich_paragraph([(_format_case_text(str(document_data.get("re_line", "")), tokens), {})]))
    body.append(_rich_paragraph([(f'Petitioner/Beneficiary: {tokens["formal_name"]}', {"bold": True})]))
    body.append(_rich_paragraph([(str(document_data.get("classification", "")), {"bold": True})]))
    body.append(_rich_paragraph([(str(document_data.get("petition_type", "")), {"bold": True})]))
    body.append(_rich_paragraph([(str(document_data.get("salutation", "Dear Immigration Officer")), {"bold": True})]))
    body.append(
        _rich_paragraph(
            [
                (
                    f'The following evidence is submitted in support of {tokens["formal_name"]}\'s application '
                    f'to qualify as an alien of extraordinary ability. The evidence presented demonstrates that '
                    f'{tokens["formal_name"]} meets the criteria of section 203(b)(1)(A) of the Immigration and '
                    'Nationality Act [8 U.S.C. 1153] as a person of extraordinary ability for the following reasons:',
                    {},
                )
            ]
        )
    )
    body.append(
        _rich_paragraph(
            [
                (f'{tokens["preferred_reference"]} has extraordinary abilities in ', {}),
                (tokens["field"], {"bold": True}),
                (", especially in ", {}),
                (tokens["specialization"], {"bold": True}),
                (" as proven by national and international accomplishments in this field, which have been extensively documented. (See Section 1)", {}),
            ],
            num_id=1,
        )
    )
    body.append(
        _rich_paragraph(
            [
                (f'{tokens["preferred_reference"]} has the experience indicating that ', {}),
                (tokens["subject_pronoun"], {}),
                (" is an outstanding professional in ", {}),
                (tokens["specialization"], {"bold": True}),
                (", whose achievements and contributions to the field are exceptional. (See Section 1 [A], [B]).", {}),
            ],
            num_id=1,
        )
    )
    roman_values = [str(item.get("roman", "")) for _, item in selected_criteria]
    count_words = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten"]
    count_text = count_words[len(selected_criteria)] if len(selected_criteria) < len(count_words) else str(len(selected_criteria))
    body.append(
        _rich_paragraph(
            [
                (f'{tokens["preferred_reference"]} has provided evidence that ', {}),
                (tokens["subject_pronoun"], {}),
                (" meets ", {}),
                (count_text, {"bold": True}),
                (f' ({", ".join(roman_values)}) of the ten criteria listed in 8 CFR section 204.5(h)(3), namely:', {}),
            ],
            num_id=1,
        )
    )
    for _, criterion in selected_criteria:
        roman = str(criterion.get("roman", ""))
        ordinal = _roman_ordinal(roman)
        cover_text = _format_case_text(str(criterion.get("cover_text", criterion.get("title", ""))), tokens)
        section = str(criterion.get("section", ""))
        body.append(
            _rich_paragraph(
                [(f"{cover_text} (Please kindly refer to Section 1 [{section}]).", {})],
                num_id=100 + ordinal,
            )
        )
    body.append(
        _rich_paragraph(
            [(f'{tokens["possessive_reference"]} work will be of significant benefit to the United States (Please kindly refer to Section 2).', {})],
            num_id=1,
        )
    )
    body.append(
        _rich_paragraph(
            [(f'In the U.S., {tokens["preferred_reference"]} plans to continue {tokens["possessive_pronoun"]} work in this area (Please kindly refer to the Beneficiary Statement on Plans to Work in the U.S.).', {})],
            num_id=1,
        )
    )
    body.append(
        _rich_paragraph(
            [(f'Citing 8 CFR section 204.5(h)(1), {tokens["preferred_reference"]} may apply for an I–140 visa to be classified under section 203(b)(1)(A) of the Act as an alien of extraordinary ability on {tokens["possessive_pronoun"]} behalf.', {})]
        )
    )
    body.append(
        _rich_paragraph(
            [("Under 8 CFR section 204.5(h)(5), neither an offer of employment in the United States nor a labor certification is required for this classification.", {})]
        )
    )
    body.append(_rich_paragraph([(tokens["formal_name"], {})]))
    body.append(_rich_paragraph([(tokens["petition_date"], {})]))
    body.append(_page_break_paragraph())

    # Memorandum body: exact stable headings; LLM fills prose in later workflow steps.
    body.append(_rich_paragraph([("Section 1. Initial Evidence of Extraordinary ability", {"bold": True})], style="Heading1"))
    body.append(
        _rich_paragraph(
            [(f'[A] General information about the beneficiary and {tokens["possessive_pronoun"]} professional biography, overview of achievements, recommendations and education', {"bold": True})],
            style="Heading2",
        )
    )
    body.append(_rich_paragraph([("Overview", {"bold": True})], style="Heading3"))
    body.append(_placeholder_paragraph("[SCRIPT PLACEHOLDER: overview drafted near the end, after criteria are systematized.]"))
    professional_draft = _draft_step_xml(case_dir, "professional_biography", strip_opening_headings=True)
    if professional_draft:
        body.append(professional_draft)
    else:
        body.append(_rich_paragraph([(f'{tokens["possessive_reference"]} LinkedIn page: [placeholder]', {"bold": True})]))
        body.append(_rich_paragraph([(f'Specialization: {tokens["specialization"]}', {"bold": True})]))
        body.append(_rich_paragraph([(f'Relevant SOC Code: {tokens["soc_code"]}', {"italic": True})]))
        body.append(_placeholder_paragraph("[LLM SECTION PLACEHOLDER: short specialization explanation]"))
        body.append(_rich_paragraph([(f'Professional Biography of {tokens["formal_name"]}', {"bold": True})], style="Heading3"))
        body.append(_placeholder_paragraph("[LLM SECTION PLACEHOLDER: education and professional_biography]"))
    body.append(_rich_paragraph([("List of recommendations and the general quotes from the letters", {"bold": True})], style="Heading3"))
    body.append(
        _draft_step_xml(case_dir, "recommendation_letters_roster")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: recommendation_letters_roster]")
    )
    body.append(
        _draft_step_xml(case_dir, "recommendation_letters_quotes")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: recommendation_letters_quotes]")
    )
    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("[B] Evidence of eligibility", {"bold": True})], style="Heading2"))

    for index, (_, criterion) in enumerate(selected_criteria, start=1):
        if index > 1:
            body.append(_page_break_paragraph())
        section = str(criterion.get("section", ""))
        title = _format_case_text(str(criterion.get("title", "")), tokens)
        body.append(_rich_paragraph([(f"[{section}] {title}", {"bold": True})], style="Heading3"))
        body.append(
            _rich_paragraph(
                [(f"(Please, kindly refer to Exhibit {index}, page XX: Criterion {criterion.get('roman', '')}. {title})", {"italic": True})]
            )
        )
        body.append(
            _rich_paragraph(
                [("Within the criterion, the following documents are attached:", {"bold": True, "italic": True, "underline": True})]
            )
        )
        step_ids = criterion.get("step_ids", [])
        if not isinstance(step_ids, list):
            step_ids = [step_ids]
        for step_id in step_ids:
            body.append(
                _draft_step_xml(case_dir, str(step_id))
                or _placeholder_paragraph(f"[LLM SECTION PLACEHOLDER: {step_id}]")
            )

    body.append(_page_break_paragraph())
    body.append(
        _rich_paragraph(
            [("Section 2. Beneficiary’s entry into the United States will substantially benefit prospectively the United States.", {})],
            style="SpecialSection",
            align="center",
        )
    )
    body.append(_rich_paragraph([(f'2.1. {tokens["specialization"]} is an area of intrinsic merit in USA', {"bold": True})], style="Heading2"))
    body.append(_rich_paragraph([("Overview of the Specialization", {"bold": True})], style="Heading3"))
    body.append(_rich_paragraph([("Economic Impact and Contributions in the U.S.", {"bold": True})], style="Heading3"))
    body.append(
        _draft_step_xml(case_dir, "specialization_essay")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: specialization_essay]")
    )
    body.append(_rich_paragraph([("Conclusion", {"bold": True})], style="Heading3"))
    body.append(_rich_paragraph([(f'2.2. {tokens["possessive_reference"]} work will ultimately benefit the United States', {"bold": True})], style="Heading2"))
    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("Conclusions / Final Merits", {"bold": True})], style="Heading1"))
    body.append(
        _draft_step_xml(case_dir, "final_overview")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: final_overview]")
    )
    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("Statement of Beneficiary on Work Plans in the United States", {})], style="Heading1"))
    body.append(_rich_paragraph([(f'Date: {tokens["petition_date"]}', {"italic": True})]))
    body.append(
        _rich_paragraph(
            [(f'Dear immigration officer, my name is {tokens["formal_name"]}. I am the beneficiary of the petition for permanent residence based on Extraordinary Ability (EB-1A). I have vast experience in the {tokens["specialization"]} field and I intend to continue to build my career in this field in the United States.', {})]
        )
    )
    body.append(
        _draft_step_xml(case_dir, "employment_plan")
        or _placeholder_paragraph("[LLM SECTION PLACEHOLDER: employment_plan]")
    )
    body.append(_page_break_paragraph())
    body.append(_rich_paragraph([("Exhibit List", {"bold": True})], style="Heading1"))
    body.append(_placeholder_paragraph("[SCRIPT PLACEHOLDER: generated from indexes/exhibit_index.csv]"))
    body.append(_section_properties())
    return _wrap_document_xml("".join(body))


def _roman_ordinal(value: str) -> int:
    mapping = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6, "vii": 7, "viii": 8, "ix": 9, "x": 10}
    return mapping.get(value.lower(), 1)


def _placeholder_paragraph(text: str) -> str:
    return _rich_paragraph([(text, {})], style="Placeholder")


def _draft_step_xml(
    case_dir: Path,
    step_id: str,
    *,
    strip_opening_headings: bool = False,
    episode_style: str = "Heading3",
    episode_label: str = "Episode:",
) -> str:
    paths = _draft_paths_for_step(case_dir, step_id)
    if not paths:
        return ""
    parts: list[str] = []
    for index, path in enumerate(paths):
        text = path.read_text(encoding="utf-8-sig", errors="replace").strip()
        if not text:
            continue
        if len(paths) > 1:
            heading = f"{episode_label} {path.stem}".strip()
            parts.append(_rich_paragraph([(heading, {"bold": True})], style=episode_style))
        parts.append(_draft_text_xml(text, strip_opening_headings=strip_opening_headings and index == 0))
    return "".join(parts)


def _draft_paths_for_step(case_dir: Path, step_id: str) -> list[Path]:
    root = case_dir / "draft_sections"
    fixed = {
        "professional_biography": root / "professional_biography.md",
        "recommendation_letters_roster": root / "recommendation_letters_roster.md",
        "recommendation_letters_quotes": root / "recommendation_letters_quotes.md",
        "specialization_essay": root / "specialization_essay.md",
        "final_overview": root / "final_overview.md",
    }
    if step_id in fixed:
        return [fixed[step_id]] if fixed[step_id].exists() else []
    patterns = {
        "criterion_awards_episode": ("criteria/awards", "*.md"),
        "criterion_memberships_episode": ("criteria/memberships", "*.md"),
        "criterion_media_episode": ("criteria/media", "*.md"),
        "criterion_judging_episode": ("criteria/judging", "*.md"),
        "criterion_original_contribution_fact": ("criteria/original_contribution", "*_phase_1.md"),
        "criterion_original_contribution_significance": ("criteria/original_contribution", "*_phase_2.md"),
        "criterion_scholarly_articles_episode": ("criteria/scholarly_articles", "*.md"),
        "criterion_exhibitions_episode": ("criteria/exhibitions", "*.md"),
        "criterion_leading_critical_role_fact": ("criteria/leading_critical_role", "*_phase_1.md"),
        "criterion_leading_critical_role_reputation": ("criteria/leading_critical_role", "*_phase_2.md"),
        "criterion_high_salary_fact": ("criteria/high_salary", "*_phase_1.md"),
        "criterion_high_salary_comparison": ("criteria/high_salary", "*_phase_2.md"),
        "criterion_commercial_success_episode": ("criteria/commercial_success", "*.md"),
        "employment_plan": ("employment_plan", "*.md"),
        "o1b_petitioner_support_letter": ("supporting_documents", "petitioner_support_letter.md"),
        "o1b_itinerary": ("supporting_documents", "itinerary.md"),
        "o1b_professional_biography": ("o1b", "professional_biography.md"),
        "o1b_recommendation_letter_episode": ("o1b/recommendation_letters", "*.md"),
        "o1b_industry_overview": ("o1b", "industry_overview.md"),
        "o1b_criterion_i_episode": ("o1b/criteria/criterion_i", "*.md"),
        "o1b_criterion_ii_episode": ("o1b/criteria/criterion_ii", "*.md"),
        "o1b_criterion_iii_role": ("o1b/criteria/criterion_iii", "*_phase_1.md"),
        "o1b_criterion_iii_reputation": ("o1b/criteria/criterion_iii", "*_phase_2.md"),
        "o1b_criterion_iv_episode": ("o1b/criteria/criterion_iv", "*.md"),
        "o1b_criterion_v_episode": ("o1b/criteria/criterion_v", "*.md"),
        "o1b_criterion_vi_compensation": ("o1b/criteria/criterion_vi", "*_phase_1.md"),
        "o1b_criterion_vi_comparison": ("o1b/criteria/criterion_vi", "*_phase_2.md"),
        "o1b_comparable_evidence_episode": ("o1b/criteria/comparable_evidence", "*.md"),
        "o1b_advisory_opinion": ("o1b", "advisory_opinion.md"),
        "o1b_continuing_to_work": ("o1b", "continuing_to_work.md"),
        "o1b_final_overview": ("o1b", "final_overview.md"),
        "o1b_conclusions": ("o1b", "conclusions.md"),
    }
    folder_pattern = patterns.get(step_id)
    if not folder_pattern:
        return []
    folder, pattern = folder_pattern
    target = root / folder
    return sorted(path for path in target.glob(pattern) if path.is_file()) if target.exists() else []


def _draft_text_xml(text: str, *, strip_opening_headings: bool = False) -> str:
    blocks = [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]
    if strip_opening_headings:
        while blocks and (
            blocks[0].lower().startswith("section 1.")
            or blocks[0].lstrip().startswith("[A]")
        ):
            blocks.pop(0)
    output: list[str] = []
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        if all(line.startswith(("- ", "* ")) for line in lines):
            for line in lines:
                output.append(_rich_paragraph(_markdown_runs(line[2:].strip()), num_id=2))
            continue
        heading_match = re.match(r"^(#{1,4})\s+(.+)$", block)
        if heading_match:
            level = min(len(heading_match.group(1)), 3)
            output.append(_rich_paragraph(_markdown_runs(heading_match.group(2)), style=f"Heading{level}"))
            continue
        compact = " ".join(lines)
        if len(compact) <= 100 and len(compact.split()) <= 12 and compact[-1:] not in ".?!;":
            output.append(_rich_paragraph(_markdown_runs(compact), style="Heading3"))
        else:
            output.append(_rich_paragraph(_markdown_runs(compact)))
    return "".join(output)


def _markdown_runs(text: str) -> list[tuple[str, dict[str, bool]]]:
    runs: list[tuple[str, dict[str, bool]]] = []
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            runs.append((part[2:-2], {"bold": True}))
        else:
            runs.append((part, {}))
    return runs


def _page_break_paragraph() -> str:
    return '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'


def _rich_paragraph(
    runs: list[tuple[str, dict[str, bool]]],
    *,
    style: str = "Normal",
    align: str = "",
    num_id: int | None = None,
) -> str:
    ppr_parts: list[str] = []
    if style and style != "Normal":
        ppr_parts.append(f'<w:pStyle w:val="{html.escape(style)}"/>')
    if align:
        ppr_parts.append(f'<w:jc w:val="{html.escape(align)}"/>')
    if num_id is not None:
        ppr_parts.append(f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>')
    ppr = f'<w:pPr>{"".join(ppr_parts)}</w:pPr>' if ppr_parts else ""
    run_xml: list[str] = []
    for text, props in runs:
        rpr: list[str] = []
        if props.get("bold"):
            rpr.append("<w:b/>")
        if props.get("italic"):
            rpr.append("<w:i/>")
        if props.get("underline"):
            rpr.append('<w:u w:val="single"/>')
        rpr_xml = f'<w:rPr>{"".join(rpr)}</w:rPr>' if rpr else ""
        fragments = text.split("\n")
        content = "<w:br/>".join(
            f'<w:t xml:space="preserve">{html.escape(fragment)}</w:t>' for fragment in fragments
        )
        run_xml.append(f"<w:r>{rpr_xml}{content}</w:r>")
    return f'<w:p>{ppr}{"".join(run_xml)}</w:p>'


def _table_xml(rows: list[list[str]], widths: list[int]) -> str:
    if not rows:
        return ""
    total = sum(widths)
    grid = "".join(f'<w:gridCol w:w="{width}"/>' for width in widths)
    row_xml: list[str] = []
    for row in rows:
        cells: list[str] = []
        for index, width in enumerate(widths):
            value = row[index] if index < len(row) else ""
            cells.append(
                f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/></w:tcPr>'
                f'{_rich_paragraph([(str(value), {})])}</w:tc>'
            )
        row_xml.append(f'<w:tr>{"".join(cells)}</w:tr>')
    return (
        '<w:tbl><w:tblPr>'
        f'<w:tblW w:w="{total}" w:type="dxa"/><w:tblInd w:w="0" w:type="dxa"/>'
        '<w:tblLayout w:type="fixed"/>'
        '<w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="0" w:type="dxa"/>'
        '<w:bottom w:w="40" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>'
        '<w:tblBorders><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>'
        f'</w:tblPr><w:tblGrid>{grid}</w:tblGrid>{"".join(row_xml)}</w:tbl>'
    )


def _wrap_document_xml(body: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:w10="urn:schemas-microsoft-com:office:word" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
        'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
        'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" '
        'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'mc:Ignorable="w14 wp14"><w:body>'
        + body
        + "</w:body></w:document>"
    )


def _document_xml(case_id: str, config: dict[str, Any], skeleton: list[dict[str, Any]]) -> str:
    body: list[str] = []
    body.append(_paragraph(f"Working memorandum: {case_id}", style="Title"))
    for line in _metadata_lines(config):
        body.append(_paragraph(line.replace("`", ""), style="Meta"))
    body.append(_paragraph("Generated working file. Replace bracketed placeholders as LLM outputs are validated.", style="Note"))
    for item in skeleton:
        level = min(max(int(item.get("level", 1)), 1), 3)
        body.append(_paragraph(str(item.get("title", "Untitled section")), style=f"Heading{level}"))
        for paragraph in item.get("paragraphs", []):
            body.append(_paragraph(str(paragraph), style="Placeholder" if "[LLM" in str(paragraph) or "[SCRIPT" in str(paragraph) else "Normal"))
        for bullet in item.get("bullets", []):
            body.append(_paragraph(str(bullet), style="Bullet"))
    body.append(_section_properties())
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:w10="urn:schemas-microsoft-com:office:word" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
        'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
        'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" '
        'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'mc:Ignorable="w14 wp14"><w:body>'
        + "".join(body)
        + "</w:body></w:document>"
    )


def _paragraph(text: str, *, style: str) -> str:
    escaped = html.escape(text)
    run_props = ""
    if style in {"Title", "Heading1", "Heading2", "Heading3"}:
        run_props = "<w:rPr><w:b/></w:rPr>"
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style != "Normal" else ""
    if style == "Bullet":
        escaped = "• " + escaped
    return f"<w:p>{ppr}<w:r>{run_props}<w:t xml:space=\"preserve\">{escaped}</w:t></w:r></w:p>"


def _section_properties() -> str:
    return (
        "<w:sectPr>"
        '<w:pgSz w:w="12240" w:h="15840"/>'
        '<w:pgMar w:top="720" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" w:gutter="0"/>'
        "</w:sectPr>"
    )


def _o1b_section_properties() -> str:
    return (
        "<w:sectPr>"
        '<w:footerReference w:type="default" r:id="rIdFooter"/>'
        '<w:pgSz w:w="12240" w:h="15840"/>'
        '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" w:gutter="0"/>'
        "</w:sectPr>"
    )


def _styles_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/><w:rPr><w:rFonts w:ascii="Aptos" w:hAnsi="Aptos" w:eastAsia="Aptos"/><w:sz w:val="24"/></w:rPr><w:pPr><w:spacing w:after="160" w:line="278" w:lineRule="auto"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="24"/></w:rPr><w:pPr><w:spacing w:before="0" w:after="160"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="24"/></w:rPr><w:pPr><w:spacing w:before="0" w:after="160"/><w:keepNext/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="24"/></w:rPr><w:pPr><w:spacing w:before="0" w:after="160"/><w:keepNext/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading3"><w:name w:val="heading 3"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="24"/></w:rPr><w:pPr><w:spacing w:before="0" w:after="160"/><w:keepNext/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="SpecialSection"><w:name w:val="Special section"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:sz w:val="32"/></w:rPr><w:pPr><w:spacing w:before="0" w:after="160"/><w:keepNext/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Meta"><w:name w:val="Memo metadata"/><w:basedOn w:val="Normal"/><w:rPr><w:i/><w:sz w:val="22"/></w:rPr><w:pPr><w:spacing w:after="60"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Note"><w:name w:val="Draft note"/><w:basedOn w:val="Normal"/><w:rPr><w:color w:val="666666"/><w:sz w:val="22"/></w:rPr><w:pPr><w:spacing w:before="120" w:after="180"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Placeholder"><w:name w:val="Script placeholder"/><w:basedOn w:val="Normal"/><w:rPr><w:color w:val="666666"/></w:rPr><w:pPr><w:spacing w:after="160"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Bullet"><w:name w:val="Script bullet"/><w:basedOn w:val="Normal"/><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr></w:style>
</w:styles>"""


def _o1b_styles_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="Times New Roman"/><w:sz w:val="24"/></w:rPr><w:pPr><w:jc w:val="both"/><w:spacing w:before="0" w:after="160" w:line="360" w:lineRule="auto"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="32"/></w:rPr><w:pPr><w:jc w:val="center"/><w:spacing w:before="0" w:after="240"/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="32"/></w:rPr><w:pPr><w:spacing w:before="240" w:after="160"/><w:keepNext/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="28"/></w:rPr><w:pPr><w:spacing w:before="200" w:after="120"/><w:keepNext/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="Heading3"><w:name w:val="heading 3"/><w:basedOn w:val="Normal"/><w:qFormat/><w:rPr><w:b/><w:sz w:val="24"/></w:rPr><w:pPr><w:spacing w:before="160" w:after="100"/><w:keepNext/></w:pPr></w:style>
  <w:style w:type="paragraph" w:styleId="SpecialSection"><w:name w:val="Special section"/><w:basedOn w:val="Heading1"/><w:qFormat/></w:style>
  <w:style w:type="paragraph" w:styleId="Meta"><w:name w:val="Memo metadata"/><w:basedOn w:val="Normal"/><w:rPr><w:i/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Note"><w:name w:val="Draft note"/><w:basedOn w:val="Normal"/><w:rPr><w:color w:val="666666"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Placeholder"><w:name w:val="Script placeholder"/><w:basedOn w:val="Normal"/><w:rPr><w:color w:val="666666"/></w:rPr></w:style>
  <w:style w:type="paragraph" w:styleId="Bullet"><w:name w:val="Script bullet"/><w:basedOn w:val="Normal"/><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr></w:pPr></w:style>
</w:styles>"""


def _content_types_xml(*, include_footer: bool = False) -> str:
    footer = (
        '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
        if include_footer
        else ""
    )
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
  <Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>""".replace("</Types>", footer + "</Types>")


def _rels_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>"""


def _document_rels_xml(*, include_footer: bool = False) -> str:
    footer = (
        '<Relationship Id="rIdFooter" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>'
        if include_footer
        else ""
    )
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
  <Relationship Id="rIdNumbering" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>
</Relationships>""".replace("</Relationships>", footer + "</Relationships>")


def _o1b_footer_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:p><w:pPr><w:jc w:val="right"/></w:pPr>
    <w:r><w:fldChar w:fldCharType="begin"/></w:r>
    <w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>
    <w:r><w:fldChar w:fldCharType="end"/></w:r>
  </w:p>
</w:ftr>"""


def _numbering_xml() -> str:
    overrides = []
    for ordinal in range(1, 11):
        overrides.append(
            f'<w:num w:numId="{100 + ordinal}"><w:abstractNumId w:val="1"/>'
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{ordinal}"/></w:lvlOverride></w:num>'
        )
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:abstractNum w:abstractNumId="0">
    <w:multiLevelType w:val="singleLevel"/>
    <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1)"/><w:lvlJc w:val="left"/><w:pPr><w:tabs><w:tab w:val="num" w:pos="360"/></w:tabs><w:ind w:left="360" w:hanging="360"/></w:pPr></w:lvl>
  </w:abstractNum>
  <w:abstractNum w:abstractNumId="1">
    <w:multiLevelType w:val="singleLevel"/>
    <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="lowerRoman"/><w:lvlText w:val="(%1)"/><w:lvlJc w:val="left"/><w:pPr><w:tabs><w:tab w:val="num" w:pos="360"/></w:tabs><w:ind w:left="360" w:hanging="360"/></w:pPr></w:lvl>
  </w:abstractNum>
  <w:abstractNum w:abstractNumId="2">
    <w:multiLevelType w:val="singleLevel"/>
    <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="&#8226;"/><w:lvlJc w:val="left"/><w:pPr><w:tabs><w:tab w:val="num" w:pos="540"/></w:tabs><w:ind w:left="540" w:hanging="270"/></w:pPr></w:lvl>
  </w:abstractNum>
  <w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
  <w:num w:numId="2"><w:abstractNumId w:val="2"/></w:num>
  """ + "".join(overrides) + """
</w:numbering>"""


def _app_xml() -> str:
    return """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><Application>Petitions Builder</Application></Properties>"""


def _core_xml(case_id: str) -> str:
    now = datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    title = html.escape(f"Working memorandum {case_id}")
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"><dc:title>{title}</dc:title><dc:creator>Petitions Builder</dc:creator><cp:lastModifiedBy>Petitions Builder</cp:lastModifiedBy><dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created><dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified></cp:coreProperties>"""
