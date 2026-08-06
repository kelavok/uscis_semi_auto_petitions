from __future__ import annotations

import csv
import hashlib
import json
import logging
import os
import re
import shutil
import subprocess
from dataclasses import dataclass
from html import escape
from pathlib import Path, PurePosixPath

from .cli_support import PROJECT_ROOT
from .evidence import INDEX_FIELDS, scan_documents
from .file_rules import prompt_sidecar_kind
from .workflow import (
    _case_path_value,
    _citation_plan_for_step,
    _normalize_slashes,
    extract_docx_text,
    find_step,
    format_missing_selected_documents_message,
    load_case,
    load_yaml_file,
    LoadedCase,
    missing_selected_documents_from_output,
    PromptOptions,
)


EXHIBIT_FIELDS = [
    "exhibit_id",
    "exhibit_number",
    "parent_exhibit_id",
    "display_title",
    "evidentiary_thesis",
    "task_type",
    "memo_section",
    "separator_title_type",
    "document_ids",
    "original_translation_order",
    "final_bundle_order",
    "user_approval_status",
    "manual_edit_lock",
    "notes",
]

TEXT_SOURCE_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}
DOCX_SOURCE_EXTENSIONS = {".docx"}
IMAGE_SOURCE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
PDF_SOURCE_EXTENSIONS = {".pdf"}
logging.getLogger("pypdf").setLevel(logging.ERROR)
AUTO_DOCUMENT_INDEX_NOTE = "Auto-assigned from validated LLM outputs."
AUTO_EXHIBIT_INDEX_NOTE = "Generated/updated from document_index.csv."
EXHIBIT_NUMBER_BY_ROLE = {
    "awards": "1",
    "memberships": "2",
    "media": "3",
    "judging": "4",
    "original_contribution": "5",
    "scholarly_articles": "6",
    "exhibitions": "7",
    "leading_critical_role": "8",
    "high_salary": "9",
    "commercial_success": "10",
    "employment_plan": "11",
}


@dataclass(frozen=True)
class ExhibitIndexSummary:
    exhibit_index_path: Path
    document_index_path: Path
    documents_seen: int
    documents_with_exhibit_number: int
    exhibits_created: int
    exhibits_updated: int
    locked_exhibits_skipped: int
    document_orders_written: int
    documents_without_exhibit_number: int


@dataclass(frozen=True)
class SeparatorSummary:
    separators_dir: Path
    exhibits_seen: int
    exhibit_pages_written: int
    episode_pages_written: int
    document_pages_written: int
    missing_document_ids: int
    manifest_path: Path


@dataclass(frozen=True)
class SeparatorPdfSummary:
    source_dir: Path
    output_dir: Path
    markdown_files_seen: int
    pdf_pages_written: int
    manifest_path: Path


@dataclass(frozen=True)
class BundlePlanSummary:
    plan_csv_path: Path
    plan_md_path: Path
    total_items: int
    ready_items: int
    missing_items: int
    unsupported_items: int
    source_documents: int
    separator_pdfs: int


@dataclass(frozen=True)
class BundleBuildSummary:
    final_pdf_path: Path
    plan_csv_path: Path
    merged_items: int
    converted_items: int
    skipped_items: int


@dataclass(frozen=True)
class FinalFilingSummary:
    final_pdf_path: Path
    numbered_memo_docx_path: Path
    numbered_memo_pdf_path: Path
    evidence_bundle_pdf_path: Path
    report_path: Path
    memo_pages: int
    bundle_pages: int
    final_pages: int
    placeholders_seen: int
    placeholders_resolved: int
    placeholders_unresolved: int
    iterations: int


@dataclass(frozen=True)
class LayoutIndexStatus:
    indexed_documents: int
    used_document_references: int
    unique_used_documents: int
    assigned_used_documents: int
    exhibit_count: int
    stale_document_ids: tuple[str, ...]
    assignment_conflicts: tuple[str, ...]
    indexed_sidecars: tuple[str, ...]
    unsupported_documents: tuple[str, ...]

    @property
    def refresh_required(self) -> bool:
        return bool(
            self.unique_used_documents
            and (
                not self.exhibit_count
                or self.indexed_sidecars
            )
        )

    @property
    def ready_for_separators(self) -> bool:
        return bool(self.unique_used_documents and not self.refresh_required)


@dataclass(frozen=True)
class LayoutIndexRefreshSummary:
    scanned_files: int
    auxiliary_rows_removed: int
    replacement_documents_rebound: int
    documents_assigned: int
    exhibit_summary: ExhibitIndexSummary
    status: LayoutIndexStatus


@dataclass(frozen=True)
class BundleSelection:
    exhibit_numbers: tuple[str, ...]
    document_ids: tuple[str, ...]


@dataclass(frozen=True)
class BundlePreparationSummary:
    selection: BundleSelection
    exhibit_pages: int
    document_pages: int
    separator_pdfs: int
    plan: BundlePlanSummary


@dataclass(frozen=True)
class BundlePreparationStatus:
    prepared: bool
    ready_to_build: bool
    selected_exhibits: int
    selected_documents: int
    expected_separator_pdfs: int
    rendered_separator_pdfs: int
    missing_items: int
    unsupported_items: int
    reason: str


@dataclass(frozen=True)
class EvidenceLayoutPlan:
    document_to_exhibit: dict[str, str]
    document_order: dict[str, int]
    exhibit_order: dict[str, int]
    document_titles: dict[str, str]
    conflicts: list[str]
    references: int
    exhibit_metadata: dict[str, tuple[str, str]]


def bundle_catalog(case_id: str) -> list[dict[str, object]]:
    loaded = load_case(case_id)
    document_rows = _read_csv(
        loaded.case_dir / _case_path_value(loaded.config, "document_index"), INDEX_FIELDS
    )
    exhibit_rows = _read_csv(
        loaded.case_dir / _case_path_value(loaded.config, "exhibit_index"), EXHIBIT_FIELDS
    )
    documents_by_id = {
        row.get("document_id", ""): row for row in document_rows if row.get("document_id")
    }
    catalog: list[dict[str, object]] = []
    for exhibit in sorted(exhibit_rows, key=_exhibit_sort_key):
        exhibit_number = exhibit.get("exhibit_number", "").strip()
        if not exhibit_number:
            continue
        document_ids = [
            part.strip()
            for part in exhibit.get("document_ids", "").split(";")
            if part.strip()
        ]
        documents = [documents_by_id[document_id] for document_id in document_ids if document_id in documents_by_id]
        catalog.append(
            {
                "exhibit_number": exhibit_number,
                "display_title": exhibit.get("display_title", "") or f"Exhibit {exhibit_number}",
                "documents": documents,
            }
        )
    return catalog


def load_bundle_selection(case_id: str) -> BundleSelection:
    loaded = load_case(case_id)
    catalog = bundle_catalog(case_id)
    all_exhibits = tuple(str(item["exhibit_number"]) for item in catalog)
    all_documents = tuple(
        str(document.get("document_id", ""))
        for exhibit in catalog
        for document in exhibit["documents"]
        if document.get("document_id", "")
    )
    path = _bundle_selection_path(loaded)
    if not path.exists():
        return BundleSelection(all_exhibits, all_documents)
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError, TypeError):
        return BundleSelection(all_exhibits, all_documents)
    valid_exhibits = set(all_exhibits)
    valid_documents = set(all_documents)
    exhibits = tuple(
        value for value in (str(item) for item in data.get("selected_exhibits", [])) if value in valid_exhibits
    )
    documents = tuple(
        value for value in (str(item) for item in data.get("selected_document_ids", [])) if value in valid_documents
    )
    return BundleSelection(exhibits, documents)


def prepare_selected_bundle(
    case_id: str, selected_exhibits: list[str], selected_document_ids: list[str]
) -> BundlePreparationSummary:
    loaded = load_case(case_id)
    selection = _validated_bundle_selection(case_id, selected_exhibits, selected_document_ids)
    _save_bundle_selection(loaded, selection)
    separators = generate_separator_pages(case_id, selection=selection)
    rendered = render_separator_pdfs(case_id)
    plan = build_bundle_plan(case_id, selection=selection)
    state = {
        "selection_digest": _selection_digest(selection),
        "index_digest": _bundle_index_digest(loaded),
        "selected_exhibits": list(selection.exhibit_numbers),
        "selected_document_ids": list(selection.document_ids),
        "expected_separator_pdfs": (
            separators.exhibit_pages_written
            + separators.episode_pages_written
            + separators.document_pages_written
        ),
        "rendered_separator_pdfs": rendered.pdf_pages_written,
        "missing_items": plan.missing_items,
        "unsupported_items": plan.unsupported_items,
    }
    state_path = _bundle_preparation_path(loaded)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(json.dumps(state, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return BundlePreparationSummary(
        selection=selection,
        exhibit_pages=separators.exhibit_pages_written,
        document_pages=separators.document_pages_written,
        separator_pdfs=rendered.pdf_pages_written,
        plan=plan,
    )


def inspect_bundle_preparation(case_id: str) -> BundlePreparationStatus:
    loaded = load_case(case_id)
    selection = load_bundle_selection(case_id)
    state_path = _bundle_preparation_path(loaded)
    if not state_path.exists():
        return BundlePreparationStatus(
            prepared=False,
            ready_to_build=False,
            selected_exhibits=len(selection.exhibit_numbers),
            selected_documents=len(selection.document_ids),
            expected_separator_pdfs=0,
            rendered_separator_pdfs=0,
            missing_items=0,
            unsupported_items=0,
            reason="Prepare the selected bundle first.",
        )
    try:
        state = json.loads(state_path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError, TypeError):
        return BundlePreparationStatus(False, False, len(selection.exhibit_numbers), len(selection.document_ids), 0, 0, 0, 0, "Preparation state is unreadable; prepare again.")
    if state.get("selection_digest") != _selection_digest(selection):
        reason = "The bundle selection changed; prepare again."
    elif state.get("index_digest") != _bundle_index_digest(loaded):
        reason = "The evidence indexes changed; prepare again."
    elif int(state.get("rendered_separator_pdfs", 0)) != int(state.get("expected_separator_pdfs", 0)):
        reason = "Not all separator PDFs were rendered; prepare again."
    else:
        reason = "Ready to build."
    ready = reason == "Ready to build."
    return BundlePreparationStatus(
        prepared=True,
        ready_to_build=ready,
        selected_exhibits=len(selection.exhibit_numbers),
        selected_documents=len(selection.document_ids),
        expected_separator_pdfs=int(state.get("expected_separator_pdfs", 0)),
        rendered_separator_pdfs=int(state.get("rendered_separator_pdfs", 0)),
        missing_items=int(state.get("missing_items", 0)),
        unsupported_items=int(state.get("unsupported_items", 0)),
        reason=reason,
    )


def inspect_layout_index_status(case_id: str) -> LayoutIndexStatus:
    loaded = load_case(case_id)
    document_index_path = loaded.case_dir / _case_path_value(loaded.config, "document_index")
    exhibit_index_path = loaded.case_dir / _case_path_value(loaded.config, "exhibit_index")
    document_rows = _read_csv(document_index_path, INDEX_FIELDS)
    plan = _desired_exhibit_assignments(loaded)
    desired = plan.document_to_exhibit
    rows_by_id = {row.get("document_id", ""): row for row in document_rows if row.get("document_id")}
    stale = sorted(document_id for document_id in desired if document_id not in rows_by_id)
    assigned = sum(
        1
        for document_id, exhibit_number in desired.items()
        if document_id in rows_by_id
        and rows_by_id[document_id].get("exhibit_number", "").strip() == exhibit_number
    )
    indexed_sidecars = sorted(
        row.get("document_id", "")
        for row in document_rows
        if prompt_sidecar_kind(Path(row.get("file_path", "")))
    )
    unsupported: list[str] = []
    for document_id in desired:
        row = rows_by_id.get(document_id)
        if not row:
            continue
        status, note = _source_document_status(loaded.case_dir / row.get("file_path", ""))
        if status in {"missing", "unsupported"}:
            title = row.get("display_title", "") or row.get("original_file_name", "")
            unsupported.append(
                f"{document_id}: {title} — {note} [{row.get('file_path', '')}]"
            )
    exhibit_rows = _read_csv(exhibit_index_path, EXHIBIT_FIELDS)
    return LayoutIndexStatus(
        indexed_documents=len(document_rows),
        used_document_references=plan.references,
        unique_used_documents=len(desired),
        assigned_used_documents=assigned,
        exhibit_count=sum(bool(row.get("exhibit_number", "").strip()) for row in exhibit_rows),
        stale_document_ids=tuple(stale),
        assignment_conflicts=tuple(plan.conflicts),
        indexed_sidecars=tuple(indexed_sidecars),
        unsupported_documents=tuple(sorted(unsupported)),
    )


def refresh_layout_indexes(case_id: str) -> LayoutIndexRefreshSummary:
    scan = scan_documents(case_id)
    loaded = load_case(case_id)
    document_index_path = loaded.case_dir / _case_path_value(loaded.config, "document_index")
    exhibit_index_path = loaded.case_dir / _case_path_value(loaded.config, "exhibit_index")
    document_rows = _read_csv(document_index_path, INDEX_FIELDS)
    plan = _desired_exhibit_assignments(loaded)
    desired = plan.document_to_exhibit
    conflicts = list(plan.conflicts)
    exhibit_metadata = plan.exhibit_metadata
    replacements_rebound = _rebind_supported_replacements(loaded, document_rows, desired)
    rows_by_id = {row.get("document_id", ""): row for row in document_rows if row.get("document_id")}

    for row in document_rows:
        if AUTO_DOCUMENT_INDEX_NOTE in row.get("notes", "") and not _truthy(row.get("manual_edit_lock", "")):
            row["exhibit_number"] = ""
            row["final_bundle_order"] = ""

    for document_id, title in plan.document_titles.items():
        row = rows_by_id.get(document_id)
        if row and title and not _truthy(row.get("manual_edit_lock", "")):
            row["display_title"] = title

    numbered_context = _numbered_document_sort_context(
        (rows_by_id[document_id] for document_id in desired if document_id in rows_by_id),
        fallback_order=plan.document_order,
        document_to_exhibit=desired,
    )
    assignments = sorted(
        desired.items(),
        key=lambda item: _assignment_sort_key(item, plan, numbered_context),
    )
    assigned = 0
    for bundle_order, (document_id, exhibit_number) in enumerate(assignments, start=1):
        row = rows_by_id.get(document_id)
        if not row:
            continue
        existing = row.get("exhibit_number", "").strip()
        if _truthy(row.get("manual_edit_lock", "")) and existing and existing != exhibit_number:
            conflicts.append(
                f"{document_id} is locked to Exhibit {existing}, but validated output requires Exhibit {exhibit_number}."
            )
            continue
        if not _truthy(row.get("manual_edit_lock", "")):
            row["exhibit_number"] = exhibit_number
            row["final_bundle_order"] = str(bundle_order)
            row["user_approval_status"] = row.get("user_approval_status") or "pending"
            row["notes"] = _merge_note(row.get("notes", ""), AUTO_DOCUMENT_INDEX_NOTE)
        assigned += 1
    _write_csv(document_index_path, document_rows, INDEX_FIELDS)

    desired_exhibits = set(desired.values())
    exhibit_rows = _read_csv(exhibit_index_path, EXHIBIT_FIELDS)
    exhibit_rows = [
        row
        for row in exhibit_rows
        if row.get("exhibit_number", "").strip() in desired_exhibits
        or _truthy(row.get("manual_edit_lock", ""))
    ]
    existing_exhibit_numbers = {
        row.get("exhibit_number", "").strip()
        for row in exhibit_rows
        if row.get("exhibit_number", "").strip()
    }
    next_exhibit_id = _next_exhibit_number(exhibit_rows)
    for exhibit_number in sorted(desired_exhibits, key=_natural_sort_key):
        if exhibit_number in existing_exhibit_numbers:
            continue
        title, memo_section = exhibit_metadata.get(exhibit_number, ("", ""))
        row = _blank_exhibit_row()
        row["exhibit_id"] = f"EXH{next_exhibit_id:03d}"
        next_exhibit_id += 1
        row["exhibit_number"] = exhibit_number
        row["display_title"] = title or f"Exhibit {exhibit_number}"
        row["memo_section"] = memo_section
        row["task_type"] = str(loaded.config.get("task_type", ""))
        row["separator_title_type"] = "exhibit"
        row["original_translation_order"] = str(
            loaded.config.get("translation_order", "original_then_translation")
        )
        row["user_approval_status"] = "pending"
        row["manual_edit_lock"] = "false"
        row["notes"] = _merge_note(
            row.get("notes", ""),
            "Generated from validated LLM output even though no indexed source documents were assigned.",
        )
        exhibit_rows.append(row)
        existing_exhibit_numbers.add(exhibit_number)
    for row in exhibit_rows:
        exhibit_number = row.get("exhibit_number", "").strip()
        title, memo_section = exhibit_metadata.get(exhibit_number, ("", ""))
        if title and not _truthy(row.get("manual_edit_lock", "")):
            row["display_title"] = title
        if memo_section and not _truthy(row.get("manual_edit_lock", "")):
            row["memo_section"] = memo_section
    _write_csv(exhibit_index_path, exhibit_rows, EXHIBIT_FIELDS)
    exhibit_summary = build_exhibit_index(case_id)
    _invalidate_bundle_preparation(loaded)
    status = inspect_layout_index_status(case_id)
    combined_conflicts = tuple(dict.fromkeys([*status.assignment_conflicts, *conflicts]))
    if combined_conflicts != status.assignment_conflicts:
        status = LayoutIndexStatus(
            indexed_documents=status.indexed_documents,
            used_document_references=status.used_document_references,
            unique_used_documents=status.unique_used_documents,
            assigned_used_documents=status.assigned_used_documents,
            exhibit_count=status.exhibit_count,
            stale_document_ids=status.stale_document_ids,
            assignment_conflicts=combined_conflicts,
            indexed_sidecars=status.indexed_sidecars,
            unsupported_documents=status.unsupported_documents,
        )
    return LayoutIndexRefreshSummary(
        scanned_files=scan.scanned_files,
        auxiliary_rows_removed=scan.removed_rows,
        replacement_documents_rebound=replacements_rebound,
        documents_assigned=assigned,
        exhibit_summary=exhibit_summary,
        status=status,
    )


def _desired_exhibit_assignments(loaded: LoadedCase) -> EvidenceLayoutPlan:
    validated_root = loaded.case_dir / _case_path_value(loaded.config, "validated_outputs")
    files = {path.stem: path for path in validated_root.glob("*.json")}
    try:
        from .stages import build_llm_stage

        ordered_stems = [unit.key for unit in build_llm_stage(loaded.case_id).units]
    except Exception:  # noqa: BLE001 - index diagnostics must survive a damaged drafting timeline.
        ordered_stems = []
    ordered_stems.extend(stem for stem in sorted(files) if stem not in ordered_stems)
    desired: dict[str, str] = {}
    document_order: dict[str, int] = {}
    exhibit_order: dict[str, int] = {}
    document_titles: dict[str, str] = {}
    conflicts: list[str] = []
    exhibit_metadata: dict[str, tuple[str, str]] = {}
    references = 0
    next_document_order = 1
    next_exhibit_order = 1
    for stem in ordered_stems:
        path = files.get(stem)
        if not path:
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, ValueError, TypeError) as exc:
            conflicts.append(f"{path.name} cannot be read: {exc}")
            continue
        used_documents = data.get("used_documents", [])
        if not isinstance(used_documents, list) or not used_documents:
            continue
        step_id = str(data.get("step_id", ""))
        episode_id = str(data.get("episode_id", ""))
        try:
            step = find_step(loaded.workflow, step_id)
            missing_selected = missing_selected_documents_from_output(
                data, loaded, step, PromptOptions(episode_id=episode_id)
            )
        except SystemExit as exc:
            missing_selected = []
            conflicts.append(f"{path.name} completeness check failed: {exc}")
        except Exception as exc:  # noqa: BLE001 - layout diagnostics should keep scanning other outputs.
            missing_selected = []
            conflicts.append(f"{path.name} completeness check failed: {exc}")
        if missing_selected:
            conflicts.append(f"{path.name}: {format_missing_selected_documents_message(missing_selected)}")
        exhibit_number = _exhibit_number_for_output(loaded, data)
        if not exhibit_number:
            conflicts.append(f"{path.name} uses documents but has no deterministic Exhibit mapping.")
            continue
        if exhibit_number not in exhibit_order:
            exhibit_order[exhibit_number] = next_exhibit_order
            next_exhibit_order += 1
        title, memo_section = _exhibit_metadata_for_output(loaded, data)
        existing_title, existing_section = exhibit_metadata.get(exhibit_number, ("", ""))
        exhibit_metadata[exhibit_number] = (
            existing_title or title,
            existing_section or memo_section,
        )
        for item in _ordered_used_documents(data):
            if not isinstance(item, dict):
                continue
            document_id = str(item.get("document_id", "")).strip()
            if not document_id:
                continue
            references += 1
            document_title = str(item.get("document_title", "")).strip()
            if document_title and document_id not in document_titles:
                document_titles[document_id] = document_title
            existing = desired.get(document_id)
            if existing and existing != exhibit_number:
                conflicts.append(
                    f"{document_id} is used by both Exhibit {existing} and Exhibit {exhibit_number}."
                )
                continue
            if document_id not in desired:
                desired[document_id] = exhibit_number
                document_order[document_id] = next_document_order
                next_document_order += 1
    return EvidenceLayoutPlan(
        document_to_exhibit=desired,
        document_order=document_order,
        exhibit_order=exhibit_order,
        document_titles=document_titles,
        conflicts=conflicts,
        references=references,
        exhibit_metadata=exhibit_metadata,
    )


def _ordered_used_documents(data: dict[str, object]) -> list[dict[str, object]]:
    used_documents = [
        item for item in data.get("used_documents", []) if isinstance(item, dict)
    ]
    if len(used_documents) < 2:
        return used_documents

    narrative_order = _document_title_order_from_draft(
        str(data.get("draft_text", "")), used_documents
    )
    if not narrative_order:
        return used_documents
    ordered_items = sorted(
        enumerate(used_documents),
        key=lambda pair: (
            narrative_order.get(str(pair[1].get("document_id", "")).strip(), 10**9),
            pair[0],
        ),
    )
    return [item for _offset, item in ordered_items]


def _document_title_order_from_draft(
    draft_text: str, used_documents: list[dict[str, object]]
) -> dict[str, int]:
    titles_by_id = {
        str(item.get("document_id", "")).strip(): _normalize_title_match_text(
            str(item.get("document_title", "")).strip()
        )
        for item in used_documents
        if str(item.get("document_id", "")).strip()
        and str(item.get("document_title", "")).strip()
    }
    if not titles_by_id:
        return {}

    ordered: dict[str, int] = {}
    for line in draft_text.splitlines():
        normalized_line = _normalize_title_match_text(line)
        if not normalized_line:
            continue
        for document_id, title in titles_by_id.items():
            if document_id in ordered or not title:
                continue
            if title in normalized_line:
                ordered[document_id] = len(ordered) + 1
    if len(ordered) >= 2:
        return ordered

    lowered = draft_text.casefold()
    positions: list[tuple[int, str]] = []
    for item in used_documents:
        document_id = str(item.get("document_id", "")).strip()
        title = str(item.get("document_title", "")).strip()
        if not document_id or not title:
            continue
        position = lowered.find(title.casefold())
        if position >= 0:
            positions.append((position, document_id))
    if len(positions) < 2:
        return {}
    return {
        document_id: index
        for index, (_position, document_id) in enumerate(sorted(positions), start=1)
    }


def _normalize_title_match_text(value: str) -> str:
    value = value.casefold()
    value = value.replace("“", '"').replace("”", '"').replace("’", "'")
    value = re.sub(
        r"\b(?:original|english translation|original and english translation)\b",
        " ",
        value,
    )
    value = re.sub(r"^\s*(?:\d+(?:\.\d+)*\.?|\d+[.)]|[-*•])\s*", " ", value)
    value = re.sub(r"[^a-z0-9а-яёіїєґ\"']+", " ", value, flags=re.IGNORECASE)
    return " ".join(value.split())


def _exhibit_metadata_for_output(
    loaded: LoadedCase, data: dict[str, object]
) -> tuple[str, str]:
    step_id = str(data.get("step_id", ""))
    episode_id = str(data.get("episode_id", ""))
    if str(loaded.config.get("task_type", "")) == "eb1a_rfe_response" and episode_id:
        from .rfe_strategy import get_strategy_unit

        unit = get_strategy_unit(loaded.case_dir, episode_id, loaded.config)
        section_title = str(unit.get("section_title", "")).strip()
        role = str(unit.get("criterion_role", "")).strip()
        return section_title or str(unit.get("title", "")).strip(), role
    title_from_structure, role_from_structure = _criterion_metadata_from_working_structure(
        loaded, step_id
    )
    if title_from_structure:
        return title_from_structure, role_from_structure
    try:
        step = find_step(loaded.workflow, step_id)
    except SystemExit:
        return "", ""
    title = str(step.get("exhibit_title", "") or step.get("title", "")).strip()
    roles = _normalize_step_roles(step)
    memo_section = str(
        step.get("memo_section", "") or step.get("criterion_role", "") or (roles[0] if len(roles) == 1 else "")
    ).strip()
    return title, memo_section


def _criterion_metadata_from_working_structure(
    loaded: LoadedCase, step_id: str
) -> tuple[str, str]:
    task_type = str(loaded.config.get("task_type", ""))
    if task_type not in {"eb1a_petition", "o1b_petition"}:
        return "", ""
    value = str(loaded.config.get("working_document_template", "")).strip()
    if not value:
        return "", ""
    path = Path(value)
    resolved = path if path.is_absolute() else PROJECT_ROOT / path
    if not resolved.exists():
        return "", ""
    try:
        structure = load_yaml_file(resolved)
    except Exception:  # noqa: BLE001 - metadata fallback should not block bundle diagnostics.
        return "", ""
    criteria = structure.get("criteria", {}) if isinstance(structure, dict) else {}
    if not isinstance(criteria, dict):
        return "", ""
    for role, criterion in criteria.items():
        if not isinstance(criterion, dict):
            continue
        step_ids = criterion.get("step_ids", [])
        if not isinstance(step_ids, list):
            step_ids = [step_ids]
        if step_id not in {str(item).strip() for item in step_ids}:
            continue
        title = str(criterion.get("title", "")).strip()
        return title, str(role).strip()
    return "", ""


def _normalize_step_roles(step: dict[str, object]) -> list[str]:
    roles = step.get("evidence_folder_roles", [])
    if not isinstance(roles, list):
        roles = [roles]
    return [str(role).strip() for role in roles if str(role).strip()]


def _exhibit_number_for_output(loaded: LoadedCase, data: dict[str, object]) -> str:
    step_id = str(data.get("step_id", ""))
    episode_id = str(data.get("episode_id", ""))
    if str(loaded.config.get("task_type", "")) == "eb1a_rfe_response" and episode_id:
        from .rfe_strategy import get_strategy_unit

        unit = get_strategy_unit(loaded.case_dir, episode_id, loaded.config)
        return EXHIBIT_NUMBER_BY_ROLE.get(str(unit.get("criterion_role", "")), "")
    try:
        step = find_step(loaded.workflow, step_id)
    except SystemExit:
        return ""
    citation = _citation_plan_for_step(loaded, step, PromptOptions(episode_id=episode_id))
    return str(citation.get("exhibit_number", "")).strip()


def _rebind_supported_replacements(
    loaded: LoadedCase, document_rows: list[dict[str, str]], desired: dict[str, str]
) -> int:
    """Keep a used DOC id when an unsupported/missing source is replaced by a sibling PDF."""
    rows_by_id = {row.get("document_id", ""): row for row in document_rows if row.get("document_id")}
    rebound = 0
    for document_id in desired:
        row = rows_by_id.get(document_id)
        if not row:
            continue
        source_path = loaded.case_dir / row.get("file_path", "")
        status, _note = _source_document_status(source_path)
        if status not in {"missing", "unsupported"}:
            continue
        source_relative = Path(row.get("file_path", ""))
        candidates: list[dict[str, str]] = []
        for candidate in document_rows:
            if candidate is row:
                continue
            candidate_relative = Path(candidate.get("file_path", ""))
            if (
                candidate_relative.parent != source_relative.parent
                or candidate_relative.stem.casefold() != source_relative.stem.casefold()
            ):
                continue
            candidate_status, _candidate_note = _source_document_status(
                loaded.case_dir / candidate.get("file_path", "")
            )
            if candidate_status == "ready_pdf":
                candidates.append(candidate)
        if len(candidates) != 1:
            continue
        replacement = candidates[0]
        old_name = row.get("original_file_name", "") or source_relative.name
        for field in (
            "original_file_name",
            "file_path",
            "document_type",
            "extraction_status",
            "text_extraction_path",
            "source_fingerprint",
            "last_scanned_at",
        ):
            row[field] = replacement.get(field, "")
        row["notes"] = _merge_note(
            row.get("notes", ""),
            f"Rebound from unsupported/missing source {old_name} to replacement PDF while preserving {document_id}.",
        )
        document_rows.remove(replacement)
        rows_by_id.pop(replacement.get("document_id", ""), None)
        rebound += 1
    return rebound


def build_exhibit_index(case_id: str) -> ExhibitIndexSummary:
    loaded = load_case(case_id)
    document_index_path = loaded.case_dir / _case_path_value(loaded.config, "document_index")
    exhibit_index_path = loaded.case_dir / _case_path_value(loaded.config, "exhibit_index")

    document_rows = _read_csv(document_index_path, INDEX_FIELDS)
    exhibit_rows = _read_csv(exhibit_index_path, EXHIBIT_FIELDS)
    exhibits_by_number = {
        row.get("exhibit_number", ""): row
        for row in exhibit_rows
        if row.get("exhibit_number")
    }

    grouped: dict[str, list[dict[str, str]]] = {}
    for row in document_rows:
        exhibit_number = row.get("exhibit_number", "").strip()
        if exhibit_number:
            grouped.setdefault(exhibit_number, []).append(row)

    exhibits_created = 0
    exhibits_updated = 0
    locked_exhibits_skipped = 0
    document_orders_written = 0

    ordered_exhibit_numbers = sorted(
        grouped,
        key=lambda number: (
            _first_document_order(grouped[number]),
            _natural_sort_key(number),
        ),
    )
    next_exhibit_id = _next_exhibit_number(exhibit_rows)
    global_document_order = 1

    for exhibit_order, exhibit_number in enumerate(ordered_exhibit_numbers, start=1):
        ordered_documents = order_documents_original_then_translation(grouped[exhibit_number])
        row = exhibits_by_number.get(exhibit_number)
        if row is None:
            row = _blank_exhibit_row()
            row["exhibit_id"] = f"EXH{next_exhibit_id:03d}"
            next_exhibit_id += 1
            row["exhibit_number"] = exhibit_number
            row["display_title"] = f"Exhibit {exhibit_number}"
            row["separator_title_type"] = "exhibit"
            row["user_approval_status"] = "pending"
            row["manual_edit_lock"] = "false"
            exhibit_rows.append(row)
            exhibits_by_number[exhibit_number] = row
            exhibits_created += 1
        elif _truthy(row.get("manual_edit_lock", "")):
            locked_exhibits_skipped += 1
            for document in ordered_documents:
                if not _truthy(document.get("manual_edit_lock", "")):
                    if not document.get("final_bundle_order"):
                        document["final_bundle_order"] = str(global_document_order)
                        document_orders_written += 1
                    global_document_order += 1
            continue
        else:
            exhibits_updated += 1

        row["task_type"] = row.get("task_type") or str(loaded.config.get("task_type", ""))
        row["memo_section"] = row.get("memo_section") or _common_value(ordered_documents, "memo_section_relevance")
        row["separator_title_type"] = row.get("separator_title_type") or "exhibit"
        row["document_ids"] = ";".join(document.get("document_id", "") for document in ordered_documents)
        row["original_translation_order"] = str(
            loaded.config.get("translation_order", "original_then_translation")
        )
        row["final_bundle_order"] = str(exhibit_order)
        row["notes"] = _merge_note(row.get("notes", ""), "Generated/updated from document_index.csv.")

        for document in ordered_documents:
            if _truthy(document.get("manual_edit_lock", "")):
                global_document_order += 1
                continue
            if not document.get("final_bundle_order"):
                document["final_bundle_order"] = str(global_document_order)
                document_orders_written += 1
            global_document_order += 1

    _write_csv(document_index_path, document_rows, INDEX_FIELDS)
    _write_csv(exhibit_index_path, exhibit_rows, EXHIBIT_FIELDS)
    return ExhibitIndexSummary(
        exhibit_index_path=exhibit_index_path,
        document_index_path=document_index_path,
        documents_seen=len(document_rows),
        documents_with_exhibit_number=sum(1 for row in document_rows if row.get("exhibit_number", "").strip()),
        exhibits_created=exhibits_created,
        exhibits_updated=exhibits_updated,
        locked_exhibits_skipped=locked_exhibits_skipped,
        document_orders_written=document_orders_written,
        documents_without_exhibit_number=sum(
            1 for row in document_rows if not row.get("exhibit_number", "").strip()
        ),
    )


def generate_separator_pages(
    case_id: str, selection: BundleSelection | None = None
) -> SeparatorSummary:
    loaded = load_case(case_id)
    document_index_path = loaded.case_dir / _case_path_value(loaded.config, "document_index")
    exhibit_index_path = loaded.case_dir / _case_path_value(loaded.config, "exhibit_index")
    bundle_root = loaded.case_dir / _case_path_value(loaded.config, "bundle_root")
    separators_dir = bundle_root / "separators" / "generated"
    separators_dir.mkdir(parents=True, exist_ok=True)
    for stale_path in separators_dir.glob("*.md"):
        stale_path.unlink()

    selection = selection or load_bundle_selection(case_id)
    selected_exhibits = set(selection.exhibit_numbers)
    selected_documents = set(selection.document_ids)

    document_rows = _read_csv(document_index_path, INDEX_FIELDS)
    exhibit_rows = _read_csv(exhibit_index_path, EXHIBIT_FIELDS)
    if not any(row.get("exhibit_number", "").strip() for row in exhibit_rows):
        raise SystemExit("Exhibit index is empty. Run 'Refresh indexes' before generating separators.")
    documents_by_id = {
        row.get("document_id", ""): row
        for row in document_rows
        if row.get("document_id")
    }

    exhibit_pages_written = 0
    episode_pages_written = 0
    document_pages_written = 0
    missing_document_ids = 0
    manifest_lines = [
        "# Generated separator pages",
        "",
        "Do not edit these generated files directly; rerunning the command may overwrite them.",
        "",
    ]

    ordered_exhibits = [
        exhibit
        for exhibit in sorted(exhibit_rows, key=_exhibit_sort_key)
        if exhibit.get("exhibit_number", "").strip() in selected_exhibits
    ]
    for exhibit_position, exhibit in enumerate(ordered_exhibits, start=1):
        exhibit_number = exhibit.get("exhibit_number", "").strip()
        if not exhibit_number:
            continue
        document_ids = [
            part.strip()
            for part in exhibit.get("document_ids", "").split(";")
            if part.strip() and part.strip() in selected_documents
        ]
        missing_document_ids += sum(1 for document_id in document_ids if document_id not in documents_by_id)
        document_groups = _logical_document_groups(document_ids, documents_by_id)
        exhibit_file = separators_dir / f"{exhibit_position:03d}_exhibit_{_safe_filename(exhibit_number)}.md"
        exhibit_file.write_text(
            _render_exhibit_separator(exhibit, document_groups, loaded.config),
            encoding="utf-8",
        )
        exhibit_pages_written += 1
        manifest_lines.append(f"- {exhibit_file.relative_to(bundle_root).as_posix()}")

        for episode_position, section in enumerate(
            _numbered_episode_document_sections(exhibit, document_groups, loaded.config),
            start=1,
        ):
            episode_title = str(section["episode_title"])
            episode_number = str(section["episode_number"])
            numbered_groups = section["documents"]
            if episode_title:
                episode_file = (
                    separators_dir
                    / f"{exhibit_position:03d}_{episode_position:03d}_000_episode_{_safe_filename(episode_number)}.md"
                )
                episode_file.write_text(
                    _render_episode_separator(exhibit, episode_number, episode_title, numbered_groups),
                    encoding="utf-8",
                )
                episode_pages_written += 1
                manifest_lines.append(f"- {episode_file.relative_to(bundle_root).as_posix()}")

            for document_position, (document_number, document, translations) in enumerate(
                numbered_groups,
                start=1,
            ):
                document_id = document.get("document_id", "")
                document_file = (
                    separators_dir
                    / f"{exhibit_position:03d}_{episode_position:03d}_{document_position:03d}_{_safe_filename(document_id)}.md"
                )
                document_file.write_text(
                    _render_document_separator(
                        exhibit,
                        episode_number,
                        episode_title,
                        document_number,
                        document,
                        translations,
                    ),
                    encoding="utf-8",
                )
                document_pages_written += 1
                manifest_lines.append(f"- {document_file.relative_to(bundle_root).as_posix()}")

    manifest_path = separators_dir / "manifest.md"
    manifest_path.write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    return SeparatorSummary(
        separators_dir=separators_dir,
        exhibits_seen=len(ordered_exhibits),
        exhibit_pages_written=exhibit_pages_written,
        episode_pages_written=episode_pages_written,
        document_pages_written=document_pages_written,
        missing_document_ids=missing_document_ids,
        manifest_path=manifest_path,
    )


def render_separator_pdfs(case_id: str) -> SeparatorPdfSummary:
    loaded = load_case(case_id)
    bundle_root = loaded.case_dir / _case_path_value(loaded.config, "bundle_root")
    source_dir = bundle_root / "separators" / "generated"
    output_dir = bundle_root / "separators" / "pdf"
    if not source_dir.exists():
        raise SystemExit(
            f"Separator markdown folder not found: {source_dir}. "
            "Run `python -m app.bundle separators --case CASE_ID` first."
        )

    try:
        from reportlab.lib import colors  # type: ignore
        from reportlab.lib.enums import TA_CENTER  # type: ignore
        from reportlab.lib.pagesizes import LETTER  # type: ignore
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore
        from reportlab.lib.units import inch  # type: ignore
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # type: ignore
        from reportlab.pdfbase import pdfmetrics  # type: ignore
        from reportlab.pdfbase.ttfonts import TTFont  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "Missing PDF dependency `reportlab`. Install with `pip install .[pdf]` "
            "or run this command with the bundled Codex Python runtime."
        ) from exc

    font_name = _register_pdf_font(pdfmetrics, TTFont, _bundle_font_family(loaded.config))
    output_dir.mkdir(parents=True, exist_ok=True)
    for stale_path in output_dir.glob("*.pdf"):
        stale_path.unlink()
    markdown_files = [
        path
        for path in sorted(source_dir.glob("*.md"))
        if path.name.lower() != "manifest.md"
    ]

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="SeparatorTitle",
            parent=styles["Title"],
            fontName=font_name,
            fontSize=22,
            leading=27,
            alignment=TA_CENTER,
            spaceAfter=0.25 * inch,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SeparatorHeading",
            parent=styles["Heading2"],
            fontName=font_name,
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#222222"),
            spaceBefore=0.12 * inch,
            spaceAfter=0.06 * inch,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SeparatorEpisode",
            parent=styles["Heading3"],
            fontName=font_name,
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#222222"),
            spaceBefore=0.1 * inch,
            spaceAfter=0.06 * inch,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SeparatorBody",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=10.5,
            leading=14,
            spaceAfter=0.06 * inch,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SeparatorMeta",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#555555"),
            spaceAfter=0.05 * inch,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SeparatorPlease",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=12,
            leading=15,
            spaceBefore=0.85 * inch,
            spaceAfter=0.05 * inch,
            leftIndent=0.35 * inch,
        )
    )

    pdf_count = 0
    manifest_lines = [
        "# Generated separator PDFs",
        "",
        "Generated from bundle/separators/generated/*.md.",
        "",
    ]
    for markdown_path in markdown_files:
        pdf_path = output_dir / (markdown_path.stem + ".pdf")
        story = _markdown_to_reportlab_story(markdown_path.read_text(encoding="utf-8-sig"), styles, Spacer)
        document = SimpleDocTemplate(
            str(pdf_path),
            pagesize=LETTER,
            leftMargin=0.75 * inch,
            rightMargin=0.75 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
            title=markdown_path.stem,
        )
        document.build(story)
        pdf_count += 1
        manifest_lines.append(f"- {pdf_path.relative_to(bundle_root).as_posix()}")

    manifest_path = output_dir / "manifest.md"
    manifest_path.write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    return SeparatorPdfSummary(
        source_dir=source_dir,
        output_dir=output_dir,
        markdown_files_seen=len(markdown_files),
        pdf_pages_written=pdf_count,
        manifest_path=manifest_path,
    )


def build_bundle_plan(
    case_id: str, selection: BundleSelection | None = None
) -> BundlePlanSummary:
    loaded = load_case(case_id)
    document_index_path = loaded.case_dir / _case_path_value(loaded.config, "document_index")
    exhibit_index_path = loaded.case_dir / _case_path_value(loaded.config, "exhibit_index")
    bundle_root = loaded.case_dir / _case_path_value(loaded.config, "bundle_root")
    plan_dir = bundle_root / "plan"
    plan_dir.mkdir(parents=True, exist_ok=True)

    document_rows = _read_csv(document_index_path, INDEX_FIELDS)
    exhibit_rows = _read_csv(exhibit_index_path, EXHIBIT_FIELDS)
    documents_by_id = {
        row.get("document_id", ""): row
        for row in document_rows
        if row.get("document_id")
    }
    selection = selection or load_bundle_selection(case_id)
    selected_exhibits = set(selection.exhibit_numbers)
    selected_documents = set(selection.document_ids)

    items: list[dict[str, str]] = []
    sequence = 1
    ordered_exhibits = sorted(exhibit_rows, key=_exhibit_sort_key)
    for exhibit_position, exhibit in enumerate(ordered_exhibits, start=1):
        exhibit_number = exhibit.get("exhibit_number", "").strip()
        if not exhibit_number or exhibit_number not in selected_exhibits:
            continue
        document_ids = [
            part.strip()
            for part in exhibit.get("document_ids", "").split(";")
            if part.strip() and part.strip() in selected_documents
        ]
        document_groups = _logical_document_groups(document_ids, documents_by_id)
        exhibit_separator = (
            bundle_root
            / "separators"
            / "pdf"
            / f"{exhibit_position:03d}_exhibit_{_safe_filename(exhibit_number)}.pdf"
        )
        items.append(
            _plan_item(
                sequence,
                "exhibit_separator",
                exhibit_number,
                "",
                exhibit_separator,
                loaded.case_dir,
                _pdf_status(exhibit_separator),
                "Generated exhibit separator PDF.",
            )
        )
        sequence += 1

        for episode_position, section in enumerate(
            _numbered_episode_document_sections(exhibit, document_groups, loaded.config),
            start=1,
        ):
            episode_title = str(section["episode_title"])
            episode_number = str(section["episode_number"])
            numbered_groups = section["documents"]
            if episode_title:
                episode_separator = (
                    bundle_root
                    / "separators"
                    / "pdf"
                    / f"{exhibit_position:03d}_{episode_position:03d}_000_episode_{_safe_filename(episode_number)}.pdf"
                )
                items.append(
                    _plan_item(
                        sequence,
                        "episode_separator",
                        exhibit_number,
                        "",
                        episode_separator,
                        loaded.case_dir,
                        _pdf_status(episode_separator),
                        "Generated episode separator PDF.",
                    )
                )
                sequence += 1

            for document_position, (_document_number, document, translations) in enumerate(
                numbered_groups,
                start=1,
            ):
                document_id = document.get("document_id", "")
                document_separator = (
                    bundle_root
                    / "separators"
                    / "pdf"
                    / f"{exhibit_position:03d}_{episode_position:03d}_{document_position:03d}_{_safe_filename(document_id)}.pdf"
                )
                items.append(
                    _plan_item(
                        sequence,
                        "document_separator",
                        exhibit_number,
                        document_id,
                        document_separator,
                        loaded.case_dir,
                        _pdf_status(document_separator),
                        "Generated document separator PDF.",
                    )
                )
                sequence += 1
                for source_document in [document, *translations]:
                    source_id = source_document.get("document_id", "")
                    source_path = loaded.case_dir / source_document.get("file_path", "")
                    status, note = _source_document_status(source_path)
                    items.append(
                        _plan_item(
                            sequence,
                            "source_document",
                            exhibit_number,
                            source_id,
                            source_path,
                            loaded.case_dir,
                            status,
                            note,
                        )
                    )
                    sequence += 1

    plan_csv_path = plan_dir / "bundle_plan.csv"
    plan_md_path = plan_dir / "bundle_plan.md"
    _write_plan_csv(plan_csv_path, items)
    plan_md_path.write_text(_render_plan_markdown(items), encoding="utf-8")
    return BundlePlanSummary(
        plan_csv_path=plan_csv_path,
        plan_md_path=plan_md_path,
        total_items=len(items),
        ready_items=sum(1 for item in items if item["status"] in {"ready_pdf", "convertible_image", "convertible_text"}),
        missing_items=sum(1 for item in items if item["status"] == "missing"),
        unsupported_items=sum(1 for item in items if item["status"] == "unsupported"),
        source_documents=sum(1 for item in items if item["item_type"] == "source_document"),
        separator_pdfs=sum(1 for item in items if item["item_type"].endswith("_separator")),
    )


def build_evidence_bundle(case_id: str) -> BundleBuildSummary:
    loaded = load_case(case_id)
    preparation = inspect_bundle_preparation(case_id)
    if not preparation.ready_to_build:
        raise SystemExit(preparation.reason)
    selection = load_bundle_selection(case_id)
    plan_summary = build_bundle_plan(case_id, selection=selection)
    items = _read_plan_csv(plan_summary.plan_csv_path)
    if not items:
        raise SystemExit("Bundle plan is empty. Run 'Refresh indexes' and generate separator PDFs first.")
    try:
        from pypdf import PdfReader, PdfWriter  # type: ignore
        from reportlab.lib.pagesizes import LETTER  # type: ignore
        from reportlab.lib.units import inch  # type: ignore
        from reportlab.lib.utils import ImageReader  # type: ignore
        from reportlab.pdfgen import canvas  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "Missing PDF dependencies. Install with `pip install .[pdf]` "
            "or run this command with the bundled Codex Python runtime."
        ) from exc

    bundle_root = loaded.case_dir / _case_path_value(loaded.config, "bundle_root")
    converted_dir = bundle_root / "converted"
    final_dir = bundle_root / "final"
    converted_dir.mkdir(parents=True, exist_ok=True)
    final_dir.mkdir(parents=True, exist_ok=True)
    catalog = bundle_catalog(case_id)
    all_document_count = sum(len(exhibit["documents"]) for exhibit in catalog)
    is_full_bundle = (
        len(selection.exhibit_numbers) == len(catalog)
        and len(selection.document_ids) == all_document_count
    )
    final_pdf_path = final_dir / (
        "evidence_bundle.pdf" if is_full_bundle else "evidence_bundle_selected.pdf"
    )

    writer = PdfWriter()
    merged_items = 0
    converted_items = 0
    skipped_items = 0
    for item in items:
        source_path = loaded.case_dir / item["path"]
        pdf_path = source_path
        if item["status"] == "convertible_image":
            pdf_path = converted_dir / f"{int(item['sequence']):04d}_{_safe_filename(item['document_id'] or item['item_type'])}.pdf"
            _image_to_pdf(source_path, pdf_path, canvas, LETTER, inch, ImageReader)
            converted_items += 1
        elif item["status"] == "convertible_text":
            pdf_path = converted_dir / f"{int(item['sequence']):04d}_{_safe_filename(item['document_id'] or item['item_type'])}.pdf"
            _text_source_to_pdf(source_path, pdf_path, _bundle_font_family(loaded.config))
            converted_items += 1
        elif item["status"] != "ready_pdf":
            skipped_items += 1
            continue

        _add_pdf_pages(writer, PdfReader, pdf_path, strip_annotations=True)
        merged_items += 1

    try:
        with final_pdf_path.open("wb") as handle:
            writer.write(handle)
    except PermissionError:
        final_pdf_path = final_pdf_path.with_name(final_pdf_path.stem + "_updated.pdf")
        with final_pdf_path.open("wb") as handle:
            writer.write(handle)
    final_pdf_path = _clean_pdf_if_possible(final_pdf_path)
    return BundleBuildSummary(
        final_pdf_path=final_pdf_path,
        plan_csv_path=plan_summary.plan_csv_path,
        merged_items=merged_items,
        converted_items=converted_items,
        skipped_items=skipped_items,
    )


def build_final_filing_pdf(case_id: str, *, memo_docx_path: str = "") -> FinalFilingSummary:
    """Stage 4: resolve memo PAGE placeholders and merge memo + evidence bundle."""
    loaded = load_case(case_id)
    preparation = inspect_bundle_preparation(case_id)
    if not preparation.ready_to_build:
        raise SystemExit(f"Stage 3 bundle is not ready: {preparation.reason}")

    bundle_summary = build_evidence_bundle(case_id)
    evidence_pdf = bundle_summary.final_pdf_path
    if not evidence_pdf.exists():
        raise SystemExit(f"Evidence bundle PDF not found: {evidence_pdf}")

    final_memo_root = loaded.case_dir / _case_path_value(loaded.config, "final_memo")
    source_memo = Path(memo_docx_path) if memo_docx_path else final_memo_root / "working_memo.docx"
    if not source_memo.is_absolute():
        source_memo = loaded.case_dir / source_memo
    if not source_memo.exists():
        raise SystemExit(f"Memo DOCX not found: {source_memo}")

    bundle_root = loaded.case_dir / _case_path_value(loaded.config, "bundle_root")
    final_dir = bundle_root / "final"
    final_dir.mkdir(parents=True, exist_ok=True)
    numbered_docx = final_memo_root / "working_memo_numbered.docx"
    numbered_pdf = final_memo_root / "working_memo_numbered.pdf"
    final_pdf = final_dir / "final_filing.pdf"
    report_path = final_dir / "final_filing_report.md"

    document_lookup = _bundle_document_citation_index(loaded)
    plan_items = _read_plan_csv(bundle_summary.plan_csv_path)
    evidence_page_count = _pdf_page_count(evidence_pdf)

    memo_pages = 0
    replacement_summary: dict[str, object] = {
        "seen": 0,
        "resolved": 0,
        "unresolved": [],
    }
    iterations = 0
    for iterations in range(1, 5):
        if iterations == 1:
            probe_pdf = final_memo_root / "working_memo_probe.pdf"
            _convert_docx_to_pdf(source_memo, probe_pdf)
            memo_pages = _pdf_page_count(probe_pdf)
        separator_pages = _document_separator_pages_from_plan(loaded, plan_items)
        absolute_pages = {
            document_id: memo_pages + bundle_page
            for document_id, bundle_page in separator_pages.items()
        }
        shutil.copy2(source_memo, numbered_docx)
        replacement_summary = _replace_memo_page_placeholders(
            numbered_docx,
            document_lookup=document_lookup,
            absolute_pages=absolute_pages,
        )
        _convert_docx_to_pdf(numbered_docx, numbered_pdf)
        next_memo_pages = _pdf_page_count(numbered_pdf)
        if next_memo_pages == memo_pages:
            break
        memo_pages = next_memo_pages

    final_pdf = _merge_memo_and_bundle(numbered_pdf, evidence_pdf, final_pdf)
    final_pages = _pdf_page_count(final_pdf)
    unresolved = replacement_summary.get("unresolved", [])
    unresolved_count = len(unresolved) if isinstance(unresolved, list) else 0
    report_path.write_text(
        _render_final_filing_report(
            case_id=case_id,
            source_memo=source_memo,
            numbered_docx=numbered_docx,
            numbered_pdf=numbered_pdf,
            evidence_pdf=evidence_pdf,
            final_pdf=final_pdf,
            memo_pages=memo_pages,
            bundle_pages=evidence_page_count,
            final_pages=final_pages,
            placeholders_seen=int(replacement_summary.get("seen", 0)),
            placeholders_resolved=int(replacement_summary.get("resolved", 0)),
            unresolved=replacement_summary.get("unresolved", []),
            iterations=iterations,
        ),
        encoding="utf-8",
    )
    return FinalFilingSummary(
        final_pdf_path=final_pdf,
        numbered_memo_docx_path=numbered_docx,
        numbered_memo_pdf_path=numbered_pdf,
        evidence_bundle_pdf_path=evidence_pdf,
        report_path=report_path,
        memo_pages=memo_pages,
        bundle_pages=evidence_page_count,
        final_pages=final_pages,
        placeholders_seen=int(replacement_summary.get("seen", 0)),
        placeholders_resolved=int(replacement_summary.get("resolved", 0)),
        placeholders_unresolved=unresolved_count,
        iterations=iterations,
    )


def _document_separator_pages_from_plan(
    loaded: LoadedCase, items: list[dict[str, str]]
) -> dict[str, int]:
    pages: dict[str, int] = {}
    bundle_page = 1
    for item in items:
        path_value = item.get("path", "").strip()
        if not path_value:
            continue
        pdf_path = loaded.case_dir / path_value
        if item.get("item_type") == "document_separator" and item.get("document_id"):
            pages.setdefault(item["document_id"], bundle_page)
        if pdf_path.exists() and item.get("status") in {"ready_pdf", "convertible_image", "convertible_text"}:
            try:
                bundle_page += _pdf_page_count(pdf_path)
            except Exception:
                bundle_page += 1
    return pages


def _bundle_document_citation_index(
    loaded: LoadedCase,
) -> dict[str, dict[object, object]]:
    document_index_path = loaded.case_dir / _case_path_value(loaded.config, "document_index")
    exhibit_index_path = loaded.case_dir / _case_path_value(loaded.config, "exhibit_index")
    document_rows = _read_csv(document_index_path, INDEX_FIELDS)
    exhibit_rows = _read_csv(exhibit_index_path, EXHIBIT_FIELDS)
    documents_by_id = {
        row.get("document_id", ""): row
        for row in document_rows
        if row.get("document_id")
    }
    by_number: dict[tuple[str, str], str] = {}
    by_title: dict[tuple[str, str], list[str]] = {}
    titles_by_id: dict[str, str] = {}
    for exhibit in sorted(exhibit_rows, key=_exhibit_sort_key):
        exhibit_number = exhibit.get("exhibit_number", "").strip()
        document_ids = [
            part.strip()
            for part in exhibit.get("document_ids", "").split(";")
            if part.strip()
        ]
        document_groups = _logical_document_groups(document_ids, documents_by_id)
        for section in _numbered_episode_document_sections(exhibit, document_groups, loaded.config):
            for document_number, document, translations in section["documents"]:
                document_id = document.get("document_id", "")
                if not document_id:
                    continue
                by_number[(exhibit_number, str(document_number).rstrip("."))] = document_id
                candidate_titles = [
                    _document_title(document),
                    document.get("display_title", ""),
                    document.get("original_file_name", ""),
                ]
                if translations:
                    candidate_titles.append(f"{_document_title(document)}; English translation")
                titles_by_id[document_id] = _document_title(document)
                for title in candidate_titles:
                    normalized = _normalize_citation_text(title)
                    if normalized:
                        by_title.setdefault((exhibit_number, normalized), []).append(document_id)
    return {
        "by_number": by_number,
        "by_title": by_title,
        "titles_by_id": titles_by_id,
    }


def _replace_memo_page_placeholders(
    docx_path: Path,
    *,
    document_lookup: dict[str, dict[object, object]],
    absolute_pages: dict[str, int],
) -> dict[str, object]:
    try:
        from docx import Document  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit("python-docx is required to update memo page placeholders.") from exc

    document = Document(docx_path)
    seen = 0
    resolved = 0
    unresolved: list[str] = []
    for paragraph in _iter_docx_paragraphs(document):
        text = paragraph.text
        if "PAGE" not in text:
            continue
        replacements: list[str | None] = []
        for match in re.finditer(r"\bPAGE\b", text):
            seen += 1
            document_id = _resolve_page_placeholder_document_id(
                text,
                match.start(),
                document_lookup=document_lookup,
            )
            if document_id and document_id in absolute_pages:
                replacements.append(str(absolute_pages[document_id]))
                resolved += 1
            else:
                replacements.append(None)
                unresolved.append(_placeholder_context(text, match.start()))
        _replace_page_tokens_in_paragraph(paragraph, replacements)
    document.save(docx_path)
    return {"seen": seen, "resolved": resolved, "unresolved": unresolved}


def _resolve_page_placeholder_document_id(
    text: str,
    page_offset: int,
    *,
    document_lookup: dict[str, dict[object, object]],
) -> str:
    start = text.rfind("Exhibit", 0, page_offset)
    if start < 0:
        start = max(0, page_offset - 240)
    end = text.find(")", page_offset)
    if end < 0:
        end = min(len(text), page_offset + 520)
    fragment = text[start:end]
    exhibit_match = re.search(r"Exhibit\s+([0-9]+(?:-[0-9]+)?)", fragment, re.IGNORECASE)
    if not exhibit_match:
        return ""
    exhibit_number = exhibit_match.group(1)
    after_page = text[page_offset + len("PAGE") : end]
    next_reference = re.search(r";\s*Exhibit\s+", after_page, re.IGNORECASE)
    if next_reference:
        after_page = after_page[: next_reference.start()]
    desc = after_page
    if ":" in desc:
        desc = desc.split(":", 1)[1]
    desc = desc.strip(" .;:")
    number_match = re.match(
        r"([0-9]+(?:-[0-9]+)?(?:\.\d+)+)\.?\s*(?:[-–—]\s*)?(.*)",
        desc,
    )
    by_number = document_lookup.get("by_number", {})
    if number_match:
        item_number = number_match.group(1).rstrip(".")
        document_id = by_number.get((exhibit_number, item_number))
        if isinstance(document_id, str) and document_id:
            return document_id
        desc = number_match.group(2).strip(" .;:")

    normalized_desc = _normalize_citation_text(desc)
    if not normalized_desc:
        return ""
    by_title = document_lookup.get("by_title", {})
    exact = by_title.get((exhibit_number, normalized_desc))
    if isinstance(exact, list) and len(set(exact)) == 1:
        return exact[0]
    candidates: list[str] = []
    for key, values in by_title.items():
        if not isinstance(key, tuple) or len(key) != 2 or key[0] != exhibit_number:
            continue
        title = str(key[1])
        if title and (title in normalized_desc or normalized_desc in title):
            if isinstance(values, list):
                candidates.extend(values)
    unique = sorted(set(candidates))
    if len(unique) == 1:
        return unique[0]

    scored: list[tuple[float, str]] = []
    desc_tokens = set(normalized_desc.split())
    if not desc_tokens:
        return ""
    for key, values in by_title.items():
        if not isinstance(key, tuple) or len(key) != 2 or key[0] != exhibit_number:
            continue
        title = str(key[1])
        title_tokens = set(title.split())
        if not title_tokens:
            continue
        overlap = len(desc_tokens & title_tokens)
        title_coverage = overlap / max(1, len(title_tokens))
        desc_coverage = overlap / max(1, len(desc_tokens))
        score = (0.75 * title_coverage) + (0.25 * desc_coverage)
        if score >= 0.62 and isinstance(values, list):
            for document_id in values:
                scored.append((score, document_id))
    if not scored:
        return ""
    scored.sort(reverse=True)
    best_score = scored[0][0]
    best_ids = sorted({document_id for score, document_id in scored if score == best_score})
    return best_ids[0] if len(best_ids) == 1 else ""


def _replace_page_tokens_in_paragraph(paragraph: object, replacements: list[str | None]) -> None:
    if not replacements:
        return
    remaining = list(replacements)
    run_page_tokens = sum(str(run.text).count("PAGE") for run in paragraph.runs)
    if run_page_tokens == len(replacements):
        for run in paragraph.runs:
            text = str(run.text)
            while "PAGE" in text and remaining:
                replacement = remaining.pop(0)
                text = text.replace("PAGE", replacement if replacement is not None else "PAGE", 1)
            run.text = text
        return

    text = paragraph.text
    for replacement in replacements:
        text = re.sub(r"\bPAGE\b", replacement if replacement is not None else "PAGE", text, count=1)
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text


def _iter_docx_paragraphs(document: object) -> list[object]:
    paragraphs: list[object] = []
    paragraphs.extend(document.paragraphs)
    for table in document.tables:
        paragraphs.extend(_iter_table_paragraphs(table))
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        for table in section.header.tables:
            paragraphs.extend(_iter_table_paragraphs(table))
        for table in section.footer.tables:
            paragraphs.extend(_iter_table_paragraphs(table))
    return paragraphs


def _iter_table_paragraphs(table: object) -> list[object]:
    paragraphs: list[object] = []
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
            for nested_table in cell.tables:
                paragraphs.extend(_iter_table_paragraphs(nested_table))
    return paragraphs


def _placeholder_context(text: str, offset: int) -> str:
    start = max(0, offset - 120)
    end = min(len(text), offset + 260)
    return " ".join(text[start:end].split())


def _normalize_citation_text(value: str) -> str:
    value = value.casefold()
    value = re.sub(r"\boriginal and english translation\b", " ", value)
    value = re.sub(r"\benglish translation\b", " ", value)
    value = re.sub(r"\boriginal\b", " ", value)
    value = re.sub(r"^\s*[0-9]+(?:-[0-9]+)?(?:\.\d+)*\.?\s*[-–—]?\s*", " ", value)
    return " ".join(
        "".join(character if character.isalnum() else " " for character in value).split()
    )


def _convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        try:
            pdf_path.unlink()
        except PermissionError:
            pdf_path = pdf_path.with_name(pdf_path.stem + "_updated.pdf")
    soffice = _find_soffice()
    if soffice:
        output_dir = pdf_path.parent
        subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(docx_path),
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        converted = output_dir / (docx_path.stem + ".pdf")
        if converted != pdf_path:
            if pdf_path.exists():
                pdf_path.unlink()
            converted.replace(pdf_path)
        if not pdf_path.exists():
            raise SystemExit(f"DOCX to PDF conversion did not create: {pdf_path}")
        return
    _convert_docx_to_pdf_with_word(docx_path, pdf_path)


def _convert_docx_to_pdf_with_word(docx_path: Path, pdf_path: Path) -> None:
    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit(
            "DOCX to PDF conversion requires LibreOffice (`soffice`) or Microsoft Word "
            "automation (`pip install pywin32`)."
        ) from exc
    word = None
    opened = None
    pythoncom.CoInitialize()
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        opened = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
        opened.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
    finally:
        if opened is not None:
            opened.Close(False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()
    if not pdf_path.exists():
        raise SystemExit(f"Microsoft Word did not create PDF: {pdf_path}")


def _find_soffice() -> str:
    on_path = shutil.which("soffice")
    if on_path:
        return on_path
    for candidate in (
        "C:/Program Files/LibreOffice/program/soffice.exe",
        "C:/Program Files (x86)/LibreOffice/program/soffice.exe",
    ):
        if Path(candidate).exists():
            return candidate
    return ""


def _pdf_page_count(path: Path) -> int:
    try:
        from pypdf import PdfReader  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit("pypdf is required to count PDF pages.") from exc
    return len(PdfReader(str(path)).pages)


def _merge_memo_and_bundle(memo_pdf: Path, evidence_pdf: Path, final_pdf: Path) -> Path:
    try:
        from pypdf import PdfReader, PdfWriter  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit("pypdf is required to merge the final filing PDF.") from exc
    writer = PdfWriter()
    for source in (memo_pdf, evidence_pdf):
        _add_pdf_pages(writer, PdfReader, source, strip_annotations=True)
    try:
        with final_pdf.open("wb") as handle:
            writer.write(handle)
    except PermissionError:
        final_pdf = final_pdf.with_name(final_pdf.stem + "_updated.pdf")
        with final_pdf.open("wb") as handle:
            writer.write(handle)
    return _clean_pdf_if_possible(final_pdf)


def _add_pdf_pages(
    writer: object,
    PdfReader: object,
    pdf_path: Path,
    *,
    strip_annotations: bool,
) -> int:
    from pypdf.generic import NameObject  # type: ignore

    reader = PdfReader(str(pdf_path))  # type: ignore[operator]
    pages_added = 0
    for page in reader.pages:  # type: ignore[attr-defined]
        if strip_annotations and "/Annots" in page:
            page.pop(NameObject("/Annots"), None)
        writer.add_page(page)  # type: ignore[attr-defined]
        pages_added += 1
    return pages_added


def _clean_pdf_if_possible(path: Path) -> Path:
    try:
        import pikepdf  # type: ignore
    except ModuleNotFoundError:
        return path
    temp_path = path.with_name(path.stem + ".cleaning.pdf")
    try:
        with pikepdf.open(path) as pdf:
            pdf.save(
                temp_path,
                linearize=True,
                compress_streams=True,
                object_stream_mode=pikepdf.ObjectStreamMode.generate,
            )
        temp_path.replace(path)
    except Exception as exc:  # noqa: BLE001
        logging.warning("Could not clean PDF %s with pikepdf: %s", path, exc)
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass
    return path


def _render_final_filing_report(
    *,
    case_id: str,
    source_memo: Path,
    numbered_docx: Path,
    numbered_pdf: Path,
    evidence_pdf: Path,
    final_pdf: Path,
    memo_pages: int,
    bundle_pages: int,
    final_pages: int,
    placeholders_seen: int,
    placeholders_resolved: int,
    unresolved: object,
    iterations: int,
) -> str:
    unresolved_items = unresolved if isinstance(unresolved, list) else []
    lines = [
        "# Final filing build report",
        "",
        f"- case_id: `{case_id}`",
        f"- source memo: `{source_memo}`",
        f"- numbered memo DOCX: `{numbered_docx}`",
        f"- numbered memo PDF: `{numbered_pdf}`",
        f"- evidence bundle PDF: `{evidence_pdf}`",
        f"- final filing PDF: `{final_pdf}`",
        f"- memo pages: {memo_pages}",
        f"- evidence bundle pages: {bundle_pages}",
        f"- final pages: {final_pages}",
        f"- PAGE placeholders seen: {placeholders_seen}",
        f"- PAGE placeholders resolved: {placeholders_resolved}",
        f"- PAGE placeholders unresolved: {len(unresolved_items)}",
        f"- page-count iterations: {iterations}",
        "",
    ]
    if unresolved_items:
        lines.extend(["## Unresolved PAGE placeholders", ""])
        for item in unresolved_items[:100]:
            lines.append(f"- {item}")
        if len(unresolved_items) > 100:
            lines.append(f"- ... {len(unresolved_items) - 100} more")
        lines.append("")
    return "\n".join(lines)


def order_documents_original_then_translation(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    by_id = {row.get("document_id", ""): row for row in rows if row.get("document_id")}
    numbered_context = _numbered_document_sort_context(rows)
    translations_by_parent: dict[str, list[dict[str, str]]] = {}
    originals: list[dict[str, str]] = []
    unpaired_translations: list[dict[str, str]] = []

    for row in rows:
        parent = row.get("parent_document_id", "").strip()
        is_translation = row.get("relationship_type") == "translation" or row.get("translation_status") == "translation"
        if is_translation and parent:
            translations_by_parent.setdefault(parent, []).append(row)
        elif is_translation:
            unpaired_translations.append(row)
        else:
            originals.append(row)

    ordered: list[dict[str, str]] = []
    for original in sorted(originals, key=lambda row: _document_sort_key(row, numbered_context)):
        ordered.append(original)
        linked = translations_by_parent.pop(original.get("document_id", ""), [])
        ordered.extend(sorted(linked, key=lambda row: _document_sort_key(row, numbered_context)))

    for parent_id in sorted(translations_by_parent, key=_natural_sort_key):
        parent = by_id.get(parent_id)
        if parent and parent not in ordered:
            ordered.append(parent)
        ordered.extend(
            sorted(
                translations_by_parent[parent_id],
                key=lambda row: _document_sort_key(row, numbered_context),
            )
        )

    ordered.extend(sorted(unpaired_translations, key=lambda row: _document_sort_key(row, numbered_context)))
    return ordered


def _first_document_order(rows: list[dict[str, str]]) -> int:
    orders: list[int] = []
    for row in rows:
        order = row.get("final_bundle_order", "").strip()
        if order.isdigit():
            orders.append(int(order))
    return min(orders, default=10**9)


def _bundle_selection_path(loaded: LoadedCase) -> Path:
    return (
        loaded.case_dir
        / _case_path_value(loaded.config, "bundle_root")
        / "selection.json"
    )


def _bundle_preparation_path(loaded: LoadedCase) -> Path:
    return (
        loaded.case_dir
        / _case_path_value(loaded.config, "bundle_root")
        / "preparation.json"
    )


def _invalidate_bundle_preparation(loaded: LoadedCase) -> None:
    path = _bundle_preparation_path(loaded)
    if path.exists():
        path.unlink()


def _validated_bundle_selection(
    case_id: str,
    selected_exhibits: list[str],
    selected_document_ids: list[str],
) -> BundleSelection:
    catalog = bundle_catalog(case_id)
    requested_exhibits = {str(value).strip() for value in selected_exhibits if str(value).strip()}
    requested_documents = {
        str(value).strip() for value in selected_document_ids if str(value).strip()
    }
    catalog_exhibits = {str(item["exhibit_number"]) for item in catalog}
    unknown_exhibits = requested_exhibits - catalog_exhibits
    if unknown_exhibits:
        raise ValueError(f"Unknown exhibit selection: {', '.join(sorted(unknown_exhibits))}.")

    exhibits: list[str] = []
    documents: list[str] = []
    known_documents: set[str] = set()
    for exhibit in catalog:
        exhibit_number = str(exhibit["exhibit_number"])
        exhibit_document_ids = [
            str(document.get("document_id", ""))
            for document in exhibit["documents"]
            if document.get("document_id", "")
        ]
        known_documents.update(exhibit_document_ids)
        selected_here = [
            document_id
            for document_id in exhibit_document_ids
            if exhibit_number in requested_exhibits and document_id in requested_documents
        ]
        if selected_here:
            exhibits.append(exhibit_number)
            documents.extend(selected_here)

    unknown_documents = requested_documents - known_documents
    if unknown_documents:
        raise ValueError(
            f"Unknown document selection: {', '.join(sorted(unknown_documents))}."
        )
    if not documents:
        raise ValueError("Select at least one document before preparing the bundle.")
    return BundleSelection(tuple(exhibits), tuple(documents))


def _save_bundle_selection(loaded: LoadedCase, selection: BundleSelection) -> None:
    path = _bundle_selection_path(loaded)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            {
                "selected_exhibits": list(selection.exhibit_numbers),
                "selected_document_ids": list(selection.document_ids),
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    _invalidate_bundle_preparation(loaded)


def _selection_digest(selection: BundleSelection) -> str:
    payload = json.dumps(
        {
            "exhibits": list(selection.exhibit_numbers),
            "documents": list(selection.document_ids),
        },
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _bundle_index_digest(loaded: LoadedCase) -> str:
    digest = hashlib.sha256()
    for key in ("document_index", "exhibit_index"):
        path = loaded.case_dir / _case_path_value(loaded.config, key)
        digest.update(key.encode("utf-8"))
        digest.update(path.read_bytes() if path.exists() else b"[missing]")
    return digest.hexdigest()


def _read_csv(path: Path, fields: list[str]) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = []
        for row in csv.DictReader(handle):
            normalized = {field: "" for field in fields}
            normalized.update({key: value or "" for key, value in row.items() if key})
            rows.append(normalized)
        return rows


def _write_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def _blank_exhibit_row() -> dict[str, str]:
    return {field: "" for field in EXHIBIT_FIELDS}


NumberedSortContext = dict[str, tuple[tuple[int, object], tuple[int, object]]]


def _assignment_sort_key(
    item: tuple[str, str],
    plan: EvidenceLayoutPlan,
    numbered_context: NumberedSortContext,
) -> tuple[object, ...]:
    document_id, exhibit_number = item
    episode_order, document_order = numbered_context.get(
        document_id,
        ((1, plan.document_order.get(document_id, 10**9)), (1, plan.document_order.get(document_id, 10**9))),
    )
    return (
        plan.exhibit_order.get(exhibit_number, 10**9),
        episode_order,
        document_order,
        plan.document_order.get(document_id, 10**9),
        _natural_sort_key(exhibit_number),
        _natural_sort_key(document_id),
    )


def _numbered_document_sort_context(
    rows: object,
    *,
    fallback_order: dict[str, int] | None = None,
    document_to_exhibit: dict[str, str] | None = None,
) -> NumberedSortContext:
    fallback_order = fallback_order or {}
    document_to_exhibit = document_to_exhibit or {}
    row_list = [row for row in rows if isinstance(row, dict) and row.get("document_id")]
    groups: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in row_list:
        document_id = row.get("document_id", "")
        exhibit_number = document_to_exhibit.get(document_id) or row.get("exhibit_number", "").strip()
        episode_title = _document_episode_title(row, {})
        groups.setdefault((exhibit_number, episode_title), []).append(row)

    context: NumberedSortContext = {}
    for (_exhibit_number, episode_title), group_rows in groups.items():
        episode_prefix = _manual_order_prefix(episode_title)
        first_fallback = min(
            (fallback_order.get(row.get("document_id", ""), _base_document_order(row)) for row in group_rows),
            default=10**9,
        )
        episode_order: tuple[int, tuple[int, ...] | int] = (
            (0, episode_prefix) if episode_prefix is not None else (1, first_fallback)
        )
        document_path_orders = {
            row.get("document_id", ""): _manual_path_order_key(row)
            for row in group_rows
        }
        has_numbered_documents = any(order is not None for order in document_path_orders.values())
        for row in group_rows:
            document_id = row.get("document_id", "")
            path_order = document_path_orders.get(document_id)
            fallback = fallback_order.get(document_id, _base_document_order(row))
            document_order: tuple[int, object]
            if has_numbered_documents and path_order is not None:
                document_order = (0, path_order)
            elif has_numbered_documents:
                document_order = (1, fallback)
            else:
                document_order = (1, fallback)
            context[document_id] = (episode_order, document_order)
    return context


def _manual_path_order_key(row: dict[str, str]) -> tuple[tuple[int, object], ...] | None:
    raw_path = _normalize_slashes(row.get("file_path", "").strip())
    if not raw_path:
        return None
    parts = PurePosixPath(raw_path).parts
    evidence_root_index = -1
    for marker in ("originals", "translations"):
        if marker in parts:
            evidence_root_index = parts.index(marker)
            break
    if evidence_root_index < 0:
        return None
    relative_parts = parts[evidence_root_index + 1 :]
    if len(relative_parts) < 3:
        return None
    sortable_parts: list[tuple[int, object]] = []
    has_numbered_component = False
    for part in relative_parts[2:]:
        prefix = _manual_order_prefix(part)
        if prefix is not None:
            sortable_parts.append((0, prefix))
            has_numbered_component = True
        else:
            sortable_parts.append((1, tuple(_natural_sort_key(part))))
    if not has_numbered_component:
        return None
    return tuple(sortable_parts)


def _manual_order_prefix(value: str) -> tuple[int, ...] | None:
    match = re.match(
        r"^\s*(\d{1,3}(?:\.\d{1,3})*)(?:\.\s+|\.(?=[^\d\s])|[\)_-]\s*)",
        value,
    )
    if not match:
        return None
    try:
        parts = tuple(int(part) for part in match.group(1).split("."))
    except ValueError:
        return None
    return parts or None


def _base_document_order(row: dict[str, str]) -> int:
    order = row.get("final_bundle_order", "").strip()
    if order.isdigit():
        return int(order)
    return 10**9


def _document_sort_key(
    row: dict[str, str], numbered_context: NumberedSortContext | None = None
) -> tuple[object, ...]:
    context = (numbered_context or {}).get(row.get("document_id", ""))
    if context:
        episode_order, document_order = context
        return (
            0,
            episode_order,
            document_order,
            _base_document_order(row),
            _natural_sort_key(row.get("file_path", "") or row.get("document_id", "")),
            row.get("document_id", ""),
        )
    return _base_document_sort_key(row)


def _base_document_sort_key(row: dict[str, str]) -> tuple[int, list[object], str]:
    order = row.get("final_bundle_order", "").strip()
    if order.isdigit():
        return (0, [int(order)], row.get("document_id", ""))
    return (1, _natural_sort_key(row.get("file_path", "") or row.get("document_id", "")), row.get("document_id", ""))


def _natural_sort_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value or "")]


def _next_exhibit_number(rows: list[dict[str, str]]) -> int:
    numbers = []
    for row in rows:
        exhibit_id = row.get("exhibit_id", "")
        digits = "".join(char for char in exhibit_id if char.isdigit())
        if digits:
            numbers.append(int(digits))
    return max(numbers, default=0) + 1


def _common_value(rows: list[dict[str, str]], key: str) -> str:
    values = sorted({row.get(key, "").strip() for row in rows if row.get(key, "").strip()})
    if len(values) == 1:
        return values[0]
    return ""


def _truthy(value: str) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y", "locked"}


def _merge_note(existing: str, note: str) -> str:
    existing = (existing or "").strip()
    note = note.strip()
    if not note:
        return existing
    if not existing:
        return note
    if note in existing:
        return existing
    return existing + " | " + note


def _logical_document_groups(
    document_ids: list[str], documents_by_id: dict[str, dict[str, str]]
) -> list[tuple[dict[str, str], list[dict[str, str]]]]:
    selected_rows = [
        documents_by_id[document_id]
        for document_id in document_ids
        if document_id in documents_by_id
    ]
    selected_ids = {row.get("document_id", "") for row in selected_rows}
    translations_by_parent: dict[str, list[dict[str, str]]] = {}
    for row in selected_rows:
        parent_id = row.get("parent_document_id", "").strip()
        if _is_translation(row) and parent_id in selected_ids:
            translations_by_parent.setdefault(parent_id, []).append(row)

    groups: list[tuple[dict[str, str], list[dict[str, str]]]] = []
    consumed: set[str] = set()
    for row in selected_rows:
        document_id = row.get("document_id", "")
        if document_id in consumed:
            continue
        parent_id = row.get("parent_document_id", "").strip()
        if _is_translation(row) and parent_id in selected_ids:
            continue
        translations = translations_by_parent.get(document_id, [])
        groups.append((row, translations))
        consumed.add(document_id)
        consumed.update(item.get("document_id", "") for item in translations)
    return groups


def _is_translation(document: dict[str, str]) -> bool:
    return (
        document.get("relationship_type", "").strip().lower() == "translation"
        or document.get("translation_status", "").strip().lower() == "translation"
    )


def _exhibit_heading(exhibit: dict[str, str]) -> str:
    exhibit_number = exhibit.get("exhibit_number", "").strip()
    title = exhibit.get("display_title", "").strip()
    generic_titles = {"", f"Exhibit {exhibit_number}".casefold()}
    if title.casefold() in generic_titles:
        return f"Exhibit {exhibit_number}"
    return f"Exhibit {exhibit_number}: {title}"


def _render_exhibit_separator(
    exhibit: dict[str, str],
    document_groups: list[tuple[dict[str, str], list[dict[str, str]]]],
    config: dict[str, object] | None = None,
) -> str:
    lines = [
        f"# {_exhibit_heading(exhibit)}",
        "",
    ]
    thesis = exhibit.get("evidentiary_thesis", "").strip()
    if thesis:
        lines.extend(["## Evidentiary thesis", "", thesis, ""])
    lines.extend(["## Documents included", ""])
    if document_groups:
        for section in _numbered_episode_document_sections(exhibit, document_groups, config or {}):
            episode_title = str(section["episode_title"])
            episode_number = str(section["episode_number"])
            numbered_groups = section["documents"]
            if episode_title:
                lines.extend([f"### {episode_number}. {episode_title}", ""])
            for document_number, document, translations in numbered_groups:
                translation_label = "; English translation" if translations else ""
                lines.append(f"{document_number}. {_document_title(document)}{translation_label}")
            lines.append("")
    else:
        lines.append("[No documents selected for this exhibit.]")
        lines.append("")
    lines.extend(_please_see_next_page_lines())
    return "\n".join(lines) + "\n"


def _render_episode_separator(
    exhibit: dict[str, str],
    episode_number: str,
    episode_title: str,
    numbered_groups: list[tuple[str, dict[str, str], list[dict[str, str]]]],
) -> str:
    lines = [
        f"**{_exhibit_heading(exhibit)}**",
        "",
        f"# {episode_number}. {episode_title}",
        "",
        "## Documents included",
        "",
    ]
    if numbered_groups:
        for document_number, document, translations in numbered_groups:
            translation_label = "; English translation" if translations else ""
            lines.append(f"{document_number}. {_document_title(document)}{translation_label}")
    else:
        lines.append("[No documents selected for this episode.]")
    lines.extend(_please_see_next_page_lines())
    return "\n".join(lines) + "\n"


def _episode_document_sections(
    document_groups: list[tuple[dict[str, str], list[dict[str, str]]]],
    config: dict[str, object],
) -> list[tuple[str, list[tuple[dict[str, str], list[dict[str, str]]]]]]:
    sections: list[tuple[str, list[tuple[dict[str, str], list[dict[str, str]]]]]] = []
    section_index: dict[str, int] = {}
    for document, translations in document_groups:
        episode_title = _document_episode_title(document, config)
        if episode_title not in section_index:
            section_index[episode_title] = len(sections)
            sections.append((episode_title, []))
        sections[section_index[episode_title]][1].append((document, translations))
    return sections


def _numbered_episode_document_sections(
    exhibit: dict[str, str],
    document_groups: list[tuple[dict[str, str], list[dict[str, str]]]],
    config: dict[str, object],
) -> list[dict[str, object]]:
    exhibit_number = exhibit.get("exhibit_number", "").strip() or "1"
    sections: list[dict[str, object]] = []
    episode_counter = 1
    direct_counter = 1
    for raw_episode_title, groups in _episode_document_sections(document_groups, config):
        overridden_episode_title = _episode_title_override(config, raw_episode_title)
        episode_title = (
            overridden_episode_title
            if overridden_episode_title != raw_episode_title
            else _display_episode_title(raw_episode_title)
        )
        numbered_groups: list[tuple[str, dict[str, str], list[dict[str, str]]]] = []
        if episode_title:
            episode_number = f"{exhibit_number}.{episode_counter}"
            for document_counter, (document, translations) in enumerate(groups, start=1):
                numbered_groups.append((f"{episode_number}.{document_counter}", document, translations))
            episode_counter += 1
        else:
            episode_number = ""
            for document, translations in groups:
                numbered_groups.append((f"{exhibit_number}.{direct_counter}", document, translations))
                direct_counter += 1
        sections.append(
            {
                "episode_title": episode_title,
                "episode_number": episode_number,
                "documents": numbered_groups,
            }
        )
    return sections


def _document_episode_title(document: dict[str, str], config: dict[str, object]) -> str:
    raw_path = _normalize_slashes(document.get("file_path", "").strip())
    if not raw_path:
        return ""
    parts = PurePosixPath(raw_path).parts
    evidence_root_index = -1
    for marker in ("originals", "translations"):
        if marker in parts:
            evidence_root_index = parts.index(marker)
            break
    if evidence_root_index < 0:
        return ""
    relative_parts = parts[evidence_root_index + 1 :]
    if len(relative_parts) < 3:
        return ""
    return relative_parts[1].strip()


def _display_episode_title(value: str) -> str:
    return re.sub(r"^\s*\d+(?:\.\d+)*[\).\s-]+", "", value).strip() or value.strip()


def _episode_title_override(config: dict[str, object], raw_title: str) -> str:
    title = raw_title.strip()
    overrides = config.get("episode_title_overrides", {})
    if not isinstance(overrides, dict) or not title:
        return title
    exact = str(overrides.get(title, "")).strip()
    if exact:
        return exact
    normalized_title = _normalize_episode_override_key(title)
    for key, value in overrides.items():
        if _normalize_episode_override_key(str(key)) == normalized_title:
            replacement = str(value).strip()
            if replacement:
                return replacement
    return title


def _normalize_episode_override_key(value: str) -> str:
    value = re.sub(r"^\s*\d+(?:\.\d+)*[\).\s-]+", "", value).strip()
    return " ".join(
        "".join(character.casefold() if character.isalnum() else " " for character in value).split()
    )


def _render_document_separator(
    exhibit: dict[str, str],
    episode_number: str,
    episode_title: str,
    document_number: str,
    document: dict[str, str],
    translations: list[dict[str, str]],
) -> str:
    title = _document_title(document)
    translation_label = "; English translation" if translations else ""
    lines = [
        f"**{_exhibit_heading(exhibit)}**",
        "",
    ]
    if episode_title:
        lines.extend([f"**{episode_number}. {episode_title}**", ""])
    lines.extend(
        [
            f"# {document_number}. {title}{translation_label}",
            "",
            *_please_see_next_page_lines(),
        ]
    )
    return "\n".join(lines) + "\n"


def _please_see_next_page_lines() -> list[str]:
    return [
        "",
        "*Please see the next page*",
        "",
    ]


def _document_title(document: dict[str, str]) -> str:
    return (
        document.get("display_title", "").strip()
        or document.get("original_file_name", "").strip()
        or document.get("document_id", "").strip()
        or "[untitled document]"
    )


def _exhibit_sort_key(row: dict[str, str]) -> tuple[int, list[object], str]:
    order = row.get("final_bundle_order", "").strip()
    if order.isdigit():
        return (0, [int(order)], row.get("exhibit_number", ""))
    return (1, _natural_sort_key(row.get("exhibit_number", "")), row.get("exhibit_id", ""))


def _safe_filename(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    safe = safe.strip("._-")
    return safe or "untitled"


def _plan_item(
    sequence: int,
    item_type: str,
    exhibit_number: str,
    document_id: str,
    path: Path,
    case_dir: Path,
    status: str,
    notes: str,
) -> dict[str, str]:
    try:
        relative_path = _normalize_slashes(path.relative_to(case_dir).as_posix())
    except ValueError:
        relative_path = _normalize_slashes(str(path))
    return {
        "sequence": str(sequence),
        "item_type": item_type,
        "exhibit_number": exhibit_number,
        "document_id": document_id,
        "path": relative_path,
        "status": status,
        "notes": notes,
    }


def _pdf_status(path: Path) -> str:
    return "ready_pdf" if path.exists() else "missing"


def _source_document_status(path: Path) -> tuple[str, str]:
    if not path.exists():
        return "missing", "Source file missing."
    suffix = path.suffix.lower()
    if suffix in PDF_SOURCE_EXTENSIONS:
        return "ready_pdf", "Source PDF."
    if suffix in IMAGE_SOURCE_EXTENSIONS:
        return "convertible_image", "Image will be converted to PDF."
    if suffix in TEXT_SOURCE_EXTENSIONS | DOCX_SOURCE_EXTENSIONS:
        return "convertible_text", "Text/DOCX source will be rendered to PDF."
    return "unsupported", f"Unsupported source type for PDF bundle: {suffix or '[no extension]'}"


PLAN_FIELDS = ["sequence", "item_type", "exhibit_number", "document_id", "path", "status", "notes"]


def _write_plan_csv(path: Path, items: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=PLAN_FIELDS)
        writer.writeheader()
        writer.writerows(items)


def _read_plan_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _render_plan_markdown(items: list[dict[str, str]]) -> str:
    lines = [
        "# Evidence bundle plan",
        "",
        "| # | Type | Exhibit | Document | Status | Path | Notes |",
        "|---:|---|---|---|---|---|---|",
    ]
    for item in items:
        lines.append(
            "| "
            + " | ".join(
                [
                    item["sequence"],
                    item["item_type"],
                    item["exhibit_number"],
                    item["document_id"],
                    item["status"],
                    "`" + item["path"] + "`",
                    item["notes"].replace("|", "\\|"),
                ]
            )
            + " |"
        )
    lines.append("")
    return "\n".join(lines)


def _image_to_pdf(
    image_path: Path,
    pdf_path: Path,
    canvas: object,
    pagesize: tuple[float, float],
    inch: float,
    ImageReader: object,
) -> None:
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    page_width, page_height = pagesize
    margin = 0.6 * inch
    c = canvas.Canvas(str(pdf_path), pagesize=pagesize)  # type: ignore[attr-defined]
    image = ImageReader(str(image_path))  # type: ignore[operator]
    image_width, image_height = image.getSize()
    max_width = page_width - 2 * margin
    max_height = page_height - 2 * margin
    scale = min(max_width / image_width, max_height / image_height)
    draw_width = image_width * scale
    draw_height = image_height * scale
    x = (page_width - draw_width) / 2
    y = (page_height - draw_height) / 2
    c.drawImage(image, x, y, width=draw_width, height=draw_height, preserveAspectRatio=True)
    c.showPage()
    c.save()


def _text_source_to_pdf(source_path: Path, pdf_path: Path, font_family: str = "Times New Roman") -> None:
    try:
        from reportlab.lib.pagesizes import LETTER  # type: ignore
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore
        from reportlab.lib.units import inch  # type: ignore
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # type: ignore
        from reportlab.pdfbase import pdfmetrics  # type: ignore
        from reportlab.pdfbase.ttfonts import TTFont  # type: ignore
    except ModuleNotFoundError as exc:
        raise SystemExit("Missing reportlab for text conversion.") from exc

    text = extract_docx_text(source_path) if source_path.suffix.lower() in DOCX_SOURCE_EXTENSIONS else source_path.read_text(encoding="utf-8-sig", errors="replace")
    font_name = _register_pdf_font(pdfmetrics, TTFont, font_family)
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="SourceText",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=10,
            leading=13,
        )
    )
    story: list[object] = [Paragraph(escape(source_path.name), styles["Title"]), Spacer(1, 0.2 * inch)]
    for paragraph in text.splitlines():
        if paragraph.strip():
            story.append(Paragraph(escape(paragraph.strip()), styles["SourceText"]))
            story.append(Spacer(1, 0.05 * inch))
    if len(story) == 2:
        story.append(Paragraph("[empty text source]", styles["SourceText"]))
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(pdf_path),
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title=source_path.name,
    )
    document.build(story)


def _bundle_font_family(config: dict[str, object]) -> str:
    value = str(config.get("bundle_font_family", "")).strip()
    if not value:
        value = str(config.get("font_family", "")).strip()
    return value or "Times New Roman"


def _register_pdf_font(pdfmetrics: object, TTFont: object, preferred_family: str = "") -> str:
    candidates = _preferred_font_candidates(preferred_family or "Times New Roman") + [
        Path("C:/Windows/Fonts/times.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("C:/Windows/Fonts/tahoma.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    ]
    for font_path in candidates:
        if font_path.exists():
            font_name = f"PetitionsFont_{re.sub(r'[^A-Za-z0-9]+', '', font_path.stem) or 'Default'}"
            try:
                pdfmetrics.registerFont(TTFont(font_name, str(font_path)))  # type: ignore[attr-defined]
                return font_name
            except Exception:
                continue
    return "Times-Roman"


def _preferred_font_candidates(family: str) -> list[Path]:
    normalized = family.strip().casefold()
    if not normalized:
        return []
    windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    known_windows = {
        "times new roman": ["times.ttf"],
        "arial": ["arial.ttf"],
        "calibri": ["calibri.ttf"],
        "tahoma": ["tahoma.ttf"],
        "georgia": ["georgia.ttf"],
        "cambria": ["cambria.ttc", "cambria.ttf"],
    }
    candidates = [windows_fonts / name for name in known_windows.get(normalized, [])]
    token = re.sub(r"[^a-z0-9]+", "", normalized)
    if token:
        candidates.extend(sorted(windows_fonts.glob(f"*{token}*.ttf")))
        candidates.extend(sorted(windows_fonts.glob(f"*{token}*.ttc")))
    seen: set[str] = set()
    unique: list[Path] = []
    for path in candidates:
        key = str(path).casefold()
        if key not in seen:
            seen.add(key)
            unique.append(path)
    return unique


def _markdown_to_reportlab_story(text: str, styles: object, Spacer: object) -> list[object]:
    story: list[object] = []
    in_front_matter = False
    front_matter_done = False
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line == "---" and not front_matter_done:
            in_front_matter = not in_front_matter
            if not in_front_matter:
                front_matter_done = True
            continue
        if in_front_matter or not line:
            continue
        story.extend(_markdown_line_to_flowables(line, styles, Spacer))
    if not story:
        from reportlab.platypus import Paragraph  # type: ignore

        story.append(Paragraph("[blank separator]", styles["SeparatorBody"]))  # type: ignore[index]
    return story


def _markdown_line_to_flowables(line: str, styles: object, Spacer: object) -> list[object]:
    from reportlab.platypus import Paragraph  # type: ignore

    if line.startswith("# "):
        return [Paragraph(_inline_markdown_to_html(line[2:]), styles["SeparatorTitle"])]  # type: ignore[index]
    if line.startswith("## "):
        return [Paragraph(_inline_markdown_to_html(line[3:]), styles["SeparatorHeading"])]  # type: ignore[index]
    if line.startswith("### "):
        return [Paragraph(_inline_markdown_to_html(line[4:]), styles["SeparatorEpisode"])]  # type: ignore[index]
    if line.strip("*").casefold() == "please see the next page":
        return [Paragraph("<i>Please see the next page</i>", styles["SeparatorPlease"])]  # type: ignore[index]
    if re.match(r"^\d+(?:\.\d+)*\.\s+", line):
        return [Paragraph(_inline_markdown_to_html(line), styles["SeparatorBody"])]  # type: ignore[index]
    if line.startswith("- "):
        return [Paragraph("• " + _inline_markdown_to_html(line[2:]), styles["SeparatorBody"])]  # type: ignore[index]
    if line.startswith("**"):
        return [Paragraph(_inline_markdown_to_html(line), styles["SeparatorMeta"])]  # type: ignore[index]
    return [Paragraph(_inline_markdown_to_html(line), styles["SeparatorBody"])]  # type: ignore[index]


def _inline_markdown_to_html(value: str) -> str:
    escaped = escape(value)
    escaped = re.sub(r"`([^`]+)`", r"<font name='Courier'>\1</font>", escaped)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", escaped)
    return escaped
