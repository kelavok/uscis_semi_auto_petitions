import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree

from docx import Document
from reportlab.pdfgen import canvas

from app import evidence, workflow
from stage0_support import SyntheticCase


class EvidenceRegressionTests(SyntheticCase):
    def test_malformed_index_is_preserved_instead_of_silently_resetting_metadata(self):
        self.create(); self.source(); evidence.scan_documents(self.case_id)
        index = self.case_dir / "indexes/document_index.csv"
        previous = index.read_bytes()
        for malformed in ("wrong_header\nmanual title\n", "document_id,file_path\nDOC0001,source.txt,extra\n", "document_id,file_path\nDOC0001,source.txt\nDOC0001,other.txt\n"):
            index.write_text(malformed, encoding="utf-8")
            with self.assertRaises(ValueError):
                evidence.scan_documents(self.case_id)
            self.assertEqual(index.read_text(), malformed)
        index.write_bytes(previous)
        evidence.scan_documents(self.case_id)
        self.assertEqual(self.rows()[0]["document_id"], "DOC0001")

    def test_rescan_reuses_valid_cache_and_repairs_corrupt_cache(self):
        self.create(); self.source(); evidence.scan_documents(self.case_id)
        row = self.rows()[0]
        extracted = self.case_dir / row["text_extraction_path"]
        with patch.object(evidence, "_extract_text", side_effect=AssertionError("Unchanged source was re-extracted")):
            evidence.scan_documents(self.case_id)
        extracted.write_text("CORRUPTED CACHE", encoding="utf-8")
        prompt = workflow.build_prompt(self.case_id, "criterion_awards_episode", [], episode_id="1", episode_folder="1. Award")
        self.assertIn("extracted text is missing or changed", prompt.read_text(encoding="utf-8"))
        evidence.scan_documents(self.case_id)
        self.assertIn("Synthetic award evidence", extracted.read_text())
        self.assertNotIn("CORRUPTED", extracted.read_text())
        self.assertEqual(self.rows()[0]["document_id"], row["document_id"])

    def test_changed_locked_source_is_flagged_without_losing_manual_fields(self):
        self.create(); source = self.source(); evidence.scan_documents(self.case_id)
        rows = self.rows()
        rows[0].update(manual_edit_lock="true", display_title="Manual title", category="manual_category", exhibit_number="7", user_approval_status="approved")
        evidence._write_index(self.case_dir / "indexes/document_index.csv", rows)
        source.write_text("CHANGED SOURCE", encoding="utf-8")
        evidence.scan_documents(self.case_id)
        row = self.rows()[0]
        for key in ("display_title", "category", "exhibit_number", "user_approval_status"):
            self.assertEqual(row[key], rows[0][key])
        self.assertEqual(row["extraction_status"], "source_changed_locked")
        prompt = workflow.build_prompt(self.case_id, "criterion_awards_episode", [], episode_id="1", episode_folder="1. Award")
        prompt_text = prompt.read_text(encoding="utf-8")
        self.assertIn("source is missing or changed", prompt_text)
        self.assertNotIn("Synthetic award evidence", prompt_text)
        row["manual_edit_lock"] = "false"
        evidence._write_index(self.case_dir / "indexes/document_index.csv", [row])
        evidence.scan_documents(self.case_id)
        self.assertEqual(self.rows()[0]["extraction_status"], "text_extracted")

    def test_rename_move_delete_restore_preserve_ids_manual_fields_and_links(self):
        self.create(); original = self.source(); self.source(translation=True)
        evidence.scan_documents(self.case_id); evidence.link_translations(self.case_id)
        rows = self.rows(); original_id = rows[0]["document_id"]
        rows[0].update(display_title="Manual title", exhibit_number="5", user_approval_status="approved")
        evidence._write_index(self.case_dir / "indexes/document_index.csv", rows)
        moved = original.parent.parent / "2. Moved" / "renamed.txt"
        moved.parent.mkdir(); original.rename(moved)
        evidence.scan_documents(self.case_id)
        after = self.rows()
        self.assertEqual(len(after), 2)
        self.assertEqual(after[0]["document_id"], original_id)
        self.assertEqual(after[0]["display_title"], "Manual title")
        self.assertEqual(after[1]["parent_document_id"], original_id)
        saved = moved.read_bytes(); moved.unlink()
        evidence.scan_documents(self.case_id)
        self.assertEqual(self.rows()[0]["extraction_status"], "source_missing")
        moved.write_bytes(saved); evidence.scan_documents(self.case_id)
        self.assertEqual(self.rows()[0]["document_id"], original_id)
        self.assertEqual(self.rows()[0]["extraction_status"], "text_extracted")

    def test_rescan_keeps_explicit_unlink_and_manual_relink_decisions(self):
        self.create(); self.source(); self.source(translation=True)
        evidence.scan_documents(self.case_id); evidence.link_translations(self.case_id)
        rows = self.rows(); original_id, translation_id = rows[0]["document_id"], rows[1]["document_id"]
        evidence.unlink_translation(self.case_id, translation_id)
        evidence.scan_documents(self.case_id); evidence.link_translations(self.case_id)
        self.assertEqual(self.rows()[1]["parent_document_id"], "")
        self.assertEqual(self.rows()[1]["translation_link_mode"], "unlinked")
        evidence.manual_link_translation(self.case_id, translation_id, original_id)
        evidence.scan_documents(self.case_id); evidence.link_translations(self.case_id)
        self.assertEqual(self.rows()[1]["parent_document_id"], original_id)

    def test_interrupted_scan_preserves_index_and_next_scan_recovers(self):
        self.create(); source = self.source(); evidence.scan_documents(self.case_id)
        index = self.case_dir / "indexes/document_index.csv"
        before = index.read_bytes()
        source.write_text("UPDATED synthetic evidence", encoding="utf-8")
        with patch.object(evidence, "_write_index", side_effect=OSError("Synthetic interrupted scan")):
            with self.assertRaises(OSError):
                evidence.scan_documents(self.case_id)
        self.assertEqual(index.read_bytes(), before)
        prompt = workflow.build_prompt(self.case_id, "criterion_awards_episode", [], episode_id="1", episode_folder="1. Award")
        self.assertIn("extracted text is missing or changed", prompt.read_text(encoding="utf-8"))
        evidence.scan_documents(self.case_id)
        self.assertEqual(self.rows()[0]["document_id"], "DOC0001")
        self.assertIn("UPDATED", (self.case_dir / self.rows()[0]["text_extraction_path"]).read_text())

    def test_editing_manual_description_is_reflected_on_rescan(self):
        self.create(); self.source("scan.pdf", "invalid PDF")
        evidence.scan_documents(self.case_id)
        description = self.case_dir / "manual_descriptions/DOC0001.md"
        description.write_text("# Description\nSynthetic first description", encoding="utf-8")
        evidence.scan_documents(self.case_id)
        self.assertEqual(self.rows()[0]["extraction_status"], "manual_description_available")
        description.write_text("# Description\nSynthetic UPDATED description", encoding="utf-8")
        evidence.scan_documents(self.case_id)
        prompt = workflow.build_prompt(self.case_id, "criterion_awards_episode", [], episode_id="1", episode_folder="1. Award")
        self.assertIn("UPDATED description", prompt.read_text(encoding="utf-8"))

    def test_full_pdf_extraction_exceeds_old_page_and_character_limits(self):
        pdf = self.root / "long.pdf"
        document = canvas.Canvas(str(pdf))
        for page in range(25):
            for line in range(100):
                document.drawString(36, 750 - (line % 70) * 9, f"Synthetic page {page} line {line}: " + "Evidence " * 12)
            document.showPage()
        document.save()
        text, status, note = evidence._extract_pdf_text(pdf)
        self.assertEqual(status, "text_extracted")
        self.assertIn("Synthetic page 24 line 99", text)
        self.assertGreater(len(text), 200_000)
        self.assertEqual(note, "")

    def test_image_only_corrupt_pdf_and_malformed_docx_have_explicit_status(self):
        self.create()
        blank = self.source("blank.pdf")
        document = canvas.Canvas(str(blank)); document.showPage(); document.save()
        self.source("broken.pdf", "invalid PDF")
        broken_docx = self.source("broken.docx")
        with zipfile.ZipFile(broken_docx, "w") as archive:
            archive.writestr("word/document.xml", "<broken")
        evidence.scan_documents(self.case_id)
        statuses = {row["original_file_name"]: row["extraction_status"] for row in self.rows()}
        self.assertEqual(statuses["blank.pdf"], "pdf_no_extractable_text")
        self.assertEqual(statuses["broken.pdf"], "pdf_extraction_error")
        self.assertEqual(statuses["broken.docx"], "docx_invalid_or_corrupt")

    def test_docx_extraction_includes_table_header_footer_and_preserves_entities(self):
        path = self.root / "synthetic.docx"
        document = Document()
        document.add_paragraph("Literal &lt; value & unicode Ж")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic table"
        document.sections[0].header.paragraphs[0].text = "Synthetic header"
        document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
        document.add_paragraph("A" * 510_000 + "TAIL-MARKER")
        document.save(path)
        text = workflow.extract_docx_text(path)
        for marker in ("Literal &lt; value & unicode Ж", "Synthetic table", "Synthetic header", "Synthetic footer", "TAIL-MARKER"):
            self.assertIn(marker, text)
