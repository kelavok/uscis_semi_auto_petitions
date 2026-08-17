from __future__ import annotations

import importlib.util
import shutil
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document

from app import cli_support, web
from app.document_layout import (
    build_original_directory_catalog,
    build_layout_bundle,
    build_layout_preview,
    list_installed_fonts,
    load_layout_status,
    refresh_layout_sources,
    save_folder_scopes,
    set_mapping_paths,
)
from app.progress import build_case_progress


PDF_STACK_AVAILABLE = bool(importlib.util.find_spec("reportlab") and importlib.util.find_spec("pypdf"))


@unittest.skipUnless(PDF_STACK_AVAILABLE, "PDF extras are required for document layout tests.")
class DocumentLayoutTests(unittest.TestCase):
    def test_document_layout_flow_supports_multi_file_mapping_and_selective_build(self) -> None:
        with TemporaryDirectory() as temp:
            root = Path(temp)
            case_root = root / "case_workspace"
            template_root = case_root / "_template"
            shutil.copytree(cli_support.CASE_TEMPLATE_ROOT, template_root)

            originals = root / "incoming" / "originals"
            translations = root / "incoming" / "translations"
            (originals / "criterion_1" / "episode_award").mkdir(parents=True)
            (originals / "support_letters").mkdir(parents=True)
            translations.mkdir(parents=True)
            _make_one_page_pdf(originals / "criterion_1/episode_award/award_certificate.pdf", "award_certificate.pdf")
            _make_one_page_pdf(originals / "criterion_1/episode_award/award_appendix.pdf", "award_appendix.pdf")
            _make_one_page_pdf(originals / "criterion_1/episode_award/jury_letter.pdf", "jury_letter.pdf")
            _make_one_page_pdf(originals / "support_letters/support_letter_a.pdf", "support_letter_a.pdf")
            _make_one_page_pdf(originals / "support_letters/support_letter_b.pdf", "support_letter_b.pdf")
            _make_one_page_pdf(translations / "jury_letter_translation.pdf", "jury translation")

            list_path = root / "layout_list.docx"
            _make_layout_docx(list_path)
            available_fonts = list_installed_fonts()
            preferred_font = (
                "Times New Roman"
                if "Times New Roman" in available_fonts
                else (available_fonts[0] if available_fonts else "Times New Roman")
            )

            with (
                patch.object(cli_support, "CASE_ROOT", case_root),
                patch.object(cli_support, "CASE_TEMPLATE_ROOT", template_root),
                patch.object(web, "CASE_ROOT", case_root),
            ):
                cli_support.create_case_from_template("layout_001", "document_layout")
                summary = refresh_layout_sources(
                    "layout_001",
                    originals_dir=str(originals),
                    translations_dir=str(translations),
                    list_document_path=str(list_path),
                    font_family=preferred_font,
                )
                status = load_layout_status("layout_001")
                save_folder_scopes(
                    "layout_001",
                    {
                        "exhibit_folder_exhibit_001": "criterion_1",
                        "episode_folder_episode_001_001": "criterion_1/episode_award",
                        "exhibit_folder_exhibit_002": "support_letters",
                    },
                )
                status = load_layout_status("layout_001")
                directory_catalog = build_original_directory_catalog("layout_001")

                first_episode = status.structure["exhibits"][0]["episodes"][0]
                award_doc = first_episode["documents"][0]
                jury_doc = first_episode["documents"][1]
                support_doc_a = status.structure["exhibits"][1]["episodes"][0]["documents"][0]
                support_doc_b = status.structure["exhibits"][1]["episodes"][0]["documents"][1]

                set_mapping_paths(
                    "layout_001",
                    award_doc["id"],
                    "original",
                    [
                        str(originals / "criterion_1/episode_award/award_certificate.pdf"),
                        str(originals / "criterion_1/episode_award/award_appendix.pdf"),
                    ],
                )
                set_mapping_paths(
                    "layout_001",
                    jury_doc["id"],
                    "original",
                    [str(originals / "criterion_1/episode_award/jury_letter.pdf")],
                )
                set_mapping_paths(
                    "layout_001",
                    jury_doc["id"],
                    "translation",
                    [str(translations / "jury_letter_translation.pdf")],
                )
                set_mapping_paths(
                    "layout_001",
                    support_doc_a["id"],
                    "original",
                    [str(originals / "support_letters/support_letter_a.pdf")],
                )
                set_mapping_paths(
                    "layout_001",
                    support_doc_b["id"],
                    "original",
                    [str(originals / "support_letters/support_letter_b.pdf")],
                )

                preview = build_layout_preview("layout_001")
                bundle = build_layout_bundle("layout_001", ["1"], [])
                progress = build_case_progress("layout_001")
                home_html = web.render_home("")
                case_html = web.render_case_page("layout_001", {})

            self.assertEqual(summary.exhibits, 2)
            self.assertEqual(summary.documents, 4)
            self.assertEqual(status.settings["font_family"], preferred_font)
            self.assertEqual(status.structure["exhibits"][0]["source_folder"], "criterion_1")
            self.assertEqual(status.structure["exhibits"][0]["episodes"][0]["source_folder"], "criterion_1/episode_award")
            self.assertEqual(status.structure["exhibits"][1]["source_folder"], "support_letters")
            self.assertEqual(
                {item["directory"] for item in directory_catalog},
                {"criterion_1", "criterion_1/episode_award", "support_letters"},
            )
            self.assertEqual(status.structure["exhibits"][0]["episodes"][0]["documents"][0]["number"], "1.1.1")
            self.assertEqual(status.structure["exhibits"][0]["episodes"][0]["documents"][1]["number"], "1.1.2")
            self.assertEqual(status.structure["exhibits"][1]["episodes"][0]["documents"][0]["number"], "2.1")
            self.assertEqual(status.structure["exhibits"][1]["episodes"][0]["documents"][1]["number"], "2.2")
            self.assertTrue(preview.pdf_path.exists())
            self.assertTrue(bundle.pdf_path.exists())
            self.assertEqual(bundle.exhibits, 1)
            self.assertEqual(bundle.documents, 2)
            self.assertEqual(bundle.source_files, 4)
            self.assertEqual(progress.task_type, "document_layout")
            self.assertIn("Document layout", home_html)
            self.assertIn("/document-layout?case=layout_001", case_html)


def _make_layout_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Exhibit 1 Criterion i. Awards")
    document.add_paragraph("National Award Packet")
    document.add_paragraph("(1. Award certificate")
    document.add_paragraph("2. Jury confirmation)")
    document.add_paragraph("Exhibit 2 Support Letters")
    document.add_paragraph("(1. Letter from Company A")
    document.add_paragraph("2. Letter from Company B)")
    document.save(path)


def _make_one_page_pdf(path: Path, title: str) -> None:
    from reportlab.lib.pagesizes import A4  # type: ignore
    from reportlab.pdfgen import canvas  # type: ignore

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setFont("Helvetica", 18)
    pdf.drawString(72, 760, title)
    pdf.showPage()
    pdf.save()
