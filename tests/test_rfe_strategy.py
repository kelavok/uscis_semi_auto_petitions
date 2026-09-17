import json
import shutil
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document

from app import cli_support, web
from app.bundle_workflow import inspect_layout_index_status, refresh_layout_indexes
from app.evidence import scan_documents
from app.memo_builder import build_working_memo
from app.rfe_strategy import (
    apply_strategy_output,
    build_strategy_bootstrap_prompt,
    import_evidence_and_scan_inputs,
)
from app.stages import build_llm_stage
from app.workflow import build_prompt


class RfeStrategyTests(unittest.TestCase):
    def test_migrator_variant_uses_three_source_bootstrap_and_front_index(self) -> None:
        with TemporaryDirectory() as temp:
            root = Path(temp)
            case_root = root / "case_workspace"
            template_root = case_root / "_template"
            shutil.copytree(cli_support.CASE_TEMPLATE_ROOT, template_root)
            strategy = root / "strategy.txt"
            rfe = root / "rfe.txt"
            initial = root / "initial.txt"
            strategy.write_text("Challenge only the awards finding; omit industry and employment.", encoding="utf-8")
            rfe.write_text("Chief Jane Smith, Officer 1234. Awards were not accepted; judging was accepted.", encoding="utf-8")
            initial.write_text("Initial filing awards argument and Initial Filing Exhibit 2.1.", encoding="utf-8")

            with (
                patch.object(cli_support, "CASE_ROOT", case_root),
                patch.object(cli_support, "CASE_TEMPLATE_ROOT", template_root),
                patch.object(web, "CASE_ROOT", case_root),
            ):
                case_dir = cli_support.create_case_from_template("migrator_rfe", "eb1a_rfe_response")
                from app.memo_builder import apply_case_intake

                apply_case_intake(
                    "migrator_rfe",
                    {
                        "beneficiary_full_name": "Ivan Ivanov",
                        "preferred_reference": "Mr. Ivanov",
                        "field": "Technology",
                        "specialization": "AI product engineering",
                        "case_number": "IOE1234567890",
                        "eb1a_rfe_template_variant": "migrator",
                    },
                )
                bootstrap = build_strategy_bootstrap_prompt(
                    "migrator_rfe", str(strategy), str(rfe), str(initial)
                )
                prompt_text = bootstrap.prompt_path.read_text(encoding="utf-8")
                self.assertIn("Full initial-filing memorandum", prompt_text)
                self.assertIn("Initial filing awards argument", prompt_text)
                self.assertIn("accepted_criteria_count", prompt_text)
                self.assertIn("officer_number", prompt_text)
                self.assertIn("EB1A_RFE_response_migrator_LLM_template.yaml", prompt_text)
                self.assertIn('"case_number": "IOE1234567890"', prompt_text)

                output = {
                    "case_id": "migrator_rfe",
                    "task_type": "eb1a_rfe_response",
                    "case_metadata": {
                        "beneficiary_full_name": "Ivan Ivanov",
                        "preferred_reference": "Mr. Ivanov",
                        "field": "Technology",
                        "specialization": "AI product engineering",
                        "case_number": "IOE1234567890",
                        "receipt_date": "January 1, 2026",
                        "rfe_date": "August 1, 2026",
                        "response_deadline": "October 1, 2026",
                        "uscis_address": "USCIS Test Address",
                        "rfe_response_date": "September 1, 2026",
                        "uscis_office_or_service_center": "Texas Service Center",
                        "officer_name": "",
                        "office_chief_name": "Ms. Jane Smith",
                        "officer_number": "1234",
                        "salutation": "Dear Ms. Jane Smith and Officer 1234:",
                        "submitter_name": "Attorney Alex Doe",
                        "submitter_title": "Attorney",
                        "petition_type": "EB-1A Form I-140",
                    },
                    "global_strategy": "Answer only the disputed national-recognition element.",
                    "accepted_criteria": ["judging"],
                    "accepted_criteria_count": 1,
                    "challenged_criteria": ["awards"],
                    "challenged_criteria_count": 1,
                    "rfe_issues": [
                        {
                            "issue_id": "awards_recognition",
                            "topic": "National recognition of the award",
                            "issue_type": "criterion",
                            "status": "not_accepted",
                            "criterion_role": "awards",
                            "exact_rfe_quote": "Awards were not accepted because national recognition was not established.",
                            "defect": "National recognition was not established.",
                            "response_strategy": "Use independent organizer and media evidence.",
                            "elements_not_disputed_do_not_discuss": ["Receipt of the award"],
                            "evidence_actions": [],
                        }
                    ],
                    "sections": [
                        {
                            "section_id": "cover",
                            "order": 1,
                            "title": "Cover Letter",
                            "section_type": "cover_letter",
                            "required": True,
                            "strategy": "Summarize the findings.",
                            "rfe_issue_ids": [],
                            "starter_text": "",
                            "episodes": [],
                        },
                        {
                            "section_id": "awards",
                            "order": 2,
                            "title": "Awards",
                            "section_type": "criterion",
                            "criterion_role": "awards",
                            "required": True,
                            "strategy": "Rebut only national recognition.",
                            "rfe_issue_ids": ["awards_recognition"],
                            "starter_text": "",
                            "episodes": [
                                {
                                    "episode_id": "award_one",
                                    "title": "National Technology Award",
                                    "source_folder": "1. Награды/1 episode",
                                    "strategy": "Lead with independent recognition.",
                                    "rfe_issue_ids": ["awards_recognition"],
                                    "planned_subheadings": ["The Award Has National Recognition"],
                                }
                            ],
                        },
                    ],
                    "template_decisions": {
                        "structure_rationale": "Front index, cover letter, challenged criterion.",
                        "include_general_response": False,
                        "include_industry_overview": False,
                        "include_employment_section": False,
                        "include_final_merits": False,
                        "include_recommendation_letters": False,
                        "attachment_label": "Exhibit",
                    },
                    "open_questions": [],
                }
                apply_strategy_output("migrator_rfe", output)
                (case_dir / "indexes/document_index.csv").write_text(
                    "document_id,original_file_name,display_title,file_path,parent_document_id,translation_status,relationship_type,episode_title\n"
                    "DOC0001,award.pdf,Independent award confirmation,source_documents/rfe_response/new_documents/originals/1. Награды/1 episode/award.pdf,,,,National Technology Award\n",
                    encoding="utf-8-sig",
                )
                (case_dir / "indexes/exhibit_index.csv").write_text(
                    "exhibit_number,display_title,document_ids\n"
                    "1,Awards,DOC0001\n",
                    encoding="utf-8-sig",
                )
                memo = build_working_memo("migrator_rfe")
                html = web.render_intake_panel("migrator_rfe", "eb1a_rfe_response")

            saved_config = (case_dir / "case_config.yaml").read_text(encoding="utf-8")
            self.assertIn("eb1a_rfe_template_variant: migrator", saved_config)
            self.assertIn("Шаблон ответа на RFE EB-1A ver.1.0.docx", saved_config)
            self.assertIn("EB-1A RFE response template", html)
            self.assertIn(">Мигратор</option>", html)
            document = Document(memo.docx_path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            self.assertEqual(paragraphs[0], "INDEX:")
            self.assertTrue(
                any("Independent award confirmation" in paragraph for paragraph in paragraphs)
            )
            self.assertIn("Dear Ms. Jane Smith and Officer 1234:", paragraphs)
            self.assertIn("the following 1 criterion", " ".join(paragraphs))
            self.assertNotIn("TEXT FORMATTING SETTINGS", " ".join(paragraphs))
            self.assertNotIn("Recommendation Letters", " ".join(paragraphs))
            self.assertNotIn("Attachments / Evidence Index", " ".join(paragraphs))
            normal = next(
                style for style in document.styles if style.name.casefold() == "normal"
            )
            self.assertEqual(normal.font.name, "Times New Roman")

    def test_strategy_manifest_drives_template_units_and_prompt_context(self) -> None:
        with TemporaryDirectory() as temp:
            root = Path(temp)
            case_root = root / "case_workspace"
            template_root = case_root / "_template"
            shutil.copytree(cli_support.CASE_TEMPLATE_ROOT, template_root)
            strategy = root / "strategy.txt"
            rfe = root / "rfe.txt"
            strategy.write_text("Argue each criterion using the mapped evidence.", encoding="utf-8")
            rfe.write_text("Case IOE123. Awards were not accepted.", encoding="utf-8")

            with (
                patch.object(cli_support, "CASE_ROOT", case_root),
                patch.object(cli_support, "CASE_TEMPLATE_ROOT", template_root),
                patch.object(web, "CASE_ROOT", case_root),
            ):
                case_dir = cli_support.create_case_from_template("rfe_test", "eb1a_rfe_response")
                bootstrap = build_strategy_bootstrap_prompt(
                    "rfe_test", str(strategy), str(rfe)
                )
                self.assertIn("Human strategy", bootstrap.prompt_path.read_text(encoding="utf-8"))
                bootstrap_html = web.render_rfe_panel("rfe_test")
                self.assertIn("RFE Stage 1A", bootstrap_html)
                self.assertIn("Generated bootstrap prompt", bootstrap_html)
                output = {
                    "case_id": "rfe_test",
                    "task_type": "eb1a_rfe_response",
                    "case_metadata": {
                        "beneficiary_full_name": "Jane Doe",
                        "preferred_reference": "Ms. Doe",
                        "field": "Business",
                        "specialization": "Software Engineering",
                        "case_number": "IOE123",
                        "receipt_date": "January 1, 2026",
                        "rfe_date": "May 1, 2026",
                        "response_deadline": "July 1, 2026",
                        "uscis_address": "USCIS Test Address",
                        "rfe_response_date": "June 15, 2026",
                        "uscis_office_or_service_center": "Nebraska Service Center",
                        "officer_name": "Jane Officer",
                        "office_chief_name": "",
                        "salutation": "Dear Ms. Jane Officer and Officer:",
                        "submitter_name": "Jane Doe",
                        "submitter_title": "Petitioner and Beneficiary",
                        "petition_type": "EB-1A Form I-140",
                    },
                    "global_strategy": "Answer the exact RFE defect and distinguish commentary from evidence.",
                    "accepted_criteria": ["judging"],
                    "challenged_criteria": ["awards"],
                    "rfe_issues": [
                        {
                            "issue_id": "awards_legitimacy",
                            "topic": "Awards",
                            "issue_type": "criterion",
                            "status": "not_accepted",
                            "criterion_role": "awards",
                            "exact_rfe_quote": "Awards were not accepted because national recognition was not established.",
                            "defect": "Recognition was not established.",
                            "response_strategy": "Use independent recognition evidence.",
                            "elements_not_disputed_do_not_discuss": ["Receipt of the award"],
                            "evidence_actions": ["Resubmit the award regulations."],
                        }
                    ],
                    "sections": [
                        {
                            "section_id": "cover",
                            "order": 1,
                            "title": "Cover Letter",
                            "section_type": "cover_letter",
                            "required": True,
                            "strategy": "Summarize the response order.",
                            "rfe_issue_ids": [],
                            "starter_text": "Dear USCIS Officer:",
                            "episodes": [],
                        },
                        {
                            "section_id": "awards",
                            "order": 2,
                            "title": "Awards",
                            "section_type": "criterion",
                            "criterion_role": "awards",
                            "required": True,
                            "strategy": "Lead with independent evidence.",
                            "rfe_issue_ids": ["awards_legitimacy"],
                            "starter_text": "",
                            "initial_filing_heading_hints": ["Awards criterion"],
                            "episodes": [
                                {
                                    "episode_id": "award_nba",
                                    "title": "Technologies and Innovations Award",
                                    "source_folder": "1. Награды/1 episode",
                                    "strategy": "Explain the selection process.",
                                    "rfe_issue_ids": ["awards_legitimacy"],
                                    "planned_subheadings": [
                                        "The Award Has National Recognition in the Field"
                                    ],
                                    "initial_filing_heading_hints": [],
                                }
                            ],
                        },
                    ],
                    "template_decisions": {
                        "structure_rationale": "Cover first, then the challenged criterion.",
                        "include_general_response": False,
                        "include_employment_section": False,
                        "include_final_merits": False,
                        "attachment_label": "Attachment",
                    },
                    "open_questions": [],
                }
                imported = apply_strategy_output("rfe_test", output)
                self.assertEqual(imported.unit_count, 2)
                accepted_html = web.render_rfe_panel("rfe_test")
                self.assertIn("Accepted: 2 drafting unit(s)", accepted_html)
                self.assertIn("RFE Stage 1C", accepted_html)

                initial_memo = root / "initial_memo.txt"
                initial_memo.write_text(
                    "Documentation of receipt of lesser nationally recognized prizes or awards\n\n"
                    "The beneficiary received the technology award.\n\n"
                    "Documentation of membership in associations\n\nMembership evidence.",
                    encoding="utf-8",
                )
                new_docs = root / "new_docs" / "originals" / "1. Награды" / "1 episode"
                new_docs.mkdir(parents=True)
                (new_docs / "organizer_letter.txt").write_text("New organizer letter", encoding="utf-8")
                (new_docs / "extracts.txt").write_text(
                    "The attached ceremony screenshot identifies Jane Doe as the category winner.",
                    encoding="utf-8",
                )
                (new_docs / "info.txt").write_text(
                    "LLM instruction: emphasize the independent organizer confirmation first.",
                    encoding="utf-8",
                )
                info_docx = Document()
                info_docx.add_paragraph("DOCX guidance: rely on the organizer letter, not this note.")
                info_docx.save(new_docs / "info.docx")
                (new_docs / "README.md").write_text(
                    "Use the ceremony screenshot and organizer letter as one evidence set.",
                    encoding="utf-8",
                )
                (new_docs / "broken_scan.docx").write_bytes(
                    b"This is an image export incorrectly named as a DOCX file."
                )
                (new_docs / "citation.xlsx").write_bytes(b"unsupported spreadsheet placeholder")
                (new_docs / "~$open_in_word.docx").write_bytes(b"Office lock file")
                evidence = import_evidence_and_scan_inputs(
                    "rfe_test",
                    str(initial_memo),
                    str(root / "new_docs"),
                )
                self.assertGreaterEqual(evidence.initial_sections, 1)
                scan_documents("rfe_test")
                document_index = (case_dir / "indexes/document_index.csv").read_text(
                    encoding="utf-8-sig"
                )
                self.assertNotIn("extracts.txt", document_index)
                self.assertNotIn("info.txt", document_index)
                self.assertNotIn("info.docx", document_index)
                self.assertNotIn("README.md", document_index)
                self.assertIn("broken_scan.docx", document_index)
                self.assertIn("docx_invalid_or_corrupt", document_index)
                self.assertNotIn("~$open_in_word.docx", document_index)
                stage = build_llm_stage("rfe_test")
                self.assertNotIn("cover", [unit.episode_id for unit in stage.units])
                award_unit = next(unit for unit in stage.units if unit.criterion == "awards")
                self.assertTrue(award_unit.selected_documents)
                prompt_path = build_prompt(
                    "rfe_test",
                    "rfe_dynamic_section",
                    [],
                    episode_id=award_unit.episode_id,
                    episode_folder="1. Награды/1 episode",
                )
                prompt_text = prompt_path.read_text(encoding="utf-8")
                self.assertIn("Explain the selection process", prompt_text)
                self.assertIn("The beneficiary received the technology award", prompt_text)
                self.assertIn("organizer_letter.txt", prompt_text)
                self.assertIn("The attached ceremony screenshot identifies Jane Doe", prompt_text)
                self.assertIn("prompt-only folder sidecar (extracts.txt)", prompt_text)
                self.assertIn("do not add to document/exhibit indexes", prompt_text)
                self.assertIn("emphasize the independent organizer confirmation first", prompt_text)
                self.assertIn("prompt-only folder sidecar (info.txt)", prompt_text)
                self.assertIn("DOCX guidance: rely on the organizer letter", prompt_text)
                self.assertIn("prompt-only folder sidecar (info.docx)", prompt_text)
                self.assertIn("additional instructions and explanations", prompt_text)
                self.assertIn("Use the ceremony screenshot and organizer letter as one evidence set", prompt_text)
                self.assertIn("prompt-only folder sidecar (README.md)", prompt_text)
                self.assertIn("Technical document selection for this unit", prompt_text)
                self.assertIn("must be copied exactly", prompt_text)

                selected_document_id = next(
                    document_id
                    for document_id, title in award_unit.selected_documents
                    if "organizer" in title.casefold()
                )
                spreadsheet_document_id = next(
                    document_id
                    for document_id, title in award_unit.selected_documents
                    if "citation" in title.casefold()
                )
                validated_path = case_dir / "validated_outputs" / f"rfe_dynamic_section.{award_unit.episode_id}.json"
                validated_path.write_text(
                    json.dumps(
                        {
                            "case_id": "rfe_test",
                            "task_type": "eb1a_rfe_response",
                            "step_id": "rfe_dynamic_section",
                            "episode_id": award_unit.episode_id,
                            "draft_text": "Awards response.",
                            "used_documents": [
                                {
                                    "document_id": document_id,
                                    "document_title": (
                                        "Organizer letter"
                                        if document_id == selected_document_id
                                        else "Citation table"
                                        if document_id == spreadsheet_document_id
                                        else title
                                    ),
                                    "used_for": "award evidence",
                                }
                                for document_id, title in award_unit.selected_documents
                            ],
                            "unsupported_claims": [],
                            "questions_for_user": [],
                            "quality_flags": [],
                            "revision_notes": [],
                        },
                        ensure_ascii=False,
                    ),
                    encoding="utf-8",
                )
                before_refresh = inspect_layout_index_status("rfe_test")
                self.assertTrue(before_refresh.refresh_required)
                before_layout = web.render_layout_page("rfe_test", {})
                self.assertIn("Index refresh required", before_layout)
                self.assertIn("Refresh indexes", before_layout)
                next(case_dir.rglob("citation.xlsx")).with_suffix(".pdf").write_bytes(
                    b"%PDF-1.4\n%%EOF\n"
                )
                refreshed = refresh_layout_indexes("rfe_test")
                self.assertEqual(refreshed.replacement_documents_rebound, 1)
                self.assertEqual(refreshed.status.assigned_used_documents, 3)
                self.assertEqual(refreshed.status.exhibit_count, 1)
                self.assertTrue(refreshed.status.ready_for_separators)
                self.assertFalse(refreshed.status.unsupported_documents)
                layout_page = web.render_layout_page("rfe_test", {})
                self.assertIn("Evidence indexes ready", layout_page)
                self.assertIn("Select & prepare", layout_page)
                self.assertIn("Prepare selected bundle", layout_page)
                self.assertIn("Build selected PDF", layout_page)
                refreshed_index = (case_dir / "indexes/document_index.csv").read_text(
                    encoding="utf-8-sig"
                )
                selected_row = next(
                    line for line in refreshed_index.splitlines() if line.startswith(f"{selected_document_id},")
                )
                self.assertIn(",1,", selected_row)
                spreadsheet_row = next(
                    line
                    for line in refreshed_index.splitlines()
                    if line.startswith(f"{spreadsheet_document_id},")
                )
                self.assertIn("citation.pdf", spreadsheet_row)

                stage = build_llm_stage("rfe_test")
                self.assertEqual([unit.criterion for unit in stage.units], ["awards"])
                self.assertNotEqual(stage.units[0].episode_id, "award_nba")
                draft_path = (
                    case_dir
                    / "draft_sections/rfe/sections"
                    / f"{stage.units[0].episode_id}.md"
                )
                draft_path.parent.mkdir(parents=True, exist_ok=True)
                draft_path.write_text("Awards response.\n", encoding="utf-8")
                memo = build_working_memo("rfe_test")
                with zipfile.ZipFile(memo.docx_path) as archive:
                    document_xml = archive.read("word/document.xml").decode("utf-8")
                self.assertIn("Dear Ms. Jane Officer and Officer:", document_xml)
                self.assertIn("Awards were not accepted because national recognition was not established", document_xml)
                self.assertIn("Awards response.", document_xml)
                self.assertIn("Attachments / Evidence Index", document_xml)
                self.assertIn("Organizer letter", document_xml)
                self.assertNotIn("Drafting direction (internal)", document_xml)
                self.assertNotIn("DRAFTING PLACEHOLDER", document_xml)
                self.assertNotIn("SCRIPT PLACEHOLDER", document_xml)


if __name__ == "__main__":
    unittest.main()
