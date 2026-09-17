import csv
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document

from app import cli_support
from app import workflow as workflow_module
from app.bundle_workflow import EXHIBIT_FIELDS
from app.evidence import INDEX_FIELDS
from app.memo_review import accept_reviewed_memo, validate_reviewed_memo


class ReviewedMemoTests(unittest.TestCase):
    def test_review_is_soft_and_acceptance_syncs_recognized_reordered_index(self) -> None:
        with TemporaryDirectory() as temp:
            root = Path(temp)
            case_root = root / "case_workspace"
            case_dir = case_root / "review_case"
            indexes = case_dir / "indexes"
            final_memo = case_dir / "final_memo"
            workflows = root / "workflows"
            indexes.mkdir(parents=True)
            final_memo.mkdir(parents=True)
            workflows.mkdir(parents=True)
            (case_dir / "case_config.yaml").write_text(
                "case_id: review_case\n"
                "task_type: eb1a_petition\n"
                "workflow: workflows/eb1a_petition.yaml\n"
                "paths:\n"
                "  document_index: indexes/document_index.csv\n"
                "  exhibit_index: indexes/exhibit_index.csv\n"
                "  final_memo: final_memo\n"
                "  bundle_root: bundle\n",
                encoding="utf-8",
            )
            (workflows / "eb1a_petition.yaml").write_text(
                "task_type: eb1a_petition\nsteps: []\n", encoding="utf-8"
            )
            rows = []
            for document_id, title, order in (
                ("DOC0001", "Known Document", "1"),
                ("DOC0002", "Removed Document", "2"),
            ):
                row = {field: "" for field in INDEX_FIELDS}
                row.update(
                    {
                        "document_id": document_id,
                        "original_file_name": title + ".pdf",
                        "display_title": title,
                        "exhibit_number": "1",
                        "final_bundle_order": order,
                    }
                )
                rows.append(row)
            with (indexes / "document_index.csv").open(
                "w", encoding="utf-8-sig", newline=""
            ) as handle:
                writer = csv.DictWriter(handle, fieldnames=INDEX_FIELDS)
                writer.writeheader()
                writer.writerows(rows)
            with (indexes / "exhibit_index.csv").open(
                "w", encoding="utf-8-sig", newline=""
            ) as handle:
                writer = csv.DictWriter(handle, fieldnames=EXHIBIT_FIELDS)
                writer.writeheader()
                exhibit = {field: "" for field in EXHIBIT_FIELDS}
                exhibit.update(
                    {
                        "exhibit_id": "EXH001",
                        "exhibit_number": "1",
                        "display_title": "Media",
                        "document_ids": "DOC0001;DOC0002",
                    }
                )
                writer.writerow(exhibit)

            candidate = root / "edited.docx"
            document = Document()
            document.add_paragraph("INDEX:")
            document.add_paragraph("Exhibit 0-1: Recommendation letters")
            document.add_paragraph("0-1.1. Recommendation Letter from Expert")
            document.add_paragraph("Exhibit 1: Media")
            document.add_paragraph("1.1. Known Document")
            document.add_paragraph("1.2. Mystery Record")
            document.add_paragraph("BODY")
            document.add_paragraph(
                "(Please refer to Exhibit 0-1, page PAGE: 0-1.1 - Recommendation Letter from Expert.)"
            )
            document.add_paragraph(
                "(Please refer to Exhibit 9, page PAGE: 9.9 - Unknown citation.)"
            )
            document.save(candidate)

            with (
                patch.object(cli_support, "CASE_ROOT", case_root),
                patch.object(workflow_module, "PROJECT_ROOT", root),
            ):
                report = validate_reviewed_memo("review_case", str(candidate))
                self.assertEqual(report.documents_matched, 1)
                self.assertTrue(report.unknown_index_items)
                self.assertTrue(report.missing_indexed_documents)
                self.assertTrue(report.unknown_citations)
                accepted = accept_reviewed_memo("review_case")

            self.assertTrue(accepted.accepted)
            self.assertTrue((final_memo / "working_memo.docx").exists())
            with (indexes / "document_index.csv").open(
                "r", encoding="utf-8-sig", newline=""
            ) as handle:
                updated = {row["document_id"]: row for row in csv.DictReader(handle)}
            self.assertEqual(updated["DOC0001"]["exhibit_number"], "1")
            self.assertEqual(updated["DOC0002"]["exhibit_number"], "")


if __name__ == "__main__":
    unittest.main()
