from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def _enum_name(value: object) -> str:
    return getattr(value, "name", str(value) if value is not None else "")


def _points(value: object) -> float | None:
    pt = getattr(value, "pt", None)
    return round(float(pt), 2) if pt is not None else None


def inspect(path: Path) -> dict[str, object]:
    document = Document(path)
    paragraphs: list[dict[str, object]] = []
    for number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text
        p_pr = paragraph._p.pPr
        page_break_before = bool(p_pr is not None and p_pr.find(qn("w:pageBreakBefore")) is not None)
        keep_with_next = bool(p_pr is not None and p_pr.find(qn("w:keepNext")) is not None)
        runs = []
        for run in paragraph.runs:
            if not run.text and not run._element.xpath(".//w:br"):
                continue
            runs.append(
                {
                    "text": run.text,
                    "bold": run.bold,
                    "italic": run.italic,
                    "underline": bool(run.underline),
                    "font": run.font.name,
                    "size_pt": _points(run.font.size),
                    "page_break": bool(run._element.xpath('.//w:br[@w:type="page"]')),
                }
            )
        paragraphs.append(
            {
                "number": number,
                "style": paragraph.style.name if paragraph.style else "",
                "alignment": _enum_name(paragraph.alignment),
                "left_indent_pt": _points(paragraph.paragraph_format.left_indent),
                "first_line_indent_pt": _points(paragraph.paragraph_format.first_line_indent),
                "space_before_pt": _points(paragraph.paragraph_format.space_before),
                "space_after_pt": _points(paragraph.paragraph_format.space_after),
                "page_break_before": page_break_before,
                "keep_with_next": keep_with_next,
                "text": text,
                "runs": runs,
            }
        )

    sections = []
    for section in document.sections:
        sections.append(
            {
                "page_width_in": round(section.page_width.inches, 3),
                "page_height_in": round(section.page_height.inches, 3),
                "top_margin_in": round(section.top_margin.inches, 3),
                "right_margin_in": round(section.right_margin.inches, 3),
                "bottom_margin_in": round(section.bottom_margin.inches, 3),
                "left_margin_in": round(section.left_margin.inches, 3),
                "header_distance_in": round(section.header_distance.inches, 3),
                "footer_distance_in": round(section.footer_distance.inches, 3),
            }
        )

    return {
        "path": str(path),
        "sections": sections,
        "paragraphs": paragraphs,
        "tables": [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ],
    }


def main() -> int:
    if len(sys.argv) not in {2, 3}:
        raise SystemExit("Usage: inspect_docx_layout.py INPUT.docx [--concise]")
    report = inspect(Path(sys.argv[1]))
    if len(sys.argv) == 3 and sys.argv[2] == "--concise":
        for paragraph in report["paragraphs"]:
            runs = paragraph["runs"]
            if not paragraph["text"] and not any(run["page_break"] for run in runs):
                continue
            flags = []
            if any(run["page_break"] for run in runs):
                flags.append("PAGE_BREAK")
            if any(run["bold"] for run in runs):
                flags.append("BOLD")
            if any(run["italic"] for run in runs):
                flags.append("ITALIC")
            if any(run["underline"] for run in runs):
                flags.append("UNDERLINE")
            print(
                f'{paragraph["number"]:03d}\t{paragraph["style"]}\t'
                f'{paragraph["alignment"]}\t{",".join(flags)}\t{paragraph["text"]}'
            )
        print("TABLES")
        print(json.dumps(report["tables"], ensure_ascii=False, indent=2))
    else:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
